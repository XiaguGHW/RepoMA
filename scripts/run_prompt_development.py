#!/usr/bin/env python3
"""Run Task 3 prompt-development classifications on the selected Baugruppen.

Expected project layout
-----------------------
Task3_Prompt_Development/
├── .env
├── llm_connector.py
├── run_prompt_development.py
├── input/classification_experiment_dataset.xlsx
├── prompts/P1.txt
└── outputs/

The dataset is created by build_experiment_dataset.py.  By default this script
processes only rows whose prompt_engineering value is yes.
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import re
import sys
import tempfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import fitz
except ImportError:
    fitz = None

try:
    from dotenv import load_dotenv
except ImportError:
    def load_dotenv(*_args, **_kwargs) -> bool:
        return False

try:
    from tqdm import tqdm
except ImportError:
    def tqdm(iterable, **_kwargs):
        return iterable


PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT_EXCEL = PROJECT_DIR / "input" / "classification_experiment_dataset.xlsx"
DEFAULT_PROMPT_DIR = PROJECT_DIR / "prompts"
DEFAULT_OUTPUT_DIR = PROJECT_DIR / "outputs"
SUPPORTED_EXTENSIONS = {".pdf", ".png", ".jpg", ".jpeg"}

# The Ground Truth workbook still contains two older label names.  Evaluation
# maps them to the Codebook V2 labels used by P1/P2/P3 before comparing.
CANONICAL_CLASS_ORDER = (
    "Lineareinheit",
    "Gantry",
    "Greifer",
    "Umsetzeinheit",
    "Roboter",
    "Rotationseinheit",
    "Keine der verfügbaren Klassen",
)
LABEL_ALIASES = {
    "lineareinheit": "Lineareinheit",
    "gantry": "Gantry",
    "multi achs system gantry": "Gantry",
    "greifer": "Greifer",
    "umsetzeinheit": "Umsetzeinheit",
    "kombinierte einheit": "Umsetzeinheit",
    "umsetzeinheit kombinierte einheit": "Umsetzeinheit",
    "roboter": "Roboter",
    "rotationseinheit": "Rotationseinheit",
    "keine der verfügbaren klassen": "Keine der verfügbaren Klassen",
    "keine der verfugbaren klassen": "Keine der verfügbaren Klassen",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-excel", type=Path, default=DEFAULT_INPUT_EXCEL)
    parser.add_argument("--sheet-name", default="Experiment_Dataset")
    parser.add_argument("--prompt-config", choices=("P1", "P2", "P3"), required=False)
    parser.add_argument(
        "--evaluate-existing",
        type=Path,
        default=None,
        help="Add or refresh Prompt_Evaluation in an existing result workbook without calling an LLM.",
    )
    parser.add_argument("--prompt-file", type=Path, default=None)
    parser.add_argument("--model", default="gemini-2.5-pro")
    parser.add_argument("--temperature", type=float, default=0.0)
    parser.add_argument(
        "--max-output-tokens", type=int, default=4096,
        help="Maximum generated tokens per response. Default: 4096.",
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--max-rows", type=int, default=None)
    parser.add_argument(
        "--all-rows",
        action="store_true",
        help="Run all dataset rows instead of only prompt_engineering = yes.",
    )
    return parser.parse_args()


def find_column(df: pd.DataFrame, expected: str, *, required: bool = True) -> str | None:
    lookup = {str(column).strip().casefold(): str(column) for column in df.columns}
    column = lookup.get(expected.strip().casefold())
    if column or not required:
        return column
    raise ValueError(f"Column '{expected}' is missing. Available: {list(df.columns)}")


def read_prompt(args: argparse.Namespace) -> tuple[Path, str]:
    prompt_path = args.prompt_file or DEFAULT_PROMPT_DIR / f"{args.prompt_config}.txt"
    if not prompt_path.is_file():
        raise FileNotFoundError(f"Prompt file not found: {prompt_path}")
    prompt = prompt_path.read_text(encoding="utf-8").strip()
    if not prompt:
        raise ValueError(f"Prompt file is empty: {prompt_path}")
    return prompt_path, prompt


def collect_supported_files(folder_path: str) -> list[str]:
    folder = Path(folder_path)
    if not folder.is_dir():
        return []
    return [
        str(path)
        for path in sorted(folder.rglob("*"))
        if path.is_file() and path.suffix.casefold() in SUPPORTED_EXTENSIONS
    ]


def prepared_base64_size(connector, file_path: str) -> int:
    """Return the size after the connector has rendered and Base64-encoded a file."""
    converter = getattr(connector, "file_to_openai_images", None)
    if not callable(converter):
        converter = getattr(connector, "_file_to_openai_images", None)
    if not callable(converter):
        raise AttributeError(
            "llm_connector.py must provide _file_to_openai_images() for HTTP 413 retries."
        )
    return sum(len(str(image)) for image in converter(file_path))


def write_short_pdf(source: Path, target: Path, page_count: int) -> None:
    if fitz is None:
        raise ImportError("PyMuPDF (fitz) is required for automatic PDF page reduction.")
    original = fitz.open(source)
    reduced = fitz.open()
    try:
        reduced.insert_pdf(original, from_page=0, to_page=min(page_count, len(original)) - 1)
        reduced.save(target)
    finally:
        reduced.close()
        original.close()


def request_is_too_large(response: object) -> bool:
    text = str(response).casefold()
    return "413" in text or "request entity too large" in text


def prepare_reduced_largest_pdf(
    connector, files: list[str], pages: int, temp_dir: Path
) -> tuple[list[str], list[str], str]:
    """Replace only the largest prepared PDF by a temporary first-N-pages copy."""
    pdf_paths = [Path(file_path) for file_path in files if Path(file_path).suffix.casefold() == ".pdf"]
    if not pdf_paths:
        raise ValueError("The request was too large, but this BG contains no PDF that can be reduced.")
    largest_pdf = max(pdf_paths, key=lambda path: prepared_base64_size(connector, str(path)))
    reduced_pdf = temp_dir / f"{largest_pdf.stem}_first_{pages}_pages.pdf"
    write_short_pdf(largest_pdf, reduced_pdf, pages)
    sent_files = [str(reduced_pdf) if Path(file_path) == largest_pdf else file_path for file_path in files]
    display_files = [
        f"{largest_pdf} [first {pages} pages used instead of first 4]"
        if Path(file_path) == largest_pdf else file_path
        for file_path in files
    ]
    return sent_files, display_files, f"HTTP 413 retry: largest PDF reduced to first {pages} pages."


def build_question(row: pd.Series) -> str:
    english_name = str(row.get("Benennung (E)", "")).strip()
    german_name = str(row.get("Benennung (D)", "")).strip()
    return (
        "Klassifiziere die beigefügte Baugruppe gemäß den Systemanweisungen.\n"
        f"Benennung (E): {english_name}\n"
        f"Benennung (D): {german_name}\n\n"
        "Berücksichtige alle beigefügten technischen Dateien. "
        "Gib ausschließlich das angeforderte JSON-Objekt zurück."
    )


def create_connector(model_name: str, api_key: str):
    try:
        from llm_connector import LLMConnector
    except ImportError as error:
        raise ImportError(
            "llm_connector.py with class LLMConnector was not found. "
            "Place it beside this script."
        ) from error
    return LLMConnector(model_name, api_key)


def parse_json_response(raw_response: object) -> tuple[dict[str, object] | None, str]:
    raw = str(raw_response).strip()
    candidates = [raw]
    if raw.startswith(chr(96) * 3):
        fence = re.escape(chr(96) * 3)
        candidates.append(re.sub(rf"^{fence}(?:json)?\\s*|\\s*{fence}$", "", raw, flags=re.IGNORECASE))
    start, end = raw.find("{"), raw.rfind("}")
    if start >= 0 and end > start:
        candidates.append(raw[start:end + 1])
    for candidate in candidates:
        try:
            parsed = json.loads(candidate)
            return (parsed, "SUCCESS") if isinstance(parsed, dict) else (None, "JSON_IS_NOT_OBJECT")
        except json.JSONDecodeError:
            continue
    return None, "INVALID_JSON"


def json_value(payload: dict[str, object] | None, *keys: str) -> object:
    if payload is None:
        return pd.NA
    normalised = {str(key).casefold(): value for key, value in payload.items()}
    for key in keys:
        if key.casefold() in normalised:
            return normalised[key.casefold()]
    return pd.NA


def normalise_possible_classes(value: object) -> str | object:
    if value is None or value is pd.NA:
        return pd.NA
    if isinstance(value, list):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def make_output_path(args: argparse.Namespace) -> Path:
    model = re.sub(r"[^A-Za-z0-9_.-]+", "_", args.model)
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    return args.output_dir / f"prompt_development_{args.prompt_config}_{model}_{timestamp}.xlsx"


def clean_cell_text(value: object) -> str:
    """Return a safe, stripped string for a spreadsheet cell."""
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def canonical_class_label(value: object) -> str:
    """Map legacy and Codebook V2 label spellings to one comparison label."""
    key = re.sub(r"[^a-z0-9äöüß]+", " ", clean_cell_text(value).casefold())
    return LABEL_ALIASES.get(" ".join(key.split()), "")


def write_prompt_evaluation_sheet(workbook, frame: pd.DataFrame) -> None:
    """Add a compact E1 prompt-development evaluation to the result workbook."""
    if "Prompt_Evaluation" in workbook.sheetnames:
        del workbook["Prompt_Evaluation"]
    sheet = workbook.create_sheet("Prompt_Evaluation")
    sheet.sheet_view.showGridLines = False

    marker = frame.get("prompt_engineering", pd.Series("", index=frame.index))
    selected = frame.loc[
        marker.fillna("").astype(str).str.strip().str.casefold().eq("yes")
    ].copy()

    cases: list[dict[str, object]] = []
    for _, row in selected.iterrows():
        ground_truth = clean_cell_text(row.get("Ground Truth"))
        predicted = clean_cell_text(row.get("Predicted_Label"))
        canonical_ground_truth = canonical_class_label(ground_truth)
        canonical_prediction = canonical_class_label(predicted)

        if not predicted:
            evaluation_status = "NOT_EVALUABLE: no parsed prediction"
        elif not canonical_ground_truth:
            evaluation_status = "NOT_EVALUABLE: unknown Ground Truth label"
        elif not canonical_prediction:
            evaluation_status = "NOT_EVALUABLE: unknown predicted label"
        elif canonical_prediction == canonical_ground_truth:
            evaluation_status = "CORRECT"
        else:
            evaluation_status = "INCORRECT"

        cases.append(
            {
                "SAP-Nummer": clean_cell_text(row.get("SAP-Nummer")),
                "Teamcenter": clean_cell_text(row.get("Teamcenter")),
                "Ground Truth": ground_truth,
                "Predicted_Label": predicted,
                "Evaluation_Status": evaluation_status,
                "Processing_Status": clean_cell_text(row.get("Processing_Status")),
                "JSON_Parse_Status": clean_cell_text(row.get("JSON_Parse_Status")),
                "_ground_truth": canonical_ground_truth,
            }
        )

    case_frame = pd.DataFrame(cases)
    comparable = case_frame["Evaluation_Status"].isin(("CORRECT", "INCORRECT")) if not case_frame.empty else pd.Series(dtype=bool)
    correct_count = int((case_frame["Evaluation_Status"] == "CORRECT").sum()) if not case_frame.empty else 0
    comparable_count = int(comparable.sum()) if not case_frame.empty else 0
    successful_count = int(
        case_frame["Processing_Status"].eq("SUCCESS").sum()
    ) if not case_frame.empty else 0
    registers = sorted(
        {
            clean_cell_text(value)
            for value in selected.get("Register", pd.Series("", index=selected.index))
            if clean_cell_text(value)
        }
    )

    title_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    header_fill = PatternFill("solid", fgColor="2F75B5")
    correct_fill = PatternFill("solid", fgColor="C6EFCE")
    incorrect_fill = PatternFill("solid", fgColor="FFC7CE")
    neutral_fill = PatternFill("solid", fgColor="FFEB9C")

    sheet["A1"] = "Prompt Development Evaluation"
    sheet["A1"].fill = title_fill
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    sheet.merge_cells("A1:D1")

    summary_rows = [
        ("Prompt configuration", clean_cell_text(selected.get("Prompt_Config", pd.Series([""])).iloc[0]) if not selected.empty else ""),
        ("Model", clean_cell_text(selected.get("Run_Model", pd.Series([""])).iloc[0]) if not selected.empty else ""),
        ("Evaluation scope", "Rows marked prompt_engineering = yes"),
        ("Register(s)", ", ".join(registers) or "not specified"),
        ("Selected BGs", len(case_frame)),
        ("Parsed and comparable", comparable_count),
        ("Correct predictions", correct_count),
        ("Overall Accuracy", correct_count / comparable_count if comparable_count else None),
        ("Not evaluable", len(case_frame) - comparable_count),
        ("Processing_Status = SUCCESS", successful_count),
    ]
    sheet["A3"] = "Overall result"
    sheet["A3"].fill = section_fill
    sheet["A3"].font = Font(bold=True)
    for row_number, (metric, value) in enumerate(summary_rows, start=4):
        sheet.cell(row=row_number, column=1, value=metric)
        value_cell = sheet.cell(row=row_number, column=2, value=value)
        if metric == "Overall Accuracy" and value is not None:
            value_cell.number_format = "0.0%"
        if metric in {"Correct predictions", "Overall Accuracy"}:
            value_cell.fill = correct_fill

    class_header_row = 16
    sheet.cell(row=class_header_row, column=1, value="Class-wise result").fill = section_fill
    sheet.cell(row=class_header_row, column=1).font = Font(bold=True)
    class_headers = ("Class", "Ground Truth rows", "Correct predictions", "Accuracy")
    for column, header in enumerate(class_headers, start=1):
        cell = sheet.cell(row=class_header_row + 1, column=column, value=header)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_number, class_label in enumerate(CANONICAL_CLASS_ORDER, start=class_header_row + 2):
        class_cases = case_frame.loc[case_frame["_ground_truth"].eq(class_label)] if not case_frame.empty else case_frame
        class_correct = int(class_cases["Evaluation_Status"].eq("CORRECT").sum()) if not class_cases.empty else 0
        class_total = len(class_cases)
        sheet.cell(row=row_number, column=1, value=class_label)
        sheet.cell(row=row_number, column=2, value=class_total)
        sheet.cell(row=row_number, column=3, value=class_correct)
        accuracy_cell = sheet.cell(
            row=row_number,
            column=4,
            value=class_correct / class_total if class_total else None,
        )
        if class_total:
            accuracy_cell.number_format = "0.0%"

    details_header_row = class_header_row + 11
    sheet.cell(row=details_header_row, column=1, value="Case-by-case comparison").fill = section_fill
    sheet.cell(row=details_header_row, column=1).font = Font(bold=True)
    detail_headers = (
        "SAP-Nummer", "Teamcenter", "Ground Truth", "Predicted_Label",
        "Evaluation_Status", "Processing_Status", "JSON_Parse_Status",
    )
    for column, header in enumerate(detail_headers, start=1):
        cell = sheet.cell(row=details_header_row + 1, column=column, value=header)
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for output_row, (_, case) in enumerate(case_frame.iterrows(), start=details_header_row + 2):
        for column, header in enumerate(detail_headers, start=1):
            cell = sheet.cell(row=output_row, column=column, value=case[header])
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if header == "Evaluation_Status":
                if case[header] == "CORRECT":
                    cell.fill = correct_fill
                elif case[header] == "INCORRECT":
                    cell.fill = incorrect_fill
                else:
                    cell.fill = neutral_fill

    widths = {"A": 28, "B": 24, "C": 30, "D": 30, "E": 34, "F": 32, "G": 22}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = f"A{details_header_row + 2}"
    if len(case_frame):
        sheet.auto_filter.ref = f"A{details_header_row + 1}:G{details_header_row + 1 + len(case_frame)}"


def write_checkpoint(frame: pd.DataFrame, output_path: Path) -> None:
    # Keep human-readable experiment results on the left.  Long technical
    # paths stay at the far right, where they cannot cover result cells.
    preferred_order = [
        "prompt_engineering", "SAP-Nummer", "Teamcenter", "Benennung (E)",
        "Benennung (D)", "Ground Truth", "Register", "Weitere zulässige Ground Truth",
        "Predicted_Label", "Reasoning", "Confidence_Percent",
        "Possible_Classes_If_Ambiguous", "Processing_Status", "JSON_Parse_Status",
    ]
    ordered_columns = [column for column in preferred_order if column in frame.columns]
    ordered_columns.extend(column for column in frame.columns if column not in ordered_columns)
    frame.loc[:, ordered_columns].to_excel(output_path, index=False, engine="openpyxl")
    workbook = load_workbook(output_path)
    sheet = workbook.active
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    result_columns = {
        "Predicted_Label", "Reasoning", "Confidence_Percent",
        "Possible_Classes_If_Ambiguous", "JSON_Parse_Status", "Processing_Status",
    }
    technical_columns = {
        "Data_Folder_Path", "Files_Used", "Raw_Model_Response", "Token_Usage_JSON",
        "Prompt_File", "Run_Timestamp",
    }
    widths = {
        "prompt_engineering": 18, "SAP-Nummer": 16, "Teamcenter": 16,
        "Benennung (E)": 22, "Benennung (D)": 22, "Ground Truth": 26,
        "Register": 10, "Weitere zulässige Ground Truth": 30,
        "Data_Folder_Path": 12, "Predicted_Label": 24, "Reasoning": 55,
        "Confidence_Percent": 18, "Possible_Classes_If_Ambiguous": 32,
        "Raw_Model_Response": 55, "JSON_Parse_Status": 20, "Processing_Status": 30,
        "Files_Used": 55, "File_Count": 12, "Run_Model": 22, "Prompt_Config": 14,
        "Prompt_File": 32, "Temperature": 14, "Run_Timestamp": 22,
        "Token_Usage_JSON": 26,
    }
    header_fill = PatternFill("solid", fgColor="1F4E78")
    result_fill = PatternFill("solid", fgColor="2F75B5")
    technical_fill = PatternFill("solid", fgColor="7F8C8D")

    for column_index, cell in enumerate(sheet[1], start=1):
        header = str(cell.value)
        cell.fill = result_fill if header in result_columns else technical_fill if header in technical_columns else header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(column_index)].width = widths.get(header, 20)

    sheet.row_dimensions[1].height = 36
    wrap_headers = {"Reasoning", "Possible_Classes_If_Ambiguous", "Raw_Model_Response", "Files_Used", "Processing_Status"}
    header_positions = {str(cell.value): cell.column for cell in sheet[1]}
    for row_index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = 75
        for header in wrap_headers:
            if header in header_positions:
                sheet.cell(row=row_index, column=header_positions[header]).alignment = Alignment(
                    vertical="top", wrap_text=True
                )

    write_prompt_evaluation_sheet(workbook, frame)
    workbook.save(output_path)


def evaluate_existing_workbook(output_path: Path) -> Path:
    """Add the evaluation sheet to a completed P1/P2/P3 result file without rerunning it."""
    if not output_path.is_file():
        raise FileNotFoundError(f"Result workbook not found: {output_path}")
    frame = pd.read_excel(output_path, sheet_name=0, dtype=str)
    workbook = load_workbook(output_path)
    write_prompt_evaluation_sheet(workbook, frame)
    workbook.save(output_path)
    return output_path


def run(args: argparse.Namespace) -> Path:
    if not 0 <= args.temperature <= 2:
        raise ValueError("--temperature must be between 0 and 2.")
    if args.max_rows is not None and args.max_rows <= 0:
        raise ValueError("--max-rows must be greater than 0.")
    if args.max_output_tokens <= 0:
        raise ValueError("--max-output-tokens must be greater than 0.")
    api_key = os.getenv("BOSCH_FARM_SUBSCRIPTION_KEY")
    if not api_key:
        raise EnvironmentError("BOSCH_FARM_SUBSCRIPTION_KEY is not set in .env.")
    if not args.input_excel.is_file():
        raise FileNotFoundError(f"Dataset not found: {args.input_excel}")

    prompt_path, system_prompt = read_prompt(args)
    dataset = pd.read_excel(args.input_excel, sheet_name=args.sheet_name, dtype=str)
    selection_column = find_column(dataset, "prompt_engineering")
    path_column = find_column(dataset, "Data_Folder_Path")

    if not args.all_rows:
        dataset = dataset.loc[
            dataset[selection_column].fillna("").astype(str).str.strip().str.casefold().eq("yes")
        ].copy()
    if dataset.empty:
        scope = "all rows" if args.all_rows else "prompt_engineering = yes"
        raise ValueError(f"No rows selected for {scope}.")
    if args.max_rows is not None:
        dataset = dataset.head(args.max_rows).copy()

    logging.info("Dataset: %s", args.input_excel.resolve())
    logging.info("Prompt: %s", prompt_path.resolve())
    logging.info("Selected rows: %d", len(dataset))

    llm = create_connector(args.model, api_key)
    output_path = make_output_path(args)
    result = dataset.copy()
    for column in (
        "Predicted_Label", "Reasoning", "Confidence_Percent",
        "Possible_Classes_If_Ambiguous", "Raw_Model_Response", "JSON_Parse_Status",
        "Processing_Status", "Files_Used", "File_Count", "Run_Model", "Prompt_Config",
        "Prompt_File", "Temperature", "Run_Timestamp", "Token_Usage_JSON",
    ):
        result[column] = pd.NA

    run_timestamp = datetime.now().isoformat(timespec="seconds")
    generation_config = {
        "temperature": args.temperature,
        "topP": 0.95,
        "candidateCount": 1,
        "maxOutputTokens": args.max_output_tokens,
    }

    for index, row in tqdm(result.iterrows(), total=len(result), desc="Classifying"):
        try:
            files = collect_supported_files(row[path_column])
            if not files:
                result.loc[index, "Processing_Status"] = "SKIPPED: no supported PDF/image files"
                result.loc[index, "File_Count"] = 0
                write_checkpoint(result, output_path)
                continue

            display_files = list(files)
            response = llm.ask_about_files(
                file_paths=files,
                question=build_question(row),
                system_prompt=system_prompt,
                generation_config=generation_config,
            )
            if request_is_too_large(response):
                with tempfile.TemporaryDirectory(prefix="prompt_development_retry_") as temp_dir_name:
                    for pages in (3, 2):
                        sent_files, display_files, retry_note = prepare_reduced_largest_pdf(
                            llm, files, pages, Path(temp_dir_name)
                        )
                        logging.info("Row %s: %s", index + 2, retry_note)
                        response = llm.ask_about_files(
                            file_paths=sent_files,
                            question=build_question(row),
                            system_prompt=system_prompt,
                            generation_config=generation_config,
                        )
                        if not request_is_too_large(response):
                            break
            payload, parse_status = parse_json_response(response)
            result.loc[index, "Raw_Model_Response"] = str(response)
            result.loc[index, "JSON_Parse_Status"] = parse_status
            result.loc[index, "Predicted_Label"] = json_value(payload, "class_label", "class")
            result.loc[index, "Reasoning"] = json_value(payload, "reasoning", "begruendung", "begründung")
            result.loc[index, "Confidence_Percent"] = json_value(
                payload, "confidence_percent", "confidence", "konfidenz"
            )
            result.loc[index, "Possible_Classes_If_Ambiguous"] = normalise_possible_classes(
                json_value(payload, "possible_classes_if_ambiguous", "possible_classes", "zweitwahl")
            )
            result.loc[index, "Files_Used"] = "\n".join(display_files)
            result.loc[index, "File_Count"] = len(files)
            result.loc[index, "Run_Model"] = args.model
            result.loc[index, "Prompt_Config"] = args.prompt_config
            result.loc[index, "Prompt_File"] = str(prompt_path)
            result.loc[index, "Temperature"] = args.temperature
            result.loc[index, "Run_Timestamp"] = run_timestamp
            result.loc[index, "Token_Usage_JSON"] = json.dumps(
                llm.get_last_token_usage() or {}, ensure_ascii=False
            )
            if request_is_too_large(response):
                result.loc[index, "Processing_Status"] = "ERROR: HTTP 413 after retries with 3 and 2 PDF pages"
                result.loc[index, "JSON_Parse_Status"] = "REQUEST_TOO_LARGE"
            else:
                result.loc[index, "Processing_Status"] = "SUCCESS" if payload else "CHECK: invalid JSON response"
        except Exception as error:
            logging.exception("Failed row %s", index + 2)
            result.loc[index, "Processing_Status"] = f"ERROR: {error}"

        write_checkpoint(result, output_path)

    logging.info("Finished. Result saved to %s", output_path.resolve())
    return output_path


if __name__ == "__main__":
    load_dotenv(PROJECT_DIR / ".env")
    arguments = parse_args()
    log_dir = arguments.output_dir / "logs"
    log_dir.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
        handlers=[
            logging.FileHandler(log_dir / f"prompt_development_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.log", encoding="utf-8"),
            logging.StreamHandler(),
        ],
        force=True,
    )
    try:
        if arguments.evaluate_existing is not None:
            result_path = evaluate_existing_workbook(arguments.evaluate_existing)
            print(f"Evaluation sheet updated: {result_path.resolve()}")
        else:
            if not arguments.prompt_config:
                raise ValueError("--prompt-config is required unless --evaluate-existing is used.")
            result_path = run(arguments)
            print(f"Done: {result_path.resolve()}")
    except Exception as error:
        logging.error("Classification did not start: %s", error)
        sys.exit(1)


# PowerShell command: Run P1 for all rows marked prompt_engineering = yes.
# python .\run_prompt_development.py --prompt-config P1 --model gemini-2.5-pro
#
# PowerShell command: Add/refresh Prompt_Evaluation in an existing result file
# without calling the LLM again.
# python .\run_prompt_development.py --evaluate-existing ".\outputs\prompt_development_P1_gemini-2.5-pro_YYYY-MM-DD_HH-MM-SS.xlsx"
#
# PowerShell commands: Add evaluation sheets to the two completed valid results
# without calling Gemini again.
# python .\run_prompt_development.py --evaluate-existing ".\outputs\valid_results\prompt_development_P1_gemini-2.5-pro_2026-09-09_17-29-42.xlsx"
# python .\run_prompt_development.py --evaluate-existing ".\outputs\valid_results\prompt_development_P2_gemini-2.5-pro_2026-09-09_18-40-47.xlsx"

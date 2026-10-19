"""Create a safe rename preview for files inside collected BG folders.

The script reads ``folder_matches`` from the coverage workbook created by
``check_129bg_folder_coverage.py``. Its default mode only writes an Excel
preview; existing files remain untouched. Pass ``--apply`` only after the
preview has been checked.
"""

from __future__ import annotations

import argparse
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DOCUMENT_TYPES = {
    "ATD": "Technische_Zeichnung",
    "ATB": "Stückliste",
    "CAT": "Datenblatt",
    "ATO": "Angebot",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}


def clean_text(value: object) -> str:
    """Return an Excel cell as clean text without turning NaN into text."""
    if pd.isna(value):
        return ""
    text = str(value).strip()
    return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def folder_code(row: pd.Series) -> str:
    """Use the BG folder name, falling back to the identifiers when needed."""
    for column in ("folder_name", "SAP-Nummer", "Teamcenter"):
        value = clean_text(row.get(column, ""))
        if value:
            return value
    return ""


def source_code_before_token(stem: str, token: str) -> str:
    """Return the alphanumeric code occurring before ATD/ATB/CAT/ATO."""
    match = re.search(re.escape(token), stem, flags=re.IGNORECASE)
    if not match:
        return ""
    before = stem[: match.start()].rstrip(" _-.")
    code_match = re.match(r"[A-Za-z0-9]+", before)
    return code_match.group(0) if code_match else ""


def rename_for_document(path: Path) -> tuple[str, str, str]:
    """Return (rule, proposed filename, note) for a root-level document."""
    if path.suffix.casefold() != ".pdf":
        return "", "", ""

    for token, label in DOCUMENT_TYPES.items():
        if token.casefold() in path.stem.casefold():
            code = source_code_before_token(path.stem, token)
            if not code:
                return label, "", f"needs review: no source code before {token}"
            return label, f"{label}_{code}{path.suffix.casefold()}", ""
    return "", "", ""


def collect_folder_rows(matches: pd.DataFrame) -> list[dict[str, str]]:
    if "folder_path" not in matches.columns:
        raise ValueError(
            "The coverage workbook sheet 'folder_matches' has no 'folder_path' column."
        )

    folders: dict[str, dict[str, str]] = {}
    for _, row in matches.iterrows():
        raw_path = clean_text(row.get("folder_path", ""))
        if not raw_path:
            continue
        path = Path(raw_path)
        folders.setdefault(
            str(path).casefold(),
            {
                "folder_path": str(path),
                "folder_code": folder_code(row) or path.name,
                "sap_number": clean_text(row.get("SAP-Nummer", "")),
                "teamcenter": clean_text(row.get("Teamcenter", "")),
            },
        )
    return list(folders.values())


def preview_one_folder(folder: dict[str, str]) -> list[dict[str, object]]:
    bg_path = Path(folder["folder_path"])
    rows: list[dict[str, object]] = []
    if not bg_path.is_dir():
        return [{
            **folder,
            "original_path": str(bg_path),
            "original_filename": "",
            "proposed_filename": "",
            "rule": "",
            "status": "folder_not_accessible",
            "note": "Folder does not exist or is not accessible.",
            "modified_time": "",
        }]

    # PDF and DFC rules apply only to files directly in the BG folder.
    for path in sorted(
        (item for item in bg_path.iterdir() if item.is_file()),
        key=lambda item: item.name.casefold(),
    ):
        modified = datetime.fromtimestamp(path.stat().st_mtime).strftime(
            "%Y-%m-%d %H:%M:%S"
        )
        if path.name.casefold() == "dfc.png":
            proposed = f"{folder['folder_code']}_DFC_Struktur.png"
            rows.append({
                **folder,
                "original_path": str(path),
                "original_filename": path.name,
                "proposed_filename": proposed,
                "rule": "DFC.png",
                "status": "ready" if path.name != proposed else "already_correct",
                "note": "",
                "modified_time": modified,
            })
            continue

        rule, proposed, note = rename_for_document(path)
        if rule:
            rows.append({
                **folder,
                "original_path": str(path),
                "original_filename": path.name,
                "proposed_filename": proposed,
                "rule": rule,
                "status": (
                    "ready" if proposed and path.name != proposed
                    else "already_correct" if proposed
                    else "needs_review"
                ),
                "note": note,
                "modified_time": modified,
            })

    # Screenshot files are only numbered inside the immediate Bilder subfolder.
    bilder = bg_path / "Bilder"
    if bilder.is_dir():
        screenshots = [
            path for path in bilder.iterdir()
            if path.is_file()
            and path.suffix.casefold() in IMAGE_EXTENSIONS
            and path.stem.casefold().startswith("screenshot")
        ]
        screenshots.sort(key=lambda path: (path.stat().st_mtime, path.name.casefold()))
        for index, path in enumerate(screenshots, start=1):
            modified = datetime.fromtimestamp(path.stat().st_mtime).strftime(
                "%Y-%m-%d %H:%M:%S"
            )
            proposed = f"{folder['folder_code']}_Screenshot_{index}{path.suffix.casefold()}"
            rows.append({
                **folder,
                "original_path": str(path),
                "original_filename": path.name,
                "proposed_filename": proposed,
                "rule": "Bilder/Screenshot by modified time",
                "status": "ready" if path.name != proposed else "already_correct",
                "note": "",
                "modified_time": modified,
            })
    return rows


def mark_conflicts(preview: pd.DataFrame) -> pd.DataFrame:
    """Never allow a proposed rename to overwrite a file."""
    if preview.empty:
        return preview

    preview["target_path"] = preview.apply(
        lambda row: (
            str(Path(str(row["original_path"])).with_name(str(row["proposed_filename"])))
            if row["proposed_filename"] else ""
        ),
        axis=1,
    )
    ready = preview[preview["status"] == "ready"]
    duplicate_targets = set(
        ready.loc[ready["target_path"].duplicated(keep=False), "target_path"]
    )

    for index, row in preview.iterrows():
        if row["status"] != "ready":
            continue
        target = Path(row["target_path"])
        source = Path(row["original_path"])
        if row["target_path"] in duplicate_targets:
            preview.at[index, "status"] = "conflict"
            preview.at[index, "note"] = "Several files would receive this same name."
        elif target.exists() and target.resolve() != source.resolve():
            preview.at[index, "status"] = "conflict"
            preview.at[index, "note"] = "Target filename already exists."
    return preview


def style_sheet(worksheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    for column in range(1, worksheet.max_column + 1):
        header = str(worksheet.cell(row=1, column=column).value or "")
        width = 18
        if header in {"folder_path", "original_path", "target_path"}:
            width = 64
        elif header == "note":
            width = 45
        elif header in {"original_filename", "proposed_filename"}:
            width = 38
        worksheet.column_dimensions[get_column_letter(column)].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_preview(preview: pd.DataFrame, output_path: Path) -> None:
    summary = (
        preview.groupby(["rule", "status"], dropna=False).size().reset_index(name="count")
        if not preview.empty
        else pd.DataFrame(columns=["rule", "status", "count"])
    )
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        preview.to_excel(writer, sheet_name="rename_preview", index=False)
        summary.to_excel(writer, sheet_name="summary", index=False)
        for sheet in writer.sheets.values():
            style_sheet(sheet)


def apply_renames(preview: pd.DataFrame) -> tuple[int, list[dict[str, str]]]:
    completed = 0
    log: list[dict[str, str]] = []
    for _, row in preview[preview["status"] == "ready"].iterrows():
        source = Path(row["original_path"])
        target = Path(row["target_path"])
        try:
            if not source.is_file():
                raise FileNotFoundError("Source file is no longer available.")
            if target.exists():
                raise FileExistsError("Target file now exists; not overwritten.")
            source.rename(target)
            completed += 1
            result, note = "renamed", ""
        except OSError as error:
            result, note = "failed", str(error)
        log.append({
            "original_path": str(source),
            "target_path": str(target),
            "result": result,
            "note": note,
        })
    return completed, log


def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Preview or apply the agreed internal BG-file rename rules."
    )
    parser.add_argument(
        "--coverage-excel",
        required=True,
        type=Path,
        help="Coverage workbook produced by check_129bg_folder_coverage.py.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help="Folder for preview/log Excel files (default: next to coverage workbook).",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Actually rename only ready rows; conflicts are never overwritten.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_arguments()
    coverage_path = args.coverage_excel.expanduser()
    if not coverage_path.is_file():
        raise FileNotFoundError(f"Coverage workbook not found: {coverage_path}")

    matches = pd.read_excel(coverage_path, sheet_name="folder_matches", dtype=str)
    folders = collect_folder_rows(matches)
    preview_rows = [row for folder in folders for row in preview_one_folder(folder)]
    preview = mark_conflicts(pd.DataFrame(preview_rows))
    if preview.empty:
        preview = pd.DataFrame(columns=[
            "folder_path", "folder_code", "sap_number", "teamcenter",
            "original_path", "original_filename", "proposed_filename",
            "rule", "status", "note", "modified_time", "target_path",
        ])

    output_dir = (args.output_dir or coverage_path.parent).expanduser()
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    preview_path = output_dir / f"70BG_internal_file_rename_preview_{timestamp}.xlsx"
    write_preview(preview, preview_path)

    print(f"BG folders read from coverage workbook: {len(folders)}")
    print(f"Rename candidates: {len(preview)}")
    print(f"Preview written: {preview_path}")
    print(preview["status"].value_counts(dropna=False).to_string())

    if not args.apply:
        print("Preview mode only: no existing file has been renamed.")
        print("After checking the Excel preview, rerun with --apply.")
        return

    completed, log = apply_renames(preview)
    log_path = output_dir / f"70BG_internal_file_rename_log_{timestamp}.xlsx"
    pd.DataFrame(log).to_excel(log_path, index=False)
    print(f"Files renamed: {completed}")
    print(f"Rename log written: {log_path}")


if __name__ == "__main__":
    main()


# PowerShell example: preview only; no existing file will be changed.
# python .\scripts\rename_bg_internal_files.py `
#   --coverage-excel "$env:USERPROFILE\Downloads\70BG\129BG_folder_coverage_2026-10-19_13-20-20.xlsx" `
#   --output-dir "$env:USERPROFILE\Downloads\70BG"
#
# After checking the preview Excel, run the same command with --apply appended.

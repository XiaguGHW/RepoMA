#!/usr/bin/env python3
"""Prepare a text-only LLM experiment with a reviewable BG file inventory.

The script does *not* classify the Baugruppe's function.  It classifies each
readable document into a document role so later scripts can extract facts from
the appropriate files.  Excel files are converted locally to a text preview;
large PDFs and images are represented by a deterministic preview attachment.

The Bosch ``llm_connector_with_prompt_caching.py`` is deliberately reused unchanged.  It must
provide ``LLMConnector(...).ask_about_files(...)`` as used by
``scripts/run_classification.py``.
"""

from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
import hashlib
import json
import logging
import os
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import pandas as pd

try:
    import fitz
except ImportError:  # pragma: no cover - clear runtime message below
    fitz = None

try:
    from PIL import Image
except ImportError:  # pragma: no cover - clear runtime message below
    Image = None

try:
    from dotenv import load_dotenv
except ImportError:
    def load_dotenv(*_args: Any, **_kwargs: Any) -> bool:
        return False

try:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pandas can still create the workbook without formatting
    Alignment = Font = PatternFill = get_column_letter = None


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent if SCRIPT_DIR.name.casefold() == "scripts" else SCRIPT_DIR

READABLE_EXTENSIONS = {".pdf", ".png", ".jpg", ".jpeg", ".xlsx", ".xls", ".csv"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
EXCEL_EXTENSIONS = {".xlsx", ".xls", ".csv"}
EXCLUDED_IMAGE_DIRECTORY_NAMES = {"screenshots", "bilder"}

DOCUMENT_TYPES = {
    "assembly_drawing",
    "bom",
    "component_datasheet",
    "assembly_structure_or_dfc",
    "cad_screenshot_or_visual_context",
    "other_technical_document",
    "irrelevant_or_unclear",
}

OUTPUT_COLUMNS = [
    "Dataset_Row",
    "BG_Folder",
    "Data_Folder_Path",
    "Relative_Path",
    "File_Name",
    "Extension",
    "File_Size_MB",
    "Readable_By_Pipeline",
    "Preview_Mode",
    "Primary_Type",
    "Secondary_Types",
    "Confidence",
    "Evidence",
    "Needs_Review",
    "Processing_Status",
    "Manual_Type",
    "Manual_Secondary_Types",
    "Review_Status",
    "Raw_Model_Response",
    "Run_Model",
    "Run_Timestamp",
]

SYSTEM_PROMPT = """Du klassifizierst technische Dokumente aus dem Maschinen- und Anlagenbau.
Deine Aufgabe ist ausschließlich die Dokumentrolle der bereitgestellten einzelnen Datei.
Du darfst NICHT die Funktionsklasse der Baugruppe ableiten und darfst keine Labels aus
einem Codebook verwenden.

Erlaubte primary_type-Werte:
- assembly_drawing: technische Baugruppen-/Montagezeichnung mit Zeichnungskopf,
  Ansichten, Bemaßungen oder Positionsreferenzen.
- bom: Stückliste, strukturierte Teileliste oder Tabellen mit Position, Menge,
  Benennung und/oder Teilenummern.
- component_datasheet: Herstellerdatenblatt eines einzelnen Funktions- oder
  Kaufteils mit Modell und technischen Parametern.
- assembly_structure_or_dfc: DFC-, Strukturbaum-, Baugruppenstruktur-,
  Hierarchie- oder Verbindungsdarstellung.
- cad_screenshot_or_visual_context: reine 3D-CAD-Ansicht, Screenshot oder
  visuelle Baugruppenübersicht ohne Strukturbaum.
- other_technical_document: sonstiges technisch relevantes Dokument.
- irrelevant_or_unclear: nicht belastbar klassifizierbar oder technisch irrelevant.

secondary_types darf leer sein oder nur weitere Werte aus derselben Liste enthalten.
Nutze nur sichtbare bzw. im Text enthaltene Hinweise. Antworte ausschließlich als JSON."""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--dataset-excel", type=Path, required=True,
        help="Existing Task3 dataset Excel containing Data_Folder_Path for every BG.",
    )
    parser.add_argument("--output-dir", type=Path, default=PROJECT_DIR / "outputs" / "text_only_preprocessing")
    parser.add_argument("--model", default="gemini-2.5-pro")
    parser.add_argument("--max-bgs", type=int, default=None, help="Process only the first N BG folders.")
    parser.add_argument("--max-workers", type=int, default=1, help="Concurrent file workers; default 1.")
    parser.add_argument(
        "--include-prompt-engineering",
        action="store_true",
        help=(
            "Include rows marked prompt_engineering = yes. By default these development "
            "BGs are excluded so the inventory is suitable for the independent final test."
        ),
    )
    parser.add_argument("--max-file-mb", type=float, default=5.0, help="Largest original PDF/image attachment in MiB.")
    parser.add_argument("--max-preview-px", type=int, default=1600, help="Maximum long edge for generated image previews.")
    parser.add_argument("--excel-preview-rows", type=int, default=20)
    parser.add_argument("--excel-preview-sheets", type=int, default=3)
    parser.add_argument("--pdf-preview-pages", type=int, default=2)
    parser.add_argument("--max-preview-chars", type=int, default=5000)
    parser.add_argument("--retries", type=int, default=1)
    parser.add_argument(
        "--no-filename-rules",
        action="store_false",
        dest="use_filename_rules",
        help="Disable high-confidence local filename/path rules and send every readable file to the LLM.",
    )
    parser.set_defaults(use_filename_rules=True)
    return parser.parse_args()


def safe_name(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", value)


def sha256_short(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()[:16]


def clean_text(value: str, max_chars: int) -> str:
    value = re.sub(r"\s+", " ", value).strip()
    return value[:max_chars]


def pdf_text_preview(path: Path, pages: int, max_chars: int) -> tuple[str, int, str]:
    if fitz is None:
        return "", 0, "PyMuPDF unavailable"
    try:
        document = fitz.open(path)
        page_total = len(document)
        snippets = []
        for index in range(min(page_total, pages)):
            text = clean_text(document[index].get_text("text"), max_chars)
            if text:
                snippets.append(f"[Page {index + 1}] {text}")
        document.close()
        return "\n".join(snippets)[:max_chars], page_total, "ok"
    except Exception as error:
        return "", 0, f"PDF text extraction failed: {error}"


def excel_text_preview(path: Path, max_sheets: int, max_rows: int, max_chars: int) -> tuple[str, str]:
    try:
        if path.suffix.casefold() == ".csv":
            frame = pd.read_csv(path, nrows=max_rows, dtype=str, encoding_errors="replace")
            return clean_text("[Sheet: CSV]\n" + frame.to_csv(index=False), max_chars), "ok"

        workbook = pd.ExcelFile(path)
        snippets = []
        for sheet_name in workbook.sheet_names[:max_sheets]:
            frame = pd.read_excel(path, sheet_name=sheet_name, nrows=max_rows, dtype=str)
            snippets.append(f"[Sheet: {sheet_name}]\n{frame.fillna('').to_csv(index=False)}")
        return clean_text("\n".join(snippets), max_chars), "ok"
    except Exception as error:
        return "", f"Excel preview failed: {error}"


def render_pdf_first_page(path: Path, target: Path, max_px: int) -> tuple[Path | None, str]:
    if fitz is None:
        return None, "PyMuPDF unavailable"
    try:
        document = fitz.open(path)
        if not document:
            return None, "PDF has no pages"
        page = document[0]
        scale = min(max_px / max(page.rect.width, page.rect.height), 2.0)
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        target.parent.mkdir(parents=True, exist_ok=True)
        pix.save(target)
        document.close()
        return target, "pdf_first_page_preview"
    except Exception as error:
        return None, f"PDF preview rendering failed: {error}"


def image_preview(path: Path, target: Path, max_px: int) -> tuple[Path | None, str]:
    if Image is None:
        return None, "Pillow unavailable"
    try:
        with Image.open(path) as image:
            image = image.convert("RGB")
            image.thumbnail((max_px, max_px))
            target.parent.mkdir(parents=True, exist_ok=True)
            image.save(target, format="JPEG", quality=88, optimize=True)
        return target, "resized_image_preview"
    except Exception as error:
        return None, f"Image preview failed: {error}"


def prepare_file_for_llm(path: Path, cache_dir: Path, args: argparse.Namespace) -> tuple[list[str], str, str, bool]:
    """Return LLM context and whether it contains actual readable file content."""
    extension = path.suffix.casefold()
    size_mib = path.stat().st_size / (1024 * 1024)
    cache_base = cache_dir / f"{sha256_short(path)}_{safe_name(path.stem)}"

    if extension in EXCEL_EXTENSIONS:
        text, status = excel_text_preview(path, args.excel_preview_sheets, args.excel_preview_rows, args.max_preview_chars)
        return [], text, f"excel_text_preview ({status})", bool(text.strip())

    if extension == ".pdf":
        text, page_total, status = pdf_text_preview(path, args.pdf_preview_pages, args.max_preview_chars)
        metadata = f"PDF pages: {page_total}. Native text preview status: {status}.\n{text}"
        if size_mib <= args.max_file_mb:
            return [str(path)], metadata, "original_pdf_attachment + native_text_preview", True
        preview, preview_status = render_pdf_first_page(path, cache_base.with_suffix(".jpg"), args.max_preview_px)
        attachments = [str(preview)] if preview else []
        return attachments, metadata, f"large_pdf_first_page_preview ({preview_status}) + native_text_preview", bool(text.strip() or attachments)

    if extension in IMAGE_EXTENSIONS:
        if size_mib <= args.max_file_mb:
            return [str(path)], "", "original_image_attachment", True
        preview, preview_status = image_preview(path, cache_base.with_suffix(".jpg"), args.max_preview_px)
        attachments = [str(preview)] if preview else []
        return attachments, "", f"large_image_preview ({preview_status})", bool(attachments)

    return [], "", "unsupported", False


def build_question(relative_path: str, path: Path, preview_text: str) -> str:
    preview = preview_text or "No locally extractable text preview is available."
    return f"""Klassifiziere genau diese eine Datei nach ihrer Dokumentrolle.

Relative path: {relative_path}
File name: {path.name}
Extension: {path.suffix.lower()}
Local text preview:
{preview}

Return exactly this JSON object:
{{
  "primary_type": "one allowed value",
  "secondary_types": ["zero or more allowed values"],
  "confidence": 0.0,
  "evidence": "short evidence based on file content or visible document characteristics",
  "needs_review": false
}}
"""


def extract_json(raw: str) -> dict[str, Any] | None:
    cleaned = str(raw).strip()
    fenced = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", cleaned, flags=re.DOTALL | re.IGNORECASE)
    candidates = [fenced.group(1)] if fenced else []
    start, end = cleaned.find("{"), cleaned.rfind("}")
    if start >= 0 and end > start:
        candidates.append(cleaned[start:end + 1])
    for candidate in candidates:
        try:
            payload = json.loads(candidate)
            return payload if isinstance(payload, dict) else None
        except json.JSONDecodeError:
            continue
    return None


def normalize_result(payload: dict[str, Any] | None) -> tuple[str, str, float | None, str, str]:
    if not payload:
        return "", "", None, "", "yes"
    primary = str(payload.get("primary_type", "")).strip().casefold()
    if primary not in DOCUMENT_TYPES:
        primary = ""
    secondary_raw = payload.get("secondary_types", [])
    if not isinstance(secondary_raw, list):
        secondary_raw = []
    secondary = [str(item).strip().casefold() for item in secondary_raw]
    secondary = [item for item in secondary if item in DOCUMENT_TYPES and item != primary]
    try:
        confidence = float(payload.get("confidence"))
        confidence = max(0.0, min(1.0, confidence))
    except (TypeError, ValueError):
        confidence = None
    evidence = str(payload.get("evidence", "")).strip()
    review = "yes" if bool(payload.get("needs_review", False)) or not primary else "no"
    return primary, "; ".join(dict.fromkeys(secondary)), confidence, evidence, review


def local_filename_rule(bg_folder: Path, file_path: Path) -> tuple[str, str] | None:
    """Return only high-confidence document-role rules; leave all other files to the LLM."""
    relative = str(file_path.relative_to(bg_folder)).replace("\\", "/").casefold()
    stem = file_path.stem.casefold()

    rules = [
        ("bom", ("stückliste", "stueckliste", "bom", "partslist", "parts_list", "teileliste")),
        ("component_datasheet", ("datenblatt", "data_sheet", "datasheet", "katalog", "catalog")),
        ("assembly_structure_or_dfc", ("strukturbaum", "baugruppenstruktur", "dfc_structure", "dfc-struktur", "dfc_struktur")),
        ("assembly_drawing", ("zeichnung", "drawing", "montageplan")),
        ("cad_screenshot_or_visual_context", ("screenshot", "screen_shot", "cad_view", "cad-view", "3d_view", "3d-view")),
    ]
    for document_type, keywords in rules:
        for keyword in keywords:
            if keyword in relative:
                return document_type, f"Local filename/path rule matched '{keyword}' in '{relative}'."
    if "dfc" in stem or "/dfc/" in f"/{relative}/":
        return "assembly_structure_or_dfc", f"Local filename/path rule matched 'dfc' in '{relative}'."
    # Engineering-drawing PDFs are often named only with a part/drawing number.
    # When an exported PDF shares its stem with a native drawing file, this is a
    # strong, deterministic signal and avoids an unnecessary LLM call.
    if file_path.suffix.casefold() == ".pdf":
        drawing_extensions = {".idw", ".dwg", ".dxf"}
        try:
            has_matching_drawing_source = any(
                candidate.is_file()
                and candidate.suffix.casefold() in drawing_extensions
                and candidate.stem.casefold() == stem
                for candidate in file_path.parent.iterdir()
            )
        except OSError:
            has_matching_drawing_source = False
        if has_matching_drawing_source:
            return (
                "assembly_drawing",
                f"Local paired drawing-source rule matched same-stem CAD drawing for '{file_path.name}'.",
            )
    if re.match(r"^dr\d{5,}", stem):
        return "assembly_drawing", f"Local drawing-number rule matched '{file_path.name}'."
    return None


def create_connector(model_name: str, api_key: str):
    try:
        from llm_connector_with_prompt_caching import LLMConnector
    except ImportError as error:
        raise ImportError(
            "llm_connector_with_prompt_caching.py was not found beside the script. Place the Bosch connector "
            "next to this script."
        ) from error
    # Compatibility for the current cached connector: its file-upload helper
    # calls _get_mime_type(), but that helper is absent in the supplied file.
    if not hasattr(LLMConnector, "_get_mime_type"):
        def _get_mime_type(file_path: str) -> str:
            mime_types = {
                ".pdf": "application/pdf",
                ".png": "image/png",
                ".jpg": "image/jpeg",
                ".jpeg": "image/jpeg",
            }
            return mime_types.get(Path(file_path).suffix.casefold(), "application/octet-stream")
        LLMConnector._get_mime_type = staticmethod(_get_mime_type)
        logging.warning("Applied MIME-type compatibility patch for llm_connector_with_prompt_caching.py")
    return LLMConnector(model_name, api_key)


def classify_one_file(llm: Any, bg_folder: Path, file_path: Path, cache_dir: Path, args: argparse.Namespace) -> dict[str, Any]:
    relative_path = str(file_path.relative_to(bg_folder))
    extension = file_path.suffix.casefold()
    base = {
        "Dataset_Row": pd.NA,
        "BG_Folder": bg_folder.name,
        "Data_Folder_Path": str(bg_folder),
        "Relative_Path": relative_path,
        "File_Name": file_path.name,
        "Extension": extension,
        "File_Size_MB": round(file_path.stat().st_size / (1024 * 1024), 3),
        "Readable_By_Pipeline": "yes" if extension in READABLE_EXTENSIONS else "no",
        "Preview_Mode": "",
        "Primary_Type": "",
        "Secondary_Types": "",
        "Confidence": None,
        "Evidence": "",
        "Needs_Review": "yes",
        "Processing_Status": "",
        "Manual_Type": "",
        "Manual_Secondary_Types": "",
        "Review_Status": "not_reviewed",
        "Raw_Model_Response": "",
        "Run_Model": args.model,
        "Run_Timestamp": datetime.now().isoformat(timespec="seconds"),
    }
    if extension not in READABLE_EXTENSIONS:
        base["Preview_Mode"] = "unsupported_local_format"
        base["Processing_Status"] = "SKIPPED_UNSUPPORTED_FORMAT"
        return base

    if args.use_filename_rules:
        rule = local_filename_rule(bg_folder, file_path)
        if rule:
            document_type, evidence = rule
            base.update({
                "Preview_Mode": "local_filename_rule",
                "Primary_Type": document_type,
                "Confidence": 0.99,
                "Evidence": evidence,
                "Needs_Review": "no",
                "Processing_Status": "SUCCESS_LOCAL_FILENAME_RULE",
                "Run_Model": "local_filename_rule",
            })
            return base

    attachments, preview_text, preview_mode, has_usable_content = prepare_file_for_llm(file_path, cache_dir, args)
    base["Preview_Mode"] = preview_mode
    if not has_usable_content:
        base.update({
            "Processing_Status": "SKIPPED_NO_USABLE_CONTENT",
            "Evidence": "No extractable text and no usable file preview were available; no LLM request was made.",
            "Needs_Review": "yes",
        })
        return base
    question = build_question(relative_path, file_path, preview_text)
    response = ""
    for attempt in range(args.retries + 1):
        response = llm.ask_about_files(
            file_paths=attachments,
            question=question,
            system_prompt=SYSTEM_PROMPT,
            generation_config={"temperature": 0.0, "topP": 0.95, "candidateCount": 1, "maxOutputTokens": 600},
        )
        payload = extract_json(str(response))
        primary, secondary, confidence, evidence, needs_review = normalize_result(payload)
        if primary:
            base.update({
                "Primary_Type": primary,
                "Secondary_Types": secondary,
                "Confidence": confidence,
                "Evidence": evidence,
                "Needs_Review": needs_review,
                "Processing_Status": "SUCCESS" if attempt == 0 else "SUCCESS_AFTER_RETRY",
                "Raw_Model_Response": str(response),
            })
            return base
    base.update({
        "Processing_Status": "CHECK_INVALID_JSON_OR_TYPE",
        "Raw_Model_Response": str(response),
    })
    return base


def save_workbook(rows: list[dict[str, Any]], output_path: Path) -> None:
    frame = pd.DataFrame(rows, columns=OUTPUT_COLUMNS)
    summary = pd.crosstab(frame["BG_Folder"], frame["Primary_Type"]).reset_index() if not frame.empty else pd.DataFrame()
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="file_classification", index=False)
        summary.to_excel(writer, sheet_name="bg_summary", index=False)
        if PatternFill is None:
            return
        for sheet_name in ("file_classification", "bg_summary"):
            sheet = writer.sheets[sheet_name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for column_number, column in enumerate(sheet.columns, start=1):
                header = sheet.cell(1, column_number).value or ""
                width = min(max(14, len(str(header)) + 2), 35)
                if header in {"Relative_Path", "Evidence", "Raw_Model_Response"}:
                    width = 55
                sheet.column_dimensions[get_column_letter(column_number)].width = width


def iter_source_files(bg_folder: Path) -> Iterable[Path]:
    """Yield original source files, excluding known duplicate conversions.

    ``converted...`` directories and the image-only ``Screenshots``/``Bilder``
    directories do not belong in the text-only source-document inventory.
    ``*_visual_context.pdf`` is also a derived PDF and is excluded.
    """
    for root, directory_names, file_names in os.walk(bg_folder):
        # Pruning prevents os.walk from ever descending into these directories.
        directory_names[:] = sorted(
            (
                name for name in directory_names
                if not name.casefold().startswith("converted")
                and name.casefold() not in EXCLUDED_IMAGE_DIRECTORY_NAMES
            ),
            key=str.casefold,
        )
        root_path = Path(root)
        for filename in sorted(file_names, key=str.casefold):
            file_path = root_path / filename
            if filename.startswith("~$"):
                continue
            if (
                file_path.suffix.casefold() == ".pdf"
                and "_visual_context" in file_path.stem.casefold()
            ):
                logging.info("Skipping duplicate visual-context PDF: %s", file_path)
                continue
            yield file_path


def run(args: argparse.Namespace) -> Path:
    if not os.getenv("BOSCH_FARM_SUBSCRIPTION_KEY"):
        raise EnvironmentError("BOSCH_FARM_SUBSCRIPTION_KEY is not set in .env or the environment.")
    if args.max_bgs is not None and args.max_bgs <= 0:
        raise ValueError("--max-bgs must be positive.")
    if args.max_workers <= 0:
        raise ValueError("--max-workers must be positive.")

    dataset = pd.read_excel(args.dataset_excel, dtype=str).fillna("")
    if "Data_Folder_Path" not in dataset.columns:
        raise ValueError("Dataset Excel must contain the column 'Data_Folder_Path'.")
    prompt_engineering_column = next(
        (
            str(column)
            for column in dataset.columns
            if str(column).strip().casefold() == "prompt_engineering"
        ),
        None,
    )
    if not args.include_prompt_engineering:
        if prompt_engineering_column is None:
            raise ValueError(
                "Dataset Excel must contain the column 'prompt_engineering' for the final-test "
                "exclusion. Use --include-prompt-engineering only when intentionally processing "
                "a dataset without this separation."
            )
        prompt_mask = (
            dataset[prompt_engineering_column]
            .fillna("")
            .astype(str)
            .str.strip()
            .str.casefold()
            .eq("yes")
        )
        excluded_count = int(prompt_mask.sum())
        dataset = dataset.loc[~prompt_mask].copy()
        logging.info(
            "Excluded %d prompt-engineering BG row(s); final-test inventory contains %d row(s).",
            excluded_count,
            len(dataset),
        )
        if excluded_count != 14:
            logging.warning(
                "Expected 14 prompt-engineering rows, but excluded %d. Check the input Excel marker column.",
                excluded_count,
            )
    bg_entries = []
    for index, row in dataset.iterrows():
        raw_path = str(row["Data_Folder_Path"]).strip()
        if not raw_path:
            logging.warning("Dataset row %s has an empty Data_Folder_Path and will be skipped.", index)
            continue
        bg_entries.append((index, Path(raw_path).expanduser()))
    if args.max_bgs:
        bg_entries = bg_entries[:args.max_bgs]
    if not bg_entries:
        raise ValueError("No non-empty Data_Folder_Path entries found in the dataset Excel.")

    args.output_dir.mkdir(parents=True, exist_ok=True)
    cache_dir = args.output_dir / "file_preview_cache"
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    output_path = args.output_dir / f"file_inventory_classified_{safe_name(args.model)}_{timestamp}.xlsx"
    llm = create_connector(args.model, os.environ["BOSCH_FARM_SUBSCRIPTION_KEY"])
    rows: list[dict[str, Any]] = []

    for dataset_index, bg_folder in bg_entries:
        if not bg_folder.is_dir():
            logging.warning("Dataset row %s points to a missing folder: %s", dataset_index, bg_folder)
            rows.append({
                "Dataset_Row": dataset_index, "BG_Folder": bg_folder.name, "Data_Folder_Path": str(bg_folder),
                "Relative_Path": "", "File_Name": "", "Extension": "", "File_Size_MB": None,
                "Readable_By_Pipeline": "no", "Preview_Mode": "", "Primary_Type": "", "Secondary_Types": "",
                "Confidence": None, "Evidence": "", "Needs_Review": "yes", "Processing_Status": "MISSING_DATA_FOLDER",
                "Manual_Type": "", "Manual_Secondary_Types": "", "Review_Status": "not_reviewed",
                "Raw_Model_Response": "", "Run_Model": args.model,
                "Run_Timestamp": datetime.now().isoformat(timespec="seconds"),
            })
            save_workbook(rows, output_path)
            continue
        files = list(iter_source_files(bg_folder))
        logging.info("BG %s: %d file(s)", bg_folder.name, len(files))
        with ThreadPoolExecutor(max_workers=args.max_workers) as executor:
            future_to_file = {
                executor.submit(classify_one_file, llm, bg_folder, file_path, cache_dir, args): file_path
                for file_path in files
            }
            for future in as_completed(future_to_file):
                file_path = future_to_file[future]
                try:
                    record = future.result()
                    record["Dataset_Row"] = dataset_index
                    rows.append(record)
                except Exception as error:
                    logging.exception("Failed to classify %s", file_path)
                    rows.append({
                        "Dataset_Row": dataset_index, "BG_Folder": bg_folder.name, "Data_Folder_Path": str(bg_folder), "Relative_Path": str(file_path.relative_to(bg_folder)),
                        "File_Name": file_path.name, "Extension": file_path.suffix.casefold(),
                        "File_Size_MB": round(file_path.stat().st_size / (1024 * 1024), 3),
                        "Readable_By_Pipeline": "yes" if file_path.suffix.casefold() in READABLE_EXTENSIONS else "no",
                        "Preview_Mode": "", "Primary_Type": "", "Secondary_Types": "", "Confidence": None,
                        "Evidence": "", "Needs_Review": "yes", "Processing_Status": f"ERROR: {error}",
                        "Manual_Type": "", "Manual_Secondary_Types": "", "Review_Status": "not_reviewed",
                        "Raw_Model_Response": "", "Run_Model": args.model,
                        "Run_Timestamp": datetime.now().isoformat(timespec="seconds"),
                    })
                save_workbook(rows, output_path)  # checkpoint after every completed file
        """
        Legacy sequential loop retained below only as source context.
        """
        for file_path in []:
            try:
                record = classify_one_file(llm, bg_folder, file_path, cache_dir, args)
                record["Dataset_Row"] = dataset_index
                rows.append(record)
            except Exception as error:
                logging.exception("Failed to classify %s", file_path)
                rows.append({
                    "Dataset_Row": dataset_index, "BG_Folder": bg_folder.name, "Data_Folder_Path": str(bg_folder), "Relative_Path": str(file_path.relative_to(bg_folder)),
                    "File_Name": file_path.name, "Extension": file_path.suffix.casefold(),
                    "File_Size_MB": round(file_path.stat().st_size / (1024 * 1024), 3),
                    "Readable_By_Pipeline": "yes" if file_path.suffix.casefold() in READABLE_EXTENSIONS else "no",
                    "Preview_Mode": "", "Primary_Type": "", "Secondary_Types": "", "Confidence": None,
                    "Evidence": "", "Needs_Review": "yes", "Processing_Status": f"ERROR: {error}",
                    "Manual_Type": "", "Manual_Secondary_Types": "", "Review_Status": "not_reviewed",
                    "Raw_Model_Response": "", "Run_Model": args.model,
                    "Run_Timestamp": datetime.now().isoformat(timespec="seconds"),
                })
            save_workbook(rows, output_path)  # checkpoint after every file

    return output_path


def main() -> None:
    load_dotenv(PROJECT_DIR / ".env")
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s - %(levelname)s - %(message)s",
        handlers=[
            logging.FileHandler(args.output_dir / f"classify_bg_files_{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.log", encoding="utf-8"),
            logging.StreamHandler(),
        ],
        force=True,
    )
    output = run(args)
    print(f"Done: {output.resolve()}")


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        logging.error("File classification did not finish: %s", error)
        sys.exit(1)

# Command examples (run from Task3_Prompt_Development; place this script beside the connector):
# 1) Check the first 5 BG folders:
# python classify_bg_files_text_only.py --dataset-excel ".\input\classification_experiment_dataset_V2.xlsx" --max-bgs 5 --max-workers 8
# 2) Process all BG folders after the check:
# python classify_bg_files_text_only.py --dataset-excel ".\input\classification_experiment_dataset_V2.xlsx" --max-workers 8
#    Rows marked prompt_engineering = yes are excluded by default (final-test split).
# 3) Only when explicitly preparing the 14 prompt-development BGs as well:
# python classify_bg_files_text_only.py --dataset-excel ".\input\classification_experiment_dataset_V2.xlsx" --include-prompt-engineering --max-workers 8


"""Batch evaluation for completed PDF classification workbooks.

This script never calls an LLM.  It recalculates metrics from the first
result-data sheet of existing .xlsx files, then writes a ``Class_Metrics``
sheet to each workbook and produces two summary workbooks:

* repeatability_<model>_summary.xlsx
* pdf_cross_model_config_comparison.xlsx

The normal classification result sheets are preserved unchanged.
"""

from __future__ import annotations

import argparse
import ast
import json
import os
import re
import shutil
import tempfile
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent if SCRIPT_DIR.name.casefold() == "scripts" else SCRIPT_DIR
DEFAULT_OUTPUT_DIR = PROJECT_DIR / "evaluation_results" / "batch"

CANONICAL_CLASS_ORDER = [
    "Lineareinheit",
    "Multi-Achs-System (Gantry)",
    "Greifer",
    "Umsetzeinheit",
    "Roboter",
    "Rotationseinheit",
    "Keine der verfügbaren Klassen",
]

HEADER_CANDIDATES = {
    "sap": ("SAP-Nummer", "SAP Nummer", "SAP_Number", "SAP"),
    "teamcenter": ("Teamcenter", "Teamcenter-ID", "Teamcenter ID", "TC"),
    "ground_truth": ("Ground Truth", "Ground_Truth", "GroundTruth"),
    "register": ("Register", "Regime"),
    "alternatives": (
        "Weitere zulässige Ground Truth",
        "Weitere zulaessige Ground Truth",
        "Alternative Ground Truth",
    ),
    "prediction": ("Predicted_Label", "Predicted Label", "Prediction"),
    "confidence": ("Confidence_Percent", "Confidence Percent", "Confidence"),
    "processing": ("Processing_Status", "Processing Status"),
    "json": ("JSON_Parse_Status", "JSON Parse Status"),
    "model": ("Run_Model", "Run Model", "Model"),
    "prompt": ("Prompt_Config", "Prompt Config", "Prompt configuration"),
    "timestamp": ("Run_Timestamp", "Run Timestamp", "Timestamp"),
}


def clean(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def normalise_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", clean(value).casefold())


def find_column(frame: pd.DataFrame, key: str, required: bool = False) -> str | None:
    normalised = {normalise_header(column): column for column in frame.columns}
    for candidate in HEADER_CANDIDATES[key]:
        match = normalised.get(normalise_header(candidate))
        if match:
            return match
    if required:
        raise ValueError(
            f"Missing required column for {key}. Available columns: {list(frame.columns)}"
        )
    return None


def canonical_label(value: Any) -> str:
    """Map spelling variants to the project's seven canonical labels."""
    raw = clean(value)
    if not raw:
        return ""
    key = re.sub(r"[\s_\-]+", " ", raw.casefold()).strip()
    compact = re.sub(r"[^a-z0-9äöüß]+", "", key)

    if "umsetzeinheit" in compact or "kombiniert" in compact or "kombinierteeinheit" in compact:
        return "Umsetzeinheit"
    if "gantry" in compact or "multi" in compact and "achs" in compact or "portal" in compact:
        return "Multi-Achs-System (Gantry)"
    if "linear" in compact or "hubeinheit" in compact or "hubgeraet" in compact or "hubgerät" in compact:
        return "Lineareinheit"
    if "greifer" in compact or "saug" in compact or "gripper" in compact:
        return "Greifer"
    if "roboter" in compact or "robot" in compact:
        return "Roboter"
    if "rotation" in compact or "dreh" in compact:
        return "Rotationseinheit"
    if "keinederverfugbarenklassen" in compact or "keinederverfügbarenklassen" in compact or "nichtklassifizierbar" in compact:
        return "Keine der verfügbaren Klassen"
    return ""


def parse_label_collection(value: Any) -> set[str]:
    """Parse explicitly listed acceptable labels without using eval()."""
    if isinstance(value, (list, tuple, set)):
        values: Iterable[Any] = value
    else:
        raw = clean(value)
        if not raw:
            return set()
        # The combined-unit alias must remain one label, not two slash-separated labels.
        direct = canonical_label(raw)
        if direct:
            return {direct}
        values = [raw]
        for parser in (json.loads, ast.literal_eval):
            try:
                parsed = parser(raw)
            except (ValueError, SyntaxError, TypeError, json.JSONDecodeError):
                continue
            if isinstance(parsed, (list, tuple, set)):
                values = parsed
            else:
                values = [parsed]
            break

    labels: set[str] = set()
    for item in values:
        item_text = clean(item)
        direct = canonical_label(item_text)
        if direct:
            labels.add(direct)
            continue
        for part in re.split(r"[\n;,|]+", item_text):
            label = canonical_label(part)
            if label:
                labels.add(label)
    return labels


def numeric_confidence(value: Any) -> float | None:
    raw = clean(value).replace("%", "").replace(",", ".")
    if not raw:
        return None
    try:
        number = float(raw)
    except ValueError:
        return None
    return number * 100 if 0 <= number <= 1 else number


def identifier(row: pd.Series, sap_col: str | None, tc_col: str | None, row_number: int) -> str:
    sap = clean(row.get(sap_col)) if sap_col else ""
    tc = clean(row.get(tc_col)) if tc_col else ""
    if sap:
        return f"SAP:{sap}"
    if tc:
        return f"TC:{tc}"
    return f"ROW:{row_number}"


def locate_data_sheet(path: Path, requested: str | None) -> str:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if requested:
            if requested not in workbook.sheetnames:
                raise ValueError(f"Data sheet '{requested}' not found in {path.name}")
            return requested
        for name in workbook.sheetnames:
            header = [cell.value for cell in next(workbook[name].iter_rows(min_row=1, max_row=1))]
            header_keys = {normalise_header(item) for item in header}
            if normalise_header("Ground Truth") in header_keys and normalise_header("Predicted_Label") in header_keys:
                return name
    finally:
        workbook.close()
    raise ValueError(f"No result-data sheet with Ground Truth and Predicted_Label found in {path}")


def read_result(path: Path, data_sheet: str | None) -> tuple[pd.DataFrame, dict[str, str | None], str]:
    sheet = locate_data_sheet(path, data_sheet)
    frame = pd.read_excel(path, sheet_name=sheet, dtype=object)
    columns = {key: find_column(frame, key, required=key in {"ground_truth", "prediction"}) for key in HEADER_CANDIDATES}
    return frame, columns, sheet


def evaluate_rows(frame: pd.DataFrame, columns: dict[str, str | None]) -> tuple[pd.DataFrame, dict[str, int]]:
    records: list[dict[str, Any]] = []
    invalid_ground_truth = invalid_prediction = missing_prediction = 0
    for row_number, (_, row) in enumerate(frame.iterrows(), start=2):
        primary_raw = clean(row.get(columns["ground_truth"]))
        prediction_raw = clean(row.get(columns["prediction"]))
        primary = canonical_label(primary_raw)
        prediction = canonical_label(prediction_raw)
        register = clean(row.get(columns["register"])) .upper() if columns["register"] else ""
        alternatives = parse_label_collection(row.get(columns["alternatives"])) if columns["alternatives"] else set()
        accepted = {primary} | alternatives if primary else set()
        if not prediction_raw:
            missing_prediction += 1
        elif not prediction:
            invalid_prediction += 1
        if primary_raw and not primary:
            invalid_ground_truth += 1
        evaluable = bool(primary and prediction)
        strict = bool(evaluable and primary == prediction)
        accepted_correct = bool(evaluable and prediction in accepted)
        records.append({
            "BG_ID": identifier(row, columns["sap"], columns["teamcenter"], row_number),
            "SAP-Nummer": clean(row.get(columns["sap"])) if columns["sap"] else "",
            "Teamcenter": clean(row.get(columns["teamcenter"])) if columns["teamcenter"] else "",
            "Register": register,
            "Ground_Truth_Primary": primary,
            "Ground_Truth_Raw": primary_raw,
            "Accepted_Labels": " | ".join(sorted(accepted)),
            "Predicted_Label": prediction,
            "Predicted_Label_Raw": prediction_raw,
            "Confidence_Percent": numeric_confidence(row.get(columns["confidence"])) if columns["confidence"] else None,
            "Processing_Status": clean(row.get(columns["processing"])) if columns["processing"] else "",
            "JSON_Parse_Status": clean(row.get(columns["json"])) if columns["json"] else "",
            "Strict_Correct": strict if evaluable else None,
            "Accepted_Set_Correct": accepted_correct if evaluable else None,
            "Evaluable": evaluable,
        })
    cases = pd.DataFrame(records)
    diagnostics = {
        "Input_Rows": len(frame),
        "Evaluable_Rows": int(cases["Evaluable"].sum()),
        "Missing_Predictions": missing_prediction,
        "Unknown_Predictions": invalid_prediction,
        "Unknown_Ground_Truth": invalid_ground_truth,
        "Duplicate_BG_IDs": int(cases["BG_ID"].duplicated().sum()),
    }
    return cases, diagnostics


def safe_divide(numerator: float, denominator: float) -> float | None:
    return numerator / denominator if denominator else None


def class_metrics(cases: pd.DataFrame, scope: str) -> tuple[pd.DataFrame, dict[str, float | int | None]]:
    if scope == "e1e2":
        data = cases.loc[cases["Evaluable"] & cases["Register"].isin({"E1", "E2"})].copy()
    else:
        data = cases.loc[cases["Evaluable"]].copy()
    labels = CANONICAL_CLASS_ORDER + [
        item for item in sorted(set(data["Ground_Truth_Primary"]) | set(data["Predicted_Label"]))
        if item not in CANONICAL_CLASS_ORDER
    ]
    rows: list[dict[str, Any]] = []
    total = len(data)
    for label in labels:
        tp = int(((data["Ground_Truth_Primary"] == label) & (data["Predicted_Label"] == label)).sum())
        fp = int(((data["Ground_Truth_Primary"] != label) & (data["Predicted_Label"] == label)).sum())
        fn = int(((data["Ground_Truth_Primary"] == label) & (data["Predicted_Label"] != label)).sum())
        tn = total - tp - fp - fn
        precision = safe_divide(tp, tp + fp)
        recall = safe_divide(tp, tp + fn)
        f1 = safe_divide(2 * precision * recall, precision + recall) if precision is not None and recall is not None else None
        rows.append({
            "Class": label, "Evaluated N": total, "Support": tp + fn,
            "TP": tp, "TN": tn, "FP": fp, "FN": fn,
            "Accuracy": safe_divide(tp + tn, total), "Precision": precision,
            "Recall": recall, "F1": f1,
        })
    result = pd.DataFrame(rows)
    supported = result.loc[result["Support"] > 0].copy()
    overall_correct = int((data["Ground_Truth_Primary"] == data["Predicted_Label"]).sum())
    support_sum = int(supported["Support"].sum())
    summary: dict[str, float | int | None] = {
        "Evaluated N": total,
        "Strict correct": overall_correct,
        "Strict accuracy": safe_divide(overall_correct, total),
        "Macro Precision": supported["Precision"].mean() if not supported.empty else None,
        "Macro Recall": supported["Recall"].mean() if not supported.empty else None,
        "Macro F1": supported["F1"].mean() if not supported.empty else None,
        "Weighted Precision": safe_divide(float((supported["Precision"] * supported["Support"]).sum()), support_sum),
        "Weighted Recall": safe_divide(float((supported["Recall"] * supported["Support"]).sum()), support_sum),
        "Weighted F1": safe_divide(float((supported["F1"] * supported["Support"]).sum()), support_sum),
    }
    return result, summary


def regime_summary(cases: pd.DataFrame) -> dict[str, float | int | None]:
    output: dict[str, float | int | None] = {}
    evaluable = cases.loc[cases["Evaluable"]].copy()
    for regime in ("E1", "E2", "M"):
        selected = evaluable.loc[evaluable["Register"] == regime]
        output[f"{regime} evaluable N"] = len(selected)
        output[f"{regime} strict accuracy"] = safe_divide(int(selected["Strict_Correct"].sum()), len(selected))
    m_cases = evaluable.loc[evaluable["Register"] == "M"]
    output["M accepted-label-set accuracy"] = safe_divide(int(m_cases["Accepted_Set_Correct"].sum()), len(m_cases))
    output["All accepted-label-set accuracy"] = safe_divide(int(evaluable["Accepted_Set_Correct"].sum()), len(evaluable))
    return output


def style_sheet(sheet) -> None:
    title = PatternFill("solid", fgColor="1F4E78")
    header = PatternFill("solid", fgColor="2F75B5")
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.fill = title
        cell.font = Font(color="FFFFFF", bold=True)
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value in {"Class", "Evaluated N", "Support", "TP", "TN", "FP", "FN", "Accuracy", "Precision", "Recall", "F1"}:
                cell.fill = header
                cell.font = Font(color="FFFFFF", bold=True)
    for column in range(1, sheet.max_column + 1):
        letter = get_column_letter(column)
        max_length = max((len(clean(sheet.cell(row=row, column=column).value)) for row in range(1, sheet.max_row + 1)), default=10)
        sheet.column_dimensions[letter].width = min(max(max_length + 2, 13), 38)
    sheet.freeze_panes = "A2"


def write_table(sheet, start_row: int, title: str, frame: pd.DataFrame, summary: dict[str, Any] | None = None) -> int:
    sheet.cell(start_row, 1, title)
    sheet.cell(start_row, 1).font = Font(bold=True, color="FFFFFF")
    sheet.cell(start_row, 1).fill = PatternFill("solid", fgColor="1F4E78")
    row = start_row + 1
    if summary:
        for key, value in summary.items():
            sheet.cell(row, 1, key)
            cell = sheet.cell(row, 2, value)
            if "accuracy" in key.casefold() or "precision" in key.casefold() or "recall" in key.casefold() or "f1" in key.casefold():
                cell.number_format = "0.0%"
            row += 1
        row += 1
    for col, name in enumerate(frame.columns, start=1):
        cell = sheet.cell(row, col, name)
        cell.fill = PatternFill("solid", fgColor="2F75B5")
        cell.font = Font(color="FFFFFF", bold=True)
    for row_offset, values in enumerate(frame.itertuples(index=False, name=None), start=row + 1):
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row_offset, col, None if pd.isna(value) else value)
            if frame.columns[col - 1] in {"Accuracy", "Precision", "Recall", "F1"} and value is not None and not pd.isna(value):
                cell.number_format = "0.0%"
    return row + len(frame) + 2


def add_class_metrics(path: Path, cases: pd.DataFrame, diagnostics: dict[str, int], *, in_place: bool, output_dir: Path) -> Path:
    all_metrics, all_summary = class_metrics(cases, "all")
    e1e2_metrics, e1e2_summary = class_metrics(cases, "e1e2")
    regime = regime_summary(cases)
    target = path
    if not in_place:
        target = output_dir / "annotated_workbooks" / path.name
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, target)
    workbook = load_workbook(target)
    if "Class_Metrics" in workbook.sheetnames:
        del workbook["Class_Metrics"]
    sheet = workbook.create_sheet("Class_Metrics")
    sheet.sheet_view.showGridLines = False
    sheet["A1"] = "Batch class metrics (strict primary-label evaluation)"
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet.merge_cells("A1:D1")
    row = 3
    for key, value in {**diagnostics, **regime}.items():
        sheet.cell(row, 1, key)
        cell = sheet.cell(row, 2, value)
        if "accuracy" in key.casefold():
            cell.number_format = "0.0%"
        row += 1
    row += 1
    row = write_table(sheet, row, "All strict primary-label cases", all_metrics, all_summary)
    row = write_table(sheet, row, "E1+E2 strict primary-label cases", e1e2_metrics, e1e2_summary)
    sheet.cell(row, 1, "Accuracy uses (TP + TN) / Evaluated N; it is not the legacy class-wise recall.")
    for column in range(1, sheet.max_column + 1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = max(sheet.column_dimensions[letter].width or 0, 14)
    sheet.column_dimensions["A"].width = 42
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = None
    sheet.sheet_view.showGridLines = False
    temporary = target.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    workbook.close()
    os.replace(temporary, target)
    return target


def discover_workbooks(directory: Path) -> list[Path]:
    if not directory.is_dir():
        raise FileNotFoundError(f"Input directory does not exist: {directory}")
    ignored = ("~$", "summary", "comparison")
    return [
        path for path in sorted(directory.rglob("*.xlsx"))
        if not path.name.startswith("~$") and not any(token in path.name.casefold() for token in ignored)
    ]


def run_metadata(frame: pd.DataFrame, columns: dict[str, str | None], path: Path) -> dict[str, str]:
    def first(column_key: str) -> str:
        column = columns[column_key]
        if not column:
            return ""
        values = frame[column].map(clean)
        values = values.loc[values != ""]
        return values.iloc[0] if not values.empty else ""
    return {
        "Source_File": path.name,
        "Run_Model": first("model") or "unknown",
        "Prompt_Config": first("prompt") or "unknown",
        "Run_Timestamp": first("timestamp"),
    }


def evaluate_directory(directory: Path, data_sheet: str | None, output_dir: Path, in_place: bool) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    for path in discover_workbooks(directory):
        try:
            frame, columns, sheet = read_result(path, data_sheet)
            cases, diagnostics = evaluate_rows(frame, columns)
            all_metrics, all_summary = class_metrics(cases, "all")
            metadata = run_metadata(frame, columns, path)
            annotated = add_class_metrics(path, cases, diagnostics, in_place=in_place, output_dir=output_dir)
            errors = cases.loc[cases["Evaluable"] & ~cases["Strict_Correct"]].copy()
            results.append({
                "ok": True, "source": path, "annotated": annotated, "cases": cases,
                "metrics": all_metrics, "diagnostics": diagnostics, "summary": all_summary,
                "regime": regime_summary(cases), "metadata": metadata, "errors": errors,
                "data_sheet": sheet,
            })
            print(f"OK  {path.name}: {diagnostics['Evaluable_Rows']} evaluable rows")
        except Exception as error:  # keep other files running
            results.append({"ok": False, "source": path, "error": str(error)})
            print(f"FAIL {path.name}: {error}")
    return results


def append_dataframe(sheet, frame: pd.DataFrame, start_row: int = 1) -> int:
    for col, name in enumerate(frame.columns, start=1):
        cell = sheet.cell(start_row, col, name)
        cell.fill = PatternFill("solid", fgColor="2F75B5")
        cell.font = Font(color="FFFFFF", bold=True)
    for row_number, values in enumerate(frame.itertuples(index=False, name=None), start=start_row + 1):
        for col, value in enumerate(values, start=1):
            sheet.cell(row_number, col, None if pd.isna(value) else value)
    return start_row + len(frame) + 1


def finalise_workbook(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for column in range(1, sheet.max_column + 1):
            max_length = max((len(clean(sheet.cell(row=row, column=column).value)) for row in range(1, sheet.max_row + 1)), default=10)
            sheet.column_dimensions[get_column_letter(column)].width = min(max(max_length + 2, 13), 42)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_repeatability_summary(results: list[dict[str, Any]], output_dir: Path) -> Path | None:
    valid = [item for item in results if item["ok"]]
    if not valid:
        return None
    model = valid[0]["metadata"]["Run_Model"].replace("/", "_").replace(" ", "_")
    rows, per_class, errors = [], [], []
    prediction_map: dict[str, dict[str, str]] = {}
    base_info: dict[str, dict[str, Any]] = {}
    for run_index, item in enumerate(valid, start=1):
        run_name = f"Run_{run_index:02d}"
        row = {**item["metadata"], **item["diagnostics"], **item["summary"], **item["regime"]}
        rows.append(row)
        metric_frame = item["metrics"].copy()
        for key, value in item["metadata"].items():
            metric_frame[key] = value
        per_class.append(metric_frame)
        error_frame = item["errors"].copy()
        error_frame["Run"] = run_name
        error_frame["Source_File"] = item["metadata"]["Source_File"]
        errors.append(error_frame)
        for _, case in item["cases"].iterrows():
            bg_id = case["BG_ID"]
            prediction_map.setdefault(bg_id, {})[run_name] = case["Predicted_Label"] or "NOT_EVALUABLE"
            base_info.setdefault(bg_id, case.to_dict())

    run_frame = pd.DataFrame(rows)
    numeric_cols = run_frame.select_dtypes(include="number").columns
    aggregate = pd.DataFrame([
        {"Source_File": label, **{column: getattr(run_frame[column], method)() for column in numeric_cols}}
        for label, method in (("MEAN", "mean"), ("STD_DEV", "std"), ("MIN", "min"), ("MAX", "max"))
    ])
    reproducibility_rows = []
    run_columns = [f"Run_{index:02d}" for index in range(1, len(valid) + 1)]
    for bg_id, predictions in sorted(prediction_map.items()):
        labels = [predictions.get(run, "MISSING") for run in run_columns]
        counts = Counter(labels)
        mode, mode_count = counts.most_common(1)[0]
        info = base_info[bg_id]
        reproducibility_rows.append({
            "BG_ID": bg_id, "SAP-Nummer": info.get("SAP-Nummer", ""), "Teamcenter": info.get("Teamcenter", ""),
            "Register": info.get("Register", ""), "Ground_Truth": info.get("Ground_Truth_Primary", ""),
            **dict(zip(run_columns, labels)), "Distinct_Predictions": len(counts),
            "Modal_Label": mode, "Modal_Count": mode_count,
            "All_Runs_Identical": len(counts) == 1, "Prediction_Set": " | ".join(sorted(counts)),
        })
    reproducibility = pd.DataFrame(reproducibility_rows)
    workbook = Workbook()
    run_sheet = workbook.active
    run_sheet.title = "Run_Summary"
    append_dataframe(run_sheet, run_frame)
    append_dataframe(run_sheet, aggregate, start_row=len(run_frame) + 4)
    metric_sheet = workbook.create_sheet("Per_Class_Metrics")
    append_dataframe(metric_sheet, pd.concat(per_class, ignore_index=True))
    repro_sheet = workbook.create_sheet("BG_Reproducibility")
    total = len(reproducibility)
    identical = int(reproducibility["All_Runs_Identical"].sum()) if total else 0
    repro_sheet["A1"] = "Total BGs"; repro_sheet["B1"] = total
    repro_sheet["A2"] = "All runs identical"; repro_sheet["B2"] = identical
    repro_sheet["A3"] = "All-runs-identical rate"; repro_sheet["B3"] = safe_divide(identical, total); repro_sheet["B3"].number_format = "0.0%"
    append_dataframe(repro_sheet, reproducibility, start_row=5)
    error_sheet = workbook.create_sheet("Errors")
    append_dataframe(error_sheet, pd.concat(errors, ignore_index=True) if errors else pd.DataFrame())
    finalise_workbook(workbook)
    path = output_dir / f"repeatability_{model}_summary.xlsx"
    workbook.save(path)
    return path


def write_cross_model_summary(results: list[dict[str, Any]], output_dir: Path) -> Path | None:
    valid = [item for item in results if item["ok"]]
    if not valid:
        return None
    rows, per_class, errors, coverage = [], [], [], []
    baseline = set(valid[0]["cases"]["BG_ID"])
    for item in valid:
        metadata = item["metadata"]
        rows.append({**metadata, **item["diagnostics"], **item["summary"], **item["regime"]})
        metrics = item["metrics"].copy()
        for key, value in metadata.items():
            metrics[key] = value
        per_class.append(metrics)
        error_frame = item["errors"].copy()
        for key, value in metadata.items():
            error_frame[key] = value
        errors.append(error_frame)
        current = set(item["cases"]["BG_ID"])
        coverage.append({
            **metadata, "BG_ID_Count": len(current), "Matches_First_Run": current == baseline,
            "Missing_vs_First_Run": " | ".join(sorted(baseline - current)),
            "Extra_vs_First_Run": " | ".join(sorted(current - baseline)),
        })
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Run_Summary"
    append_dataframe(sheet, pd.DataFrame(rows))
    metric_sheet = workbook.create_sheet("Per_Class_Metrics")
    append_dataframe(metric_sheet, pd.concat(per_class, ignore_index=True))
    error_sheet = workbook.create_sheet("Errors")
    append_dataframe(error_sheet, pd.concat(errors, ignore_index=True) if errors else pd.DataFrame())
    coverage_sheet = workbook.create_sheet("Coverage_Check")
    append_dataframe(coverage_sheet, pd.DataFrame(coverage))
    finalise_workbook(workbook)
    path = output_dir / "pdf_cross_model_config_comparison.xlsx"
    workbook.save(path)
    return path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Batch-evaluate completed PDF classification Excel workbooks.")
    parser.add_argument("--repeatability-dir", type=Path, required=True, help="Folder with exactly the 10 Gemini P3-optimized result .xlsx files.")
    parser.add_argument("--cross-model-dir", type=Path, required=True, help="Folder with the 16 cross-model/configuration result .xlsx files.")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--data-sheet", default=None, help="Optional explicit name of the raw result sheet.")
    parser.add_argument("--in-place", action="store_true", help="Add/replace Class_Metrics directly in the source workbooks. Without it, annotated copies are created under output-dir.")
    parser.add_argument("--expected-repeatability-files", type=int, default=10)
    parser.add_argument("--expected-cross-model-files", type=int, default=16)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    repeatability_files = discover_workbooks(args.repeatability_dir)
    cross_model_files = discover_workbooks(args.cross_model_dir)
    if len(repeatability_files) != args.expected_repeatability_files:
        print(f"WARNING: expected {args.expected_repeatability_files} repeatability workbooks, found {len(repeatability_files)}")
    if len(cross_model_files) != args.expected_cross_model_files:
        print(f"WARNING: expected {args.expected_cross_model_files} cross-model workbooks, found {len(cross_model_files)}")
    repeatability = evaluate_directory(args.repeatability_dir, args.data_sheet, args.output_dir, args.in_place)
    cross_model = evaluate_directory(args.cross_model_dir, args.data_sheet, args.output_dir, args.in_place)
    repeatability_summary = write_repeatability_summary(repeatability, args.output_dir)
    cross_model_summary = write_cross_model_summary(cross_model, args.output_dir)
    print("\nFinished")
    print(f"Repeatability: {sum(item['ok'] for item in repeatability)}/{len(repeatability)} files succeeded")
    print(f"Cross-model: {sum(item['ok'] for item in cross_model)}/{len(cross_model)} files succeeded")
    print(f"Repeatability summary: {repeatability_summary}")
    print(f"Cross-model summary: {cross_model_summary}")


if __name__ == "__main__":
    main()


# PowerShell command (run from the Task3_Prompt_Development project folder).
# Close all result Excel files first.  Add the two missing Gemini P1/P2 result
# files to outputs\pdf_cross_model before running, so this folder contains 16
# workbooks in total.
#
# python .\scripts\batch_evaluate_experiments.py `
#   --repeatability-dir ".\outputs\valid_results_10_runs\gemini-2.5-pro_pdf_P3_optimized_10_runs" `
#   --cross-model-dir ".\outputs\pdf_cross_model" `
#   --output-dir ".\evaluation_results\batch" `
#   --in-place

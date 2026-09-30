from __future__ import annotations

"""Create a two-metric / four-label-config comparison chart from an .xlsx file.

The source worksheet is detected automatically. It must contain the columns
``Label_Config``, ``E1_E2_Strict_Accuracy`` and ``Allowed_Label_Accuracy``.
The original workbook is preserved; a new workbook is written by default.
"""

import argparse
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


OUTPUT_SHEET = "Label_Config_Comparison"
REQUIRED_HEADERS = (
    "Label_Config",
    "E1_E2_Strict_Accuracy",
    "Allowed_Label_Accuracy",
)
CONFIG_ORDER = ("L1", "L2", "L3", "L4")
CONFIG_COLORS = {
    "L1": "F7D36A",  # warm gold
    "L2": "EFA05A",  # muted orange
    "L3": "C96A45",  # terracotta
    "L4": "8B3A3A",  # deep wine red
}
HEADER_FILL = PatternFill("solid", fgColor="123D63")
HEADER_FONT = Font(bold=True, size=12, color="FFFFFF")
THIN_GRAY = Side(style="thin", color="D9E1F2")
TABLE_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
PERCENT_FMT = "0.0%"


def header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def chart_text(size: int) -> RichText:
    run = CharacterProperties(sz=size)  # Chart text is expressed in 1/100 pt.
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=run), endParaRPr=run)])


def find_source_sheet(wb):
    required = {header_key(header) for header in REQUIRED_HEADERS}
    matches = []
    for ws in wb.worksheets:
        found = {header_key(cell.value) for cell in ws[1]}
        if required.issubset(found):
            matches.append(ws)
    if len(matches) != 1:
        names = ", ".join(ws.title for ws in matches) or "none"
        raise ValueError(
            "Could not uniquely detect the source sheet. "
            f"Matching sheets: {names}."
        )
    return matches[0]


def column_for(ws, header: str) -> int:
    wanted = header_key(header)
    for cell in ws[1]:
        if header_key(cell.value) == wanted:
            return cell.column
    raise ValueError(f"Header {header!r} is missing from {ws.title!r}.")


def read_metrics(input_path: Path) -> dict[str, dict[str, float]]:
    wb = load_workbook(input_path, read_only=True, data_only=True)
    try:
        ws = find_source_sheet(wb)
        config_col = column_for(ws, "Label_Config")
        strict_col = column_for(ws, "E1_E2_Strict_Accuracy")
        allowed_col = column_for(ws, "Allowed_Label_Accuracy")
        metrics: dict[str, dict[str, float]] = {}

        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) < max(config_col, strict_col, allowed_col):
                continue
            raw_config = row[config_col - 1]
            if raw_config is None:
                continue
            config = str(raw_config).strip().upper()
            if config not in CONFIG_ORDER:
                continue
            strict, allowed = row[strict_col - 1], row[allowed_col - 1]
            if strict is None or allowed is None:
                raise ValueError(f"Missing Strict or Allowed value for {config}.")
            if config in metrics:
                raise ValueError(f"Duplicate row for {config}.")
            metrics[config] = {"strict": float(strict), "allowed": float(allowed)}

        missing = [config for config in CONFIG_ORDER if config not in metrics]
        if missing:
            raise ValueError("Missing label configurations: " + ", ".join(missing))
        for config, values in metrics.items():
            for name, value in values.items():
                if not 0 <= value <= 1:
                    raise ValueError(f"Invalid {name} value for {config}: {value!r}")
        return metrics
    finally:
        wb.close()


def add_comparison_sheet(output_path: Path, metrics: dict[str, dict[str, float]]) -> None:
    wb = load_workbook(output_path)
    if OUTPUT_SHEET in wb.sheetnames:
        del wb[OUTPUT_SHEET]
    ws = wb.create_sheet(OUTPUT_SHEET, 0)
    ws.sheet_view.showGridLines = False

    ws["A1"] = "E1/E2 Accuracy by Label Configuration"
    ws["A1"].font = Font(bold=True, size=18, color="123D63")
    ws.merge_cells("A1:F1")

    for col, value in enumerate(["Metric", *CONFIG_ORDER], start=1):
        cell = ws.cell(3, col, value)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, TABLE_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    metric_rows = (
        ("E1/E2 Strict Accuracy", "strict"),
        ("Allowed-Label Accuracy", "allowed"),
    )
    for row, (label, key) in enumerate(metric_rows, start=4):
        ws.cell(row, 1, label)
        for col, config in enumerate(CONFIG_ORDER, start=2):
            ws.cell(row, col, metrics[config][key])
        for cell in ws[row][:5]:
            cell.font = Font(size=12)
            cell.border = TABLE_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in range(2, 6):
            ws.cell(row, col).number_format = PERCENT_FMT

    ws.column_dimensions["A"].width = 27
    for col in "BCDE":
        ws.column_dimensions[col].width = 13
    ws.row_dimensions[3].height = 26
    for row in range(4, 6):
        ws.row_dimensions[row].height = 22

    chart = BarChart()
    chart.type, chart.grouping, chart.overlap, chart.gapWidth = "col", "clustered", -10, 120
    chart.title = None
    chart.y_axis.title = "Accuracy"
    chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 1
    chart.y_axis.numFmt = "0%"
    chart.legend.position = "b"
    chart.height, chart.width, chart.style = 12, 28, 10
    chart.x_axis.txPr = chart_text(1400)
    chart.y_axis.txPr = chart_text(1400)
    chart.legend.txPr = chart_text(1300)
    chart.y_axis.majorGridlines = ChartLines(
        spPr=GraphicalProperties(ln=LineProperties(solidFill="D9E2F3", w=6350))
    )
    chart.graphical_properties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))

    data = Reference(ws, min_col=2, max_col=5, min_row=3, max_row=5)
    chart.add_data(data, titles_from_data=True)
    category_formula = f"'{ws.title}'!$A$4:$A$5"
    for series, config in zip(chart.series, CONFIG_ORDER):
        series.cat = AxDataSource(strRef=StrRef(f=category_formula))
        series.graphicalProperties = GraphicalProperties(solidFill=CONFIG_COLORS[config])
    chart.x_axis.tickLblPos, chart.x_axis.tickLblSkip = "low", 1
    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    chart.dLbls.showCatName = False
    chart.dLbls.showLegendKey = False
    chart.dLbls.showSerName = False
    chart.dLbls.numFmt, chart.dLbls.dLblPos = PERCENT_FMT, "outEnd"
    chart.dLbls.txPr = chart_text(1200)
    ws.add_chart(chart, "A8")
    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Create an L1-L4 Strict/Allowed comparison chart from a processed .xlsx workbook."
    )
    parser.add_argument("input_file", type=Path, help="Absolute path to the source .xlsx workbook.")
    parser.add_argument("--output", type=Path, default=None, help="Optional absolute output .xlsx path.")
    args = parser.parse_args()
    input_path = args.input_file.expanduser().resolve()
    if not input_path.is_file():
        parser.error(f"Input file does not exist: {input_path}")
    output_path = (
        args.output.expanduser().resolve()
        if args.output
        else input_path.with_name(f"{input_path.stem}_with_Label_Config_Comparison.xlsx")
    )
    if input_path == output_path:
        parser.error("Output path must differ from input path to preserve the original workbook.")

    metrics = read_metrics(input_path)
    shutil.copy2(input_path, output_path)
    add_comparison_sheet(output_path, metrics)
    print(f"DONE: {output_path}")


if __name__ == "__main__":
    main()


# Examples (Windows PowerShell)
# The source worksheet name is detected automatically. It must contain:
# Label_Config, E1_E2_Strict_Accuracy and Allowed_Label_Accuracy.
# python .\create_label_config_comparison_chart.py `
#   "C:\Users\<Name>\RepoMA\outputs\label_config_comparison.xlsx"
#
# With an explicit absolute output path:
# python .\create_label_config_comparison_chart.py `
#   "C:\Users\<Name>\RepoMA\outputs\label_config_comparison.xlsx" `
#   --output "C:\Users\<Name>\RepoMA\outputs\label_config_comparison_with_chart.xlsx"

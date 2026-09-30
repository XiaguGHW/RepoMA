from __future__ import annotations

"""Create a P1-P4 / model comparison chart from a processed workbook.

The source worksheet name is detected automatically.  It only needs to
contain the headers ``Sheet`` and ``Strict Overall Accuracy``.
"""

from pathlib import Path
import argparse
import re
import shutil

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


OUTPUT_SHEET = "Cross_Model_Comparison"
PROMPTS = ("P1", "P2", "P3", "P4")
PROMPT_COLORS = {
    "P1": "B7DDF1",  # light blue
    "P2": "6EADD5",  # sky blue
    "P3": "2F75A7",  # medium blue
    "P4": "123D63",  # dark blue
}
KNOWN_MODELS = (
    ("claude-haiku-4-5", "Claude Haiku 4.5"),
    ("claude-opus-4-8", "Claude Opus 4.8"),
    ("gemini-2-5-flash", "Gemini 2.5 Flash"),
    ("gemini-2-5-pro", "Gemini 2.5 Pro"),
)

HEADER_FILL = PatternFill("solid", fgColor="123D63")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN_GRAY = Side(style="thin", color="D9E1F2")
TABLE_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
PERCENT_FMT = "0.0%"
CHART_TEXT_SIZE = 1400  # 14 pt; Excel chart text uses 1/100 pt.
DATA_LABEL_TEXT_SIZE = 1200  # 12 pt.


def chart_text(size: int) -> RichText:
    """Return a RichText style that Excel applies to chart labels."""
    run = CharacterProperties(sz=size)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=run), endParaRPr=run)])


def header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def find_source_sheet(wb):
    wanted = {header_key("Sheet"), header_key("Strict Overall Accuracy")}
    matches = []
    for ws in wb.worksheets:
        found = {header_key(cell.value) for cell in ws[1]}
        if wanted.issubset(found):
            matches.append(ws)
    if len(matches) != 1:
        names = ", ".join(ws.title for ws in matches) or "none"
        raise ValueError(
            "Could not uniquely detect the source summary sheet. "
            f"Matching sheets: {names}."
        )
    return matches[0]


def column_for(ws, header: str) -> int:
    wanted = header_key(header)
    for cell in ws[1]:
        if header_key(cell.value) == wanted:
            return cell.column
    raise ValueError(f"Header {header!r} is missing from {ws.title!r}.")


def prompt_and_model(sheet_value: object) -> tuple[str, str]:
    name = str(sheet_value or "").strip()
    match = re.match(r"^(P[1-4])_(.+)$", name, re.IGNORECASE)
    if not match:
        raise ValueError(f"Cannot parse P1-P4 from Sheet value {name!r}.")
    prompt = match.group(1).upper()
    raw = match.group(2).casefold().replace("_", "-").replace(".", "-")
    for key, display in KNOWN_MODELS:
        if key in raw:
            return prompt, display
    raw = re.sub(r"-20\d{6}(?:-\d{2}-\d{2}-\d{2})?$", "", raw)
    return prompt, re.sub(r"[-_]+", " ", raw).strip().title()


def read_matrix(input_path: Path) -> tuple[list[str], dict[str, dict[str, float]]]:
    wb = load_workbook(input_path, read_only=True, data_only=True)
    try:
        ws = find_source_sheet(wb)
        sheet_col = column_for(ws, "Sheet")
        accuracy_col = column_for(ws, "Strict Overall Accuracy")
        matrix: dict[str, dict[str, float]] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) < max(sheet_col, accuracy_col):
                continue
            sheet_value, accuracy_value = row[sheet_col - 1], row[accuracy_col - 1]
            if sheet_value is None or accuracy_value is None:
                continue
            prompt, model = prompt_and_model(sheet_value)
            accuracy = float(accuracy_value)
            if not 0 <= accuracy <= 1:
                raise ValueError(f"Invalid accuracy for {sheet_value!r}: {accuracy_value!r}")
            if prompt in matrix.setdefault(model, {}):
                raise ValueError(f"Duplicate result for {model} / {prompt}.")
            matrix[model][prompt] = accuracy

        missing = [f"{model} / {prompt}" for model, values in matrix.items()
                   for prompt in PROMPTS if prompt not in values]
        if missing:
            raise ValueError("Missing P1-P4 results: " + ", ".join(missing))
        preferred = [name for _, name in KNOWN_MODELS if name in matrix]
        return preferred + sorted(set(matrix).difference(preferred)), matrix
    finally:
        wb.close()


def add_comparison_sheet(output_path: Path, models: list[str], matrix: dict[str, dict[str, float]]) -> None:
    wb = load_workbook(output_path)
    if OUTPUT_SHEET in wb.sheetnames:
        del wb[OUTPUT_SHEET]
    ws = wb.create_sheet(OUTPUT_SHEET, 0)
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Strict Overall Accuracy – Model and Prompt Comparison"
    ws["A1"].font = Font(bold=True, size=18, color="123D63")
    ws.merge_cells("A1:H1")

    for col, value in enumerate(["Model", *PROMPTS], start=1):
        cell = ws.cell(3, col, value)
        cell.fill, cell.font, cell.border = HEADER_FILL, Font(bold=True, size=12, color="FFFFFF"), TABLE_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row, model in enumerate(models, start=4):
        ws.cell(row, 1, model)
        for col, prompt in enumerate(PROMPTS, start=2):
            ws.cell(row, col, matrix[model][prompt])
        for cell in ws[row][:5]:
            cell.border = TABLE_BORDER
            cell.font = Font(size=12)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in range(2, 6):
            ws.cell(row, col).number_format = PERCENT_FMT

    ws.column_dimensions["A"].width = 24
    for col in "BCDE":
        ws.column_dimensions[col].width = 13
    ws.row_dimensions[3].height = 26
    for row in range(4, 4 + len(models)):
        ws.row_dimensions[row].height = 22
    ws.freeze_panes = "B4"

    chart = BarChart()
    # More white space between model groups: 125 is over 1.5× the previous
    # setting (75). Prompt columns within each group remain clearly separated.
    chart.type, chart.grouping, chart.overlap, chart.gapWidth = "col", "clustered", -20, 125
    chart.title = None
    chart.y_axis.title = "Strict Overall Accuracy"
    chart.y_axis.scaling.min, chart.y_axis.scaling.max = 0, 1
    chart.y_axis.numFmt = "0%"
    # A right-side legend has enough room for four large labels and cannot
    # overlap the columns or the model names below the plot area.
    chart.legend.position = "r"
    chart.height, chart.width, chart.style = 15, 35, 10
    chart.x_axis.txPr = chart_text(CHART_TEXT_SIZE)
    chart.y_axis.txPr = chart_text(CHART_TEXT_SIZE)
    chart.legend.txPr = chart_text(CHART_TEXT_SIZE)
    chart.y_axis.majorGridlines = ChartLines(
        spPr=GraphicalProperties(ln=LineProperties(solidFill="D9E2F3", w=6350))
    )
    # Let the worksheet remain the visual canvas: no chart-area fill or frame.
    chart.graphical_properties = GraphicalProperties(
        noFill=True,
        ln=LineProperties(noFill=True),
    )
    data = Reference(ws, min_col=2, max_col=5, min_row=3, max_row=3 + len(models))
    chart.add_data(data, titles_from_data=True)
    # ``Reference`` serialises category labels as a numeric reference. Excel
    # can then omit text model names on the x-axis. A string reference makes
    # the four labels explicit: one underneath each model group.
    category_formula = "'{}'!$A$4:$A${}".format(ws.title, 3 + len(models))
    for series in chart.series:
        series.cat = AxDataSource(strRef=StrRef(f=category_formula))
    chart.x_axis.tickLblPos = "low"
    chart.x_axis.tickLblSkip = 1
    chart.dLbls = DataLabelList()
    chart.dLbls.showVal, chart.dLbls.showCatName = True, False
    chart.dLbls.showLegendKey, chart.dLbls.showSerName = False, False
    chart.dLbls.numFmt, chart.dLbls.dLblPos = PERCENT_FMT, "outEnd"
    chart.dLbls.txPr = chart_text(DATA_LABEL_TEXT_SIZE)
    for series, prompt in zip(chart.series, PROMPTS):
        series.graphicalProperties = GraphicalProperties(solidFill=PROMPT_COLORS[prompt])
    ws.add_chart(chart, "A10")
    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Create a P1-P4 comparison chart from a processed .xlsx workbook.")
    parser.add_argument("input_file", type=Path, help="Absolute path to the processed .xlsx workbook.")
    parser.add_argument("--output", type=Path, default=None, help="Optional absolute output .xlsx path.")
    args = parser.parse_args()
    input_path = args.input_file.expanduser().resolve()
    if not input_path.is_file():
        parser.error(f"Input file does not exist: {input_path}")
    output_path = args.output.expanduser().resolve() if args.output else input_path.with_name(f"{input_path.stem}_with_Comparison.xlsx")
    if input_path == output_path:
        parser.error("Output path must differ from input path to preserve the original workbook.")

    models, matrix = read_matrix(input_path)
    shutil.copy2(input_path, output_path)
    add_comparison_sheet(output_path, models, matrix)
    print(f"DONE: {output_path}")


if __name__ == "__main__":
    main()


# Examples (Windows PowerShell)
# The Summary worksheet name is detected automatically. Only the ABSOLUTE
# workbook path is required; the original workbook remains unchanged.
# python .\create_cross_model_comparison_chart.py `
#   "C:\Users\<Name>\RepoMA\outputs\pdf_cross_model\Evaluation_Summary.xlsx"
#
# With an explicit ABSOLUTE output path:
# python .\create_cross_model_comparison_chart.py `
#   "C:\Users\<Name>\RepoMA\outputs\pdf_cross_model\Evaluation_Summary.xlsx" `
#   --output "C:\Users\<Name>\RepoMA\outputs\pdf_cross_model\Evaluation_Summary_with_Comparison.xlsx"

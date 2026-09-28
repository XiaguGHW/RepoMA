from __future__ import annotations

"""
Batch evaluation for independent Excel classification experiments.

Place this script in the SAME folder as the experiment .xlsx files and run:
    python batch_evaluate_excel_experiments.py

Or pass one or more folders containing experiment .xlsx files.  Each folder
receives its own Evaluation_Summary.xlsx:
    python batch_evaluate_excel_experiments.py "C:\\path\\to\\experiment_A" "D:\\path\\to\\experiment_B"

Rules:
- Every .xlsx file is one independent experiment.
- Only worksheet "Sheet1" is read; all other worksheets are ignored.
- Ground Truth: header lookup first, fallback column F.
- Regime E1/E2/M: header lookup first, fallback column G.
- Prediction: header lookup first, fallback column I.
- Blank prediction = nicht auswertbar and excluded from classification metrics.
- E1/E2/M/Gesamt use strict primary Ground Truth.
- M confusion matrix also uses strict primary Ground Truth.
- Strict Overall Accuracy and Accepted-Set Overall Accuracy are separate metrics.
- For M, "Weitere zulässige Ground Truth" may list multiple allowed classes,
  separated by commas, semicolons, pipes, or line breaks.
- All official metrics are calculated with custom pure-Python logic.
- sklearn independently recalculates confusion matrix and P/R/F1 metrics as a cross-check.
"""

from pathlib import Path
from math import isclose
import argparse
import re

from openpyxl import load_workbook, Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    import numpy as np
    from sklearn.metrics import accuracy_score
    from sklearn.metrics import confusion_matrix as sk_confusion_matrix
    from sklearn.metrics import precision_recall_fscore_support
except ImportError as exc:
    raise SystemExit(
        "This script requires numpy and scikit-learn for mandatory metric verification. "
        "Install them with: pip install numpy scikit-learn"
    ) from exc


OUTPUT_FILENAME = "Evaluation_Summary.xlsx"
INPUT_SHEET_NAME = "Sheet1"

# Fixed class order for every confusion matrix and class-wise result.
CANONICAL_LABELS = [
    "Lineareinheit",
    "Greifer",
    "Roboter",
    "Rotationseinheit",
    "Multi-Achs-System (Gantry)",
    "Kombinierte Einheit",
    "Keine der verfügbaren Klassen",
]

LABEL_ALIASES_RAW = {
    "Lineareinheit": "Lineareinheit",
    "Greifer": "Greifer",
    "Roboter": "Roboter",
    "Rotationseinheit": "Rotationseinheit",

    # Same class:
    "Gantry": "Multi-Achs-System (Gantry)",
    "Multi-Achs-System (Gantry)": "Multi-Achs-System (Gantry)",

    # Same class:
    "Umsetzeinheit": "Kombinierte Einheit",
    "Kombinierte Einheit": "Kombinierte Einheit",

    "Keine der verfügbaren Klassen": "Keine der verfügbaren Klassen",
}

GROUND_TRUTH_HEADERS = ["Ground Truth", "Groundtruth", "Ground_Truth", "GT"]
REGIME_HEADERS = ["Register", "Regime", "Eindeutigkeit", "GT Regime", "Ground Truth Regime"]
PREDICTION_HEADERS = ["Predicted_Label", "Predicted Label", "Prediction", "LLM Output", "Output"]
ALTERNATIVE_GT_HEADERS = [
    "Weitere zulässige Ground Truth",
    "Weitere zulaessige Ground Truth",
    "Alternative Ground Truth",
    "Alternative GT",
    "Accepted Alternative",
    "Weitere Ground Truth",
]

FALLBACK_GT_COL = 6
FALLBACK_REGIME_COL = 7
FALLBACK_PRED_COL = 9
VALID_REGIMES = ("E1", "E2", "M")
SCOPES = ("E1", "E2", "M", "Gesamt")

# Muted, report-friendly colors. The white outlines are configured directly
# on each data point so adjacent pie segments remain visually distinct.
PIE_COLORS = {
    "Strict Correct": "278C3F",  # Bosch green
    "Strict Wrong": "D51317",    # Bosch red
    "Nicht auswertbar": "AEB5BF",
}
PIE_OUTLINE_COLOR = "FFFFFF"
PIE_OUTLINE_WIDTH = 19050  # 1.5 pt (OOXML units: 1 pt = 12,700)


def norm_text(value) -> str:
    if value is None:
        return ""
    text = str(value).strip().replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", text)


def key(value) -> str:
    return re.sub(r"[^a-z0-9äöüß]+", "", norm_text(value).casefold())


LABEL_ALIASES = {key(k): v for k, v in LABEL_ALIASES_RAW.items()}
for label in CANONICAL_LABELS:
    LABEL_ALIASES.setdefault(key(label), label)


def normalize_label(value):
    raw = norm_text(value)
    if not raw:
        return None

    k = key(raw)
    if k in LABEL_ALIASES:
        return LABEL_ALIASES[k]

    # Conservative fallback: only accept if exactly one known class alias
    # occurs in the raw text.
    matched = {
        canonical
        for alias_key, canonical in LABEL_ALIASES.items()
        if alias_key and alias_key in k
    }
    if len(matched) == 1:
        return next(iter(matched))

    return f"__UNKNOWN__:{raw}"


ALT_LABEL_SEPARATORS = re.compile(r"[,;|\r\n]+")


def normalize_label_set(value):
    """Parse the zero or more allowed alternative labels from one Excel cell."""
    if value is None or not norm_text(value):
        return set()

    labels = set()
    for raw_label in ALT_LABEL_SEPARATORS.split(str(value)):
        label = normalize_label(raw_label)
        if label is not None:
            labels.add(label)
    return labels


def normalize_regime(value):
    value = norm_text(value).upper().replace(" ", "")
    return value if value in VALID_REGIMES else None


def find_header_column(ws, candidates):
    wanted = {key(x) for x in candidates}
    for cell in ws[1]:
        if key(cell.value) in wanted:
            return cell.column
    return None


def read_experiment(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if INPUT_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f'worksheet "{INPUT_SHEET_NAME}" not found')

        ws = wb[INPUT_SHEET_NAME]
        gt_col = find_header_column(ws, GROUND_TRUTH_HEADERS) or FALLBACK_GT_COL
        regime_col = find_header_column(ws, REGIME_HEADERS) or FALLBACK_REGIME_COL
        pred_col = find_header_column(ws, PREDICTION_HEADERS) or FALLBACK_PRED_COL
        alt_col = find_header_column(ws, ALTERNATIVE_GT_HEADERS)

        rows = []
        for r, row_cells in enumerate(ws.iter_rows(min_row=2), start=2):
            def row_value(col_idx):
                if not col_idx or col_idx < 1 or col_idx > len(row_cells):
                    return None
                return row_cells[col_idx - 1].value

            raw_gt = row_value(gt_col)
            raw_regime = row_value(regime_col)
            raw_pred = row_value(pred_col)
            raw_alt = row_value(alt_col) if alt_col else None

            if all(norm_text(x) == "" for x in (raw_gt, raw_regime, raw_pred, raw_alt)):
                continue

            gt = normalize_label(raw_gt)
            regime = normalize_regime(raw_regime)
            pred = normalize_label(raw_pred)
            alt_labels = normalize_label_set(raw_alt)

            row_valid = gt is not None and regime is not None
            rows.append({
                "excel_row": r,
                "raw_gt": raw_gt,
                "raw_regime": raw_regime,
                "raw_pred": raw_pred,
                "raw_alt": raw_alt,
                "gt": gt,
                "regime": regime,
                "pred": pred,
                "alt_labels": alt_labels,
                "row_valid": row_valid,
                "issue": "" if row_valid else "Missing/invalid Ground Truth or Regime",
            })

        return rows, {
            "gt_col": gt_col,
            "regime_col": regime_col,
            "pred_col": pred_col,
            "alt_col": alt_col,
        }
    finally:
        wb.close()


def safe_div(a, b):
    return a / b if b else 0.0


def build_confusion_matrix(rows):
    idx = {label: i for i, label in enumerate(CANONICAL_LABELS)}
    n = len(CANONICAL_LABELS)
    cm = [[0] * n for _ in range(n)]

    for row in rows:
        gt, pred = row["gt"], row["pred"]
        if gt not in idx:
            raise ValueError(f"Unknown Ground Truth after normalization: {gt!r}")
        if pred not in idx:
            raise ValueError(f"Unknown prediction after normalization: {pred!r}")
        cm[idx[gt]][idx[pred]] += 1

    return cm


def metrics_from_cm(cm):
    n_classes = len(CANONICAL_LABELS)
    n_samples = sum(sum(r) for r in cm)
    class_metrics = []

    for c, label in enumerate(CANONICAL_LABELS):
        tp = cm[c][c]
        fn = sum(cm[c][j] for j in range(n_classes) if j != c)
        fp = sum(cm[i][c] for i in range(n_classes) if i != c)
        tn = n_samples - tp - fn - fp

        assert tp + tn + fp + fn == n_samples

        support = tp + fn
        precision = safe_div(tp, tp + fp)
        recall = safe_div(tp, tp + fn)
        f1 = safe_div(2 * precision * recall, precision + recall)
        accuracy_ovr = safe_div(tp + tn, n_samples)

        class_metrics.append({
            "class": label,
            "support": support,
            "TP": tp,
            "TN": tn,
            "FP": fp,
            "FN": fn,
            "accuracy_ovr": accuracy_ovr,
            "precision": precision,
            "recall": recall,
            "f1": f1,
        })

    total_support = sum(x["support"] for x in class_metrics)
    correct = sum(cm[i][i] for i in range(n_classes))

    macro_precision = sum(x["precision"] for x in class_metrics) / n_classes
    macro_recall = sum(x["recall"] for x in class_metrics) / n_classes
    macro_f1 = sum(x["f1"] for x in class_metrics) / n_classes

    weighted_precision = safe_div(
        sum(x["precision"] * x["support"] for x in class_metrics), total_support
    )
    weighted_recall = safe_div(
        sum(x["recall"] * x["support"] for x in class_metrics), total_support
    )
    weighted_f1 = safe_div(
        sum(x["f1"] * x["support"] for x in class_metrics), total_support
    )

    sum_tp = sum(x["TP"] for x in class_metrics)
    sum_tn = sum(x["TN"] for x in class_metrics)
    sum_fp = sum(x["FP"] for x in class_metrics)
    sum_fn = sum(x["FN"] for x in class_metrics)
    assert sum_tp + sum_tn + sum_fp + sum_fn == n_samples * n_classes

    return {
        "n": n_samples,
        "correct": correct,
        "accuracy": safe_div(correct, n_samples),
        "macro_precision": macro_precision,
        "macro_recall": macro_recall,
        "macro_f1": macro_f1,
        "weighted_precision": weighted_precision,
        "weighted_recall": weighted_recall,
        "weighted_f1": weighted_f1,
        "sum_TP": sum_tp,
        "sum_TN": sum_tn,
        "sum_FP": sum_fp,
        "sum_FN": sum_fn,
        "class_metrics": class_metrics,
    }


def evaluate_scope(rows, regime=None):
    population = [
        r for r in rows
        if r["row_valid"] and (regime is None or r["regime"] == regime)
    ]
    evaluated = [r for r in population if r["pred"] is not None]

    # Unknown nonblank labels must never silently disappear.
    for r in evaluated:
        if str(r["gt"]).startswith("__UNKNOWN__"):
            raise ValueError(f"Row {r['excel_row']}: unknown GT {r['raw_gt']!r}")
        if str(r["pred"]).startswith("__UNKNOWN__"):
            raise ValueError(f"Row {r['excel_row']}: unknown prediction {r['raw_pred']!r}")

    cm = build_confusion_matrix(evaluated)
    metrics = metrics_from_cm(cm)
    return population, evaluated, cm, metrics


def overall_accuracies(rows):
    population = [r for r in rows if r["row_valid"]]
    evaluated = [r for r in population if r["pred"] is not None]

    for r in evaluated:
        if str(r["gt"]).startswith("__UNKNOWN__"):
            raise ValueError(f"Row {r['excel_row']}: unknown GT {r['raw_gt']!r}")
        if str(r["pred"]).startswith("__UNKNOWN__"):
            raise ValueError(f"Row {r['excel_row']}: unknown prediction {r['raw_pred']!r}")
        unknown_alternatives = [
            label for label in r["alt_labels"] if str(label).startswith("__UNKNOWN__")
        ]
        if r["regime"] == "M" and unknown_alternatives:
            raise ValueError(
                f"Row {r['excel_row']}: unknown alternative GT "
                f"{unknown_alternatives!r} in {r['raw_alt']!r}"
            )

    strict_correct = sum(r["pred"] == r["gt"] for r in evaluated)

    accepted_correct = 0
    for r in evaluated:
        accepted = {r["gt"]}
        # Alternative GT changes correctness only for M.
        if r["regime"] == "M":
            accepted.update(r["alt_labels"])
        if r["pred"] in accepted:
            accepted_correct += 1

    n_total = len(population)
    n_eval = len(evaluated)

    return {
        "n_total": n_total,
        "n_auswertbar": n_eval,
        "n_nicht_auswertbar": n_total - n_eval,
        "coverage": safe_div(n_eval, n_total),
        "strict_correct": strict_correct,
        "strict_accuracy": safe_div(strict_correct, n_eval),
        "accepted_correct": accepted_correct,
        "accepted_accuracy": safe_div(accepted_correct, n_eval),
    }


def sklearn_cross_check(rows, custom_cm, custom_metrics, context):
    """Independent sklearn validation of the custom strict multiclass metrics."""
    # sklearn's multiclass helpers reject empty inputs.  An empty regime is a
    # valid experiment outcome, however: the custom definition is an all-zero
    # matrix with all metrics set to 0.0.  Validate that definition directly.
    if not rows:
        if any(any(value for value in matrix_row) for matrix_row in custom_cm):
            raise AssertionError(f"{context}: non-zero confusion matrix for empty scope")
        if custom_metrics["n"] != 0:
            raise AssertionError(f"{context}: non-zero sample count for empty scope")
        metric_names = (
            "accuracy", "macro_precision", "macro_recall", "macro_f1",
            "weighted_precision", "weighted_recall", "weighted_f1",
        )
        if any(custom_metrics[name] != 0.0 for name in metric_names):
            raise AssertionError(f"{context}: non-zero metric for empty scope")
        return "EMPTY: PASS"

    y_true = [r["gt"] for r in rows]
    y_pred = [r["pred"] for r in rows]

    sk_cm = sk_confusion_matrix(y_true, y_pred, labels=CANONICAL_LABELS)
    if not np.array_equal(np.array(custom_cm, dtype=int), sk_cm):
        raise AssertionError(f"{context}: custom confusion matrix != sklearn")

    sk_acc = float(accuracy_score(y_true, y_pred))
    if not isclose(custom_metrics["accuracy"], sk_acc, abs_tol=1e-12):
        raise AssertionError(f"{context}: accuracy mismatch")

    p, r, f1, support = precision_recall_fscore_support(
        y_true,
        y_pred,
        labels=CANONICAL_LABELS,
        average=None,
        zero_division=0,
    )

    for i, custom in enumerate(custom_metrics["class_metrics"]):
        if custom["support"] != int(support[i]):
            raise AssertionError(f"{context}: support mismatch for {custom['class']}")
        if not isclose(custom["precision"], float(p[i]), abs_tol=1e-12):
            raise AssertionError(f"{context}: precision mismatch for {custom['class']}")
        if not isclose(custom["recall"], float(r[i]), abs_tol=1e-12):
            raise AssertionError(f"{context}: recall mismatch for {custom['class']}")
        if not isclose(custom["f1"], float(f1[i]), abs_tol=1e-12):
            raise AssertionError(f"{context}: F1 mismatch for {custom['class']}")

    for avg, names in [
        ("macro", ("macro_precision", "macro_recall", "macro_f1")),
        ("weighted", ("weighted_precision", "weighted_recall", "weighted_f1")),
    ]:
        sp, sr, sf, _ = precision_recall_fscore_support(
            y_true,
            y_pred,
            labels=CANONICAL_LABELS,
            average=avg,
            zero_division=0,
        )
        for actual, name in zip((sp, sr, sf), names):
            if not isclose(custom_metrics[name], float(actual), abs_tol=1e-12):
                raise AssertionError(f"{context}: {name} mismatch")

    return "PASS"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
SUBHEADER_FONT = Font(bold=True)
DIAGONAL_FILL = PatternFill("solid", fgColor="DDEBF7")
SECTION_FILL = PatternFill("solid", fgColor="DDEBF7")
SECTION_FONT = Font(bold=True, size=12)
PERCENT_FMT = "0.0%"
THIN_GRAY = Side(style="thin", color="D9E1F2")
TABLE_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)


def style_header_range(ws, row, start_col, end_col, dark=True):
    fill = HEADER_FILL if dark else SUBHEADER_FILL
    font = HEADER_FONT if dark else SUBHEADER_FONT
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row, col)
        if cell.value is not None:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = TABLE_BORDER


def style_table_cells(ws, min_row, max_row, min_col, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = TABLE_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def section_title(ws, row, col, text, span=6):
    ws.cell(row, col, text)
    for c in range(col, col + span):
        ws.cell(row, c).fill = SECTION_FILL
        ws.cell(row, c).font = SECTION_FONT
    return row


def auto_width(ws, min_width=10, max_width=32):
    for col_cells in ws.columns:
        col_idx = col_cells[0].column
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(
            max(max_len + 2, min_width), max_width
        )


def fit_wrapped_rows(ws, min_row=1, max_row=None, min_height=15, max_height=75):
    """Set practical row heights after column widths and text wrapping are applied.

    openpyxl stores wrapped text but does not calculate Excel's automatic row
    height.  This keeps long headers and notes readable in Excel and LibreOffice
    without making short rows unnecessarily tall.
    """
    if max_row is None:
        max_row = ws.max_row

    for row in range(min_row, max_row + 1):
        required_lines = 1
        for cell in ws[row]:
            if cell.value is None or not cell.alignment.wrap_text:
                continue

            column_letter = get_column_letter(cell.column)
            if ws.column_dimensions[column_letter].hidden:
                continue

            width = ws.column_dimensions[column_letter].width or 10
            chars_per_line = max(int(width * 1.1), 1)
            required_lines = max(
                required_lines,
                max(
                    (len(line) + chars_per_line - 1) // chars_per_line
                    for line in str(cell.value).splitlines() or [""]
                ),
            )

        ws.row_dimensions[row].height = min(
            max(min_height, required_lines * 15),
            max_height,
        )


def print_progress(current, total, label, width=28):
    """Render one compact terminal progress bar without extra dependencies."""
    fraction = current / total if total else 1.0
    filled = round(width * fraction)
    bar = "#" * filled + "-" * (width - filled)
    print(
        f"\r{label} [{bar}] {current}/{total} ({fraction:.0%})",
        end="\n" if current >= total else "",
        flush=True,
    )


def extract_experiment_sheet_base(filename_or_stem):
    """
    Extract a compact sheet name from experiment filenames like:
      prompt_development_P1_gemini-2.5-pro_caching_2026-09-15_17-15-50.xlsx
      -> P1_gemini-2.5-pro

      prompt_development_P4_claude-haiku-4-5_20251001_caching_...
      -> P4_claude-haiku-4-5_20251001

    If the expected pattern is not found, fall back to the original stem.
    """
    name = Path(filename_or_stem).stem
    match = re.search(
        r"(?:^|_)P([1-4])_(.+?)_caching(?:_|$)",
        name,
        flags=re.IGNORECASE,
    )
    if not match:
        return name

    prompt_id = f"P{match.group(1)}"
    model_name = match.group(2).strip("_")
    return f"{prompt_id}_{model_name}"


def make_unique_sheet_name(raw_name, used_names):
    cleaned = re.sub(r"[\[\]:*?/\\]", "_", raw_name).strip().strip("'")
    if not cleaned:
        cleaned = "Experiment"
    cleaned = cleaned[:31]

    candidate = cleaned
    counter = 2
    while candidate.casefold() in used_names:
        suffix = f"_{counter}"
        candidate = f"{cleaned[:31-len(suffix)]}{suffix}"
        counter += 1
    used_names.add(candidate.casefold())
    return candidate


def add_overview_counts(ws, res, start_row=1, start_col=1):
    overall = res["overall"]
    total_metrics = res["scopes"]["Gesamt"]["metrics"]
    section_title(ws, start_row, start_col, "Experiment Overview / Counts", span=4)
    header_row = start_row + 1
    headers = ["Metric", "Value", "Notes"]
    for j, h in enumerate(headers, start=start_col):
        ws.cell(header_row, j, h)
    style_header_range(ws, header_row, start_col, start_col + 2)

    values = [
        ("N total", overall["n_total"], "Valid Ground Truth + Regime rows"),
        ("N auswertbar", overall["n_auswertbar"], "Nonblank prediction"),
        ("Total TP", total_metrics["sum_TP"], "Sum across 7 one-vs-rest classes; strict primary GT"),
        ("Total TN", total_metrics["sum_TN"], "Sum across 7 one-vs-rest classes; strict primary GT"),
        ("Total FP", total_metrics["sum_FP"], "Sum across 7 one-vs-rest classes; strict primary GT"),
        ("Total FN", total_metrics["sum_FN"], "Sum across 7 one-vs-rest classes; strict primary GT"),
        ("Strict Correct", overall["strict_correct"], "Prediction exactly matches primary Ground Truth"),
        (
            "Accepted-Set Correct",
            overall["accepted_correct"],
            "Primary Ground Truth plus all listed M alternative labels",
        ),
        ("N nicht auswertbar", overall["n_nicht_auswertbar"], "Usually blank prediction / technical failure"),
        ("Coverage", overall["coverage"], "N auswertbar / N total"),
    ]
    row = header_row + 1
    for metric, value, notes in values:
        ws.cell(row, start_col, metric)
        ws.cell(row, start_col + 1, value)
        ws.cell(row, start_col + 2, notes)
        if metric == "Coverage":
            ws.cell(row, start_col + 1).number_format = PERCENT_FMT
        row += 1

    style_table_cells(ws, header_row, row - 1, start_col, start_col + 2)
    return row + 1


def add_primary_overall_metrics(ws, res, start_row=1, start_col=1):
    overall = res["overall"]
    m = res["scopes"]["Gesamt"]["metrics"]

    section_title(ws, start_row, start_col, "Primary Overall Metrics", span=4)
    header_row = start_row + 1
    headers = ["Metric", "Value", "Details"]
    for j, h in enumerate(headers, start=start_col):
        ws.cell(header_row, j, h)
    style_header_range(ws, header_row, start_col, start_col + 2)

    values = [
        (
            "Strict Overall Accuracy",
            overall["strict_accuracy"],
            f'{overall["strict_correct"]}/{overall["n_auswertbar"]}',
        ),
        (
            "Accepted-Set Overall Accuracy",
            overall["accepted_accuracy"],
            f'{overall["accepted_correct"]}/{overall["n_auswertbar"]}',
        ),
        (
            "Gesamt Macro Precision",
            m["macro_precision"],
            "Mean of the 7 class Prec. values: (1/7) * Σ TP/(TP+FP); strict primary GT; zero division = 0",
        ),
        (
            "Gesamt Macro Recall",
            m["macro_recall"],
            "Mean of the 7 class Rec. values: (1/7) * Σ TP/(TP+FN); strict primary GT; zero division = 0",
        ),
        (
            "Gesamt Macro F1",
            m["macro_f1"],
            "Mean of the 7 per-class F1 values: (1/7) * Σ 2*Prec.*Rec./(Prec.+Rec.); strict primary GT; zero division = 0",
        ),
        ("sklearn verification (Gesamt)", res["scopes"]["Gesamt"]["sklearn_status"], ""),
    ]

    row = header_row + 1
    for metric, value, details in values:
        ws.cell(row, start_col, metric)
        ws.cell(row, start_col + 1, value)
        ws.cell(row, start_col + 2, details)
        if metric != "sklearn verification (Gesamt)":
            ws.cell(row, start_col + 1).number_format = PERCENT_FMT
        row += 1

    style_table_cells(ws, header_row, row - 1, start_col, start_col + 2)
    return row + 1


def add_weighted_overall_metrics(ws, res, start_row=1, start_col=1):
    m = res["scopes"]["Gesamt"]["metrics"]

    section_title(ws, start_row, start_col, "Gesamt Weighted Metrics", span=4)
    header_row = start_row + 1
    headers = ["Metric", "Value"]
    for j, h in enumerate(headers, start=start_col):
        ws.cell(header_row, j, h)
    style_header_range(ws, header_row, start_col, start_col + 1)

    values = [
        ("Gesamt Weighted Precision", m["weighted_precision"]),
        ("Gesamt Weighted Recall", m["weighted_recall"]),
        ("Gesamt Weighted F1", m["weighted_f1"]),
    ]

    row = header_row + 1
    for metric, value in values:
        ws.cell(row, start_col, metric)
        ws.cell(row, start_col + 1, value)
        ws.cell(row, start_col + 1).number_format = PERCENT_FMT
        row += 1

    style_table_cells(ws, header_row, row - 1, start_col, start_col + 1)
    return row + 1


def add_overview_pie_chart(ws, res, anchor="F2"):
    overall = res["overall"]
    strict_wrong = overall["n_auswertbar"] - overall["strict_correct"]

    start_row = 2
    label_col = 27
    value_col = 28
    pie_rows = [
        ("Strict Correct", overall["strict_correct"]),
        ("Strict Wrong", strict_wrong),
        ("Nicht auswertbar", overall["n_nicht_auswertbar"]),
    ]

    for i, (label, value) in enumerate(pie_rows, start=start_row):
        ws.cell(i, label_col, label)
        ws.cell(i, value_col, value)

    labels = Reference(ws, min_col=label_col, min_row=start_row, max_row=start_row + 2)
    data = Reference(ws, min_col=value_col, min_row=start_row, max_row=start_row + 2)

    chart = PieChart()
    chart.add_data(data, titles_from_data=False)
    chart.set_categories(labels)
    chart.title = None
    # Keep the chart compact; Excel retains the data-label font size.
    # A wide canvas leaves dedicated empty space for the lower-right legend.
    # 14.0 / 5.3 = 2.64, i.e. wider than the requested 2.5:1 aspect ratio.
    chart.height = 5.3
    chart.width = 14.0
    chart.firstSliceAng = 270
    chart.legend.position = "r"
    chart.legend.layout = Layout(
        manualLayout=ManualLayout(
            layoutTarget="outer",
            xMode="factor",
            yMode="factor",
            wMode="factor",
            hMode="factor",
            x=0.62,
            y=0.70,
            w=0.34,
            h=0.25,
        )
    )
    # Remove Excel's default chart-area border: the worksheet remains the
    # visual canvas and the pie reads as a clean dashboard element.
    chart.graphical_properties = GraphicalProperties(
        noFill=True,
        ln=LineProperties(noFill=True),
    )
    # The source cells are hidden in AA:AB.  Excel otherwise omits hidden
    # cells from the chart and renders an empty chart area.
    chart.visible_cells_only = False
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.showLegendKey = False
    chart.dataLabels.showPercent = False
    chart.dataLabels.showCatName = False
    chart.dataLabels.showSerName = False
    # Values are placed directly in the pie segments; the legend below
    # provides the category names without adding overlapping long labels.
    chart.dataLabels.showLeaderLines = False
    chart.dataLabels.dLblPos = "ctr"
    chart.series[0].data_points = [
        DataPoint(
            idx=i,
            explosion=0,
            spPr=GraphicalProperties(
                solidFill=PIE_COLORS[label],
                ln=LineProperties(
                    solidFill=PIE_OUTLINE_COLOR,
                    w=PIE_OUTLINE_WIDTH,
                ),
            ),
        )
        for i, (label, _) in enumerate(pie_rows)
    ]
    ws.add_chart(chart, anchor)

    ws.column_dimensions["AA"].hidden = True
    ws.column_dimensions["AB"].hidden = True


def add_scope_metrics(ws, res, scope, start_row, start_col=1):
    m = res["scopes"][scope]["metrics"]

    section_title(ws, start_row, start_col, f"{scope} Evaluation", span=9)
    header_row = start_row + 1
    headers = ["Metric", "Value"]
    for j, h in enumerate(headers, start=start_col):
        ws.cell(header_row, j, h)
    style_header_range(ws, header_row, start_col, start_col + 1)

    values = [
        (f"{scope} N auswertbar", m["n"]),
        (f"{scope} Accuracy", m["accuracy"]),
        (f"{scope} Macro Precision", m["macro_precision"]),
        (f"{scope} Macro Recall", m["macro_recall"]),
        (f"{scope} Macro F1", m["macro_f1"]),
        (f"{scope} Weighted Precision", m["weighted_precision"]),
        (f"{scope} Weighted Recall", m["weighted_recall"]),
        (f"{scope} Weighted F1", m["weighted_f1"]),
        (f"{scope} sklearn verification", res["scopes"][scope]["sklearn_status"]),
    ]

    row = header_row + 1
    for metric, value in values:
        ws.cell(row, start_col, metric)
        ws.cell(row, start_col + 1, value)
        if metric not in (f"{scope} N auswertbar", f"{scope} sklearn verification"):
            ws.cell(row, start_col + 1).number_format = PERCENT_FMT
        row += 1

    style_table_cells(ws, header_row, row - 1, start_col, start_col + 1)
    return row + 1


def add_confusion_matrix(ws, scope, cm, start_row, start_col=1):
    section_title(ws, start_row, start_col, f"{scope} Confusion Matrix", span=8)
    header_row = start_row + 1
    ws.cell(header_row, start_col, "True \\ Pred")
    for j, label in enumerate(CANONICAL_LABELS, start=start_col + 1):
        ws.cell(header_row, j, label)
    style_header_range(
        ws,
        header_row,
        start_col,
        start_col + len(CANONICAL_LABELS),
        dark=False,
    )

    for i, label in enumerate(CANONICAL_LABELS):
        r = header_row + 1 + i
        ws.cell(r, start_col, label)
        ws.cell(r, start_col).font = SUBHEADER_FONT
        ws.cell(r, start_col).fill = SUBHEADER_FILL
        for j, value in enumerate(cm[i], start=start_col + 1):
            cell = ws.cell(r, j, value)
            if i == j - (start_col + 1):
                cell.fill = DIAGONAL_FILL

    last_row = header_row + len(CANONICAL_LABELS)
    style_table_cells(
        ws,
        header_row,
        last_row,
        start_col,
        start_col + len(CANONICAL_LABELS),
    )
    return last_row + 2


def add_class_metrics(ws, scope, class_metrics, start_row, start_col=1):
    section_title(ws, start_row, start_col, f"{scope} Class Metrics", span=8)
    header_row = start_row + 1
    headers = [
        "Class",
        "Support",
        "TP",
        "FP",
        "FN",
        "Precision",
        "Recall",
        "F1",
    ]
    for j, h in enumerate(headers, start=start_col):
        ws.cell(header_row, j, h)
    style_header_range(ws, header_row, start_col, start_col + len(headers) - 1)

    for i, c in enumerate(class_metrics, start=1):
        r = header_row + i
        vals = [
            c["class"],
            c["support"],
            c["TP"],
            c["FP"],
            c["FN"],
            c["precision"],
            c["recall"],
            c["f1"],
        ]
        for j, value in enumerate(vals, start=start_col):
            ws.cell(r, j, value)

        for j in range(start_col + 5, start_col + 8):
            ws.cell(r, j).number_format = PERCENT_FMT

    last_row = header_row + len(class_metrics)
    style_table_cells(
        ws,
        header_row,
        last_row,
        start_col,
        start_col + len(headers) - 1,
    )
    return last_row + 2


def add_non_evaluable(ws, res, start_row):
    bad_rows = []
    for r in res["rows"]:
        if not r["row_valid"]:
            reason = r["issue"]
        elif r["pred"] is None:
            reason = "Predicted_Label blank -> nicht auswertbar"
        else:
            continue
        bad_rows.append((r, reason))

    section_title(ws, start_row, 1, f"Nicht auswertbar ({len(bad_rows)})", span=7)
    header_row = start_row + 1
    headers = [
        "Excel row",
        "Regime",
        "Ground Truth",
        "Prediction",
        "Alternative GT",
        "Reason",
        "Source file",
    ]
    for j, h in enumerate(headers, start=1):
        ws.cell(header_row, j, h)
    style_header_range(ws, header_row, 1, len(headers))

    row = header_row + 1
    for r, reason in bad_rows:
        vals = [
            r["excel_row"],
            r["raw_regime"],
            r["raw_gt"],
            r["raw_pred"],
            r["raw_alt"],
            reason,
            res["source_filename"],
        ]
        for j, value in enumerate(vals, start=1):
            ws.cell(row, j, value)
        row += 1

    if not bad_rows:
        ws.cell(row, 1, "None")
        row += 1

    style_table_cells(ws, header_row, row - 1, 1, len(headers))
    return row + 1


def write_experiment_sheet(wb, res, used_sheet_names):
    sheet_base = extract_experiment_sheet_base(res["source_filename"])
    sheet_name = make_unique_sheet_name(sheet_base, used_sheet_names)
    res["output_sheet"] = sheet_name

    ws = wb.create_sheet(sheet_name)
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    ws["A1"] = res["experiment"]
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(vertical="center")

    # 1) Experiment Overview / Counts
    row = add_overview_counts(ws, res, start_row=3, start_col=1)

    # 2) Primary Overall Metrics
    row = add_primary_overall_metrics(ws, res, start_row=row, start_col=1)

    # 3) Gesamt Weighted Metrics
    row = add_weighted_overall_metrics(ws, res, start_row=row, start_col=1)

    # 4) Overall Overview Chart (right side; does not interrupt vertical flow)
    add_overview_pie_chart(ws, res, anchor="F2")

    # 5) Gesamt Confusion Matrix
    row = add_confusion_matrix(
        ws,
        "Gesamt",
        res["scopes"]["Gesamt"]["cm"],
        start_row=row,
        start_col=1,
    )

    # 6) Gesamt Class Metrics
    row = add_class_metrics(
        ws,
        "Gesamt",
        res["scopes"]["Gesamt"]["metrics"]["class_metrics"],
        start_row=row,
        start_col=1,
    )

    # 7-9) E1 / E2 / M: metrics -> confusion matrix -> class metrics
    for scope in ("E1", "E2", "M"):
        row = add_scope_metrics(ws, res, scope, start_row=row, start_col=1)
        row = add_confusion_matrix(
            ws,
            scope,
            res["scopes"][scope]["cm"],
            start_row=row,
            start_col=1,
        )
        row = add_class_metrics(
            ws,
            scope,
            res["scopes"][scope]["metrics"]["class_metrics"],
            start_row=row,
            start_col=1,
        )

    # 10) Nicht auswertbar details
    add_non_evaluable(ws, res, row)

    auto_width(ws, 10, 28)
    ws.column_dimensions["A"].width = 30
    for col in range(2, 16):
        ws.column_dimensions[get_column_letter(col)].width = min(
            ws.column_dimensions[get_column_letter(col)].width or 12,
            22,
        )
    fit_wrapped_rows(ws)
    ws.row_dimensions[1].height = 24

    return sheet_name


def write_output(results, output_path):
    wb = Workbook()
    wb.remove(wb.active)
    used_sheet_names = {"experiment_summary", "config"}

    ws = wb.create_sheet("Experiment_Summary")
    headers = [
        "Experiment", "Sheet", "N total", "N auswertbar", "N nicht auswertbar", "Coverage",
        "Strict Overall Correct", "Strict Overall Accuracy",
        "Accepted-Set Correct", "Accepted-Set Overall Accuracy",
        "E1 N", "E1 Accuracy", "E1 Macro F1", "E1 Weighted F1",
        "E2 N", "E2 Accuracy", "E2 Macro F1", "E2 Weighted F1",
        "M N", "M Accuracy", "M Macro F1", "M Weighted F1",
        "Gesamt N", "Gesamt Accuracy", "Gesamt Macro F1", "Gesamt Weighted F1",
        "sklearn verification",
    ]
    ws.append(headers)
    style_header_range(ws, 1, 1, len(headers))
    # Keep experiment identity visible while horizontally reading the 27 metrics.
    ws.freeze_panes = "C2"
    ws.sheet_view.showGridLines = False

    for res in results:
        write_experiment_sheet(wb, res, used_sheet_names)

    for res in results:
        overall = res["overall"]
        row = [
            res["experiment"], res["output_sheet"], overall["n_total"],
            overall["n_auswertbar"], overall["n_nicht_auswertbar"], overall["coverage"],
            overall["strict_correct"], overall["strict_accuracy"],
            overall["accepted_correct"], overall["accepted_accuracy"],
        ]
        for scope in SCOPES:
            m = res["scopes"][scope]["metrics"]
            row.extend([m["n"], m["accuracy"], m["macro_f1"], m["weighted_f1"]])
        row.append(res["sklearn_status"])
        ws.append(row)
        current = ws.max_row
        ws.cell(current, 2).hyperlink = f"#'{res['output_sheet']}'!A1"
        ws.cell(current, 2).style = "Hyperlink"

    percent_cols = [6, 8, 10, 12, 13, 14, 16, 17, 18, 20, 21, 22, 24, 25, 26]
    for col in percent_cols:
        for r in range(2, ws.max_row + 1):
            ws.cell(r, col).number_format = PERCENT_FMT
    style_table_cells(ws, 1, ws.max_row, 1, len(headers))
    auto_width(ws, 10, 28)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 28
    fit_wrapped_rows(ws)
    ws.row_dimensions[1].height = max(ws.row_dimensions[1].height or 0, 30)

    ws = wb.create_sheet("Config")
    ws.append(["Setting", "Value"])
    style_header_range(ws, 1, 1, 2)
    config_rows = [
        ("Input worksheet", INPUT_SHEET_NAME),
        ("Output file", OUTPUT_FILENAME),
        ("Canonical labels", " | ".join(CANONICAL_LABELS)),
        ("Gantry mapping", "Gantry = Multi-Achs-System (Gantry)"),
        ("Umsetzeinheit mapping", "Umsetzeinheit = Kombinierte Einheit"),
        ("Blank prediction", "nicht auswertbar; excluded from all classification metrics"),
        ("E1/E2/M/Gesamt confusion matrices", "Strict primary Ground Truth"),
        ("Strict Overall Accuracy", "Primary GT only; denominator = all auswertbar rows"),
        ("Accepted-Set Overall Accuracy", "Primary GT + all listed M alternative GT labels; denominator = all auswertbar rows"),
        ("Calculation", "Custom pure-Python metrics are authoritative"),
        ("sklearn", "Mandatory independent verification for confusion matrix, accuracy, class P/R/F1, macro and weighted P/R/F1 for E1/E2/M/Gesamt"),
        ("Pie chart", "Wide 2.64:1 pie chart with values inside segments and a fixed lower-right vertical legend; colors Bosch green #278C3F, Bosch red #D51317, gray #AEB5BF"),
        ("Experiment sheets", "Order: Counts -> Primary Overall -> Gesamt Weighted -> Gesamt CM -> Gesamt Class Metrics -> E1 -> E2 -> M -> Nicht auswertbar"),
        ("Experiment sheet naming", "Extract P1/P2/P3/P4 + model name from filename; fallback to original stem; duplicate names get _2, _3, ..."),
    ]
    for setting, value in config_rows:
        ws.append([setting, value])
    style_table_cells(ws, 1, ws.max_row, 1, 2)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 100
    ws.sheet_view.showGridLines = False
    fit_wrapped_rows(ws)
    ws.row_dimensions[1].height = max(ws.row_dimensions[1].height or 0, 24)

    wb.save(output_path)


def process_directory(base_dir: Path, directory_index: int, directory_total: int) -> bool:
    output_path = base_dir / OUTPUT_FILENAME

    input_files = sorted(
        p for p in base_dir.glob("*.xlsx")
        if p.name != OUTPUT_FILENAME and not p.name.startswith("~$")
    )
    if not input_files:
        print(f"SKIP [{directory_index}/{directory_total}]: no .xlsx experiment files in {base_dir}")
        return False

    print(f"\nFolder [{directory_index}/{directory_total}]: {base_dir}")
    print(f"Found {len(input_files)} experiment file(s).")
    print(f"Only worksheet {INPUT_SHEET_NAME!r} will be read.")

    results = []
    for i, path in enumerate(input_files, 1):
        print_progress(i - 1, len(input_files), "Progress")
        print(f"\n[{i}/{len(input_files)}] {path.name}")

        rows, columns = read_experiment(path)
        overall = overall_accuracies(rows)

        scopes = {}
        checks = []
        for scope, regime in (("E1","E1"), ("E2","E2"), ("M","M"), ("Gesamt",None)):
            population, evaluated, cm, metrics = evaluate_scope(rows, regime)
            status = sklearn_cross_check(evaluated, cm, metrics, f"{path.name}/{scope}")
            checks.append(f"{scope}:{status}")
            scopes[scope] = {
                "population": population,
                "evaluated": evaluated,
                "cm": cm,
                "metrics": metrics,
                "sklearn_status": status,
            }

        # Cross-definition consistency checks for single-label multiclass evaluation.
        gesamt = scopes["Gesamt"]["metrics"]
        if overall["strict_correct"] != gesamt["correct"]:
            raise AssertionError(f"{path.name}: Strict correct count != Gesamt correct count")
        if not isclose(overall["strict_accuracy"], gesamt["accuracy"], abs_tol=1e-12):
            raise AssertionError(f"{path.name}: Strict Overall Accuracy != Gesamt Accuracy")
        if not isclose(gesamt["weighted_recall"], gesamt["accuracy"], abs_tol=1e-12):
            raise AssertionError(f"{path.name}: Weighted Recall != Accuracy")
        if overall["accepted_accuracy"] + 1e-12 < overall["strict_accuracy"]:
            raise AssertionError(f"{path.name}: Accepted-Set Accuracy < Strict Accuracy")

        results.append({
            "experiment": path.stem,
            "source_filename": path.name,
            "rows": rows,
            "columns": columns,
            "overall": overall,
            "scopes": scopes,
            "sklearn_status": "; ".join(checks),
        })

        print(
            f"    total={overall['n_total']} | "
            f"auswertbar={overall['n_auswertbar']} | "
            f"strict={overall['strict_accuracy']:.1%} | "
            f"accepted={overall['accepted_accuracy']:.1%} | "
            f"sklearn={'PASS' if all(x.endswith('PASS') for x in checks) else 'CHECK'}"
        )
        print_progress(i, len(input_files), "Progress")

    # Output is written only after all experiments pass the mandatory sklearn checks.
    write_output(results, output_path)
    print(f"DONE [{directory_index}/{directory_total}]: {output_path}")
    return True


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Evaluate every experiment .xlsx in one or more folders. "
            "Each selected folder receives its own Evaluation_Summary.xlsx."
        )
    )
    parser.add_argument(
        "input_dirs",
        nargs="*",
        metavar="INPUT_DIR",
        help=(
            "Folder containing experiment .xlsx files. Use absolute paths for "
            "folders outside the script directory."
        ),
    )
    args = parser.parse_args()

    raw_directories = args.input_dirs or [str(Path(__file__).resolve().parent)]
    directories = []
    seen_directories = set()
    for raw_directory in raw_directories:
        directory = Path(raw_directory).expanduser().resolve()
        if not directory.is_dir():
            parser.error(f"Input directory does not exist or is not a folder: {directory}")
        if directory not in seen_directories:
            directories.append(directory)
            seen_directories.add(directory)

    completed = sum(
        process_directory(directory, index, len(directories))
        for index, directory in enumerate(directories, 1)
    )
    if not completed:
        raise SystemExit("No Evaluation_Summary.xlsx was created because no input .xlsx files were found.")
    print(f"\nCOMPLETE: {completed}/{len(directories)} folder(s) processed successfully.")


if __name__ == "__main__":
    main()


# Examples (Windows PowerShell)
# Default: process .xlsx files in this script's own folder.
# python .\batch_evaluate_excel_experiments.py
#
# Process several absolute folder paths; each folder receives its own
# Evaluation_Summary.xlsx.
# python .\batch_evaluate_excel_experiments.py `
#   "C:\\Users\\<Name>\\RepoMA\\outputs\\valid_results_10_runs\\gemini-2.5-pro_pdf_P3_optimized_10_runs" `
#   "C:\\Users\\<Name>\\RepoMA\\outputs\\pdf_cross_model"

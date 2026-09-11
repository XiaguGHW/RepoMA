#!/usr/bin/env python3
"""Run and evaluate the label-configuration experiment without changing the input Excel.

The runner selects only ``prompt_engineering = yes`` rows and uses the existing
columns Ground Truth, Register and Weitere zulässige Ground Truth for evaluation.
It never writes to the input workbook.
"""
from __future__ import annotations

import argparse, json, os, re, sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SCRIPT_DIR = Path(__file__).resolve().parent
SUPPORTED = {".pdf", ".png", ".jpg", ".jpeg"}
BASE = ("Lineareinheit", "Gantry", "Greifer", "Umsetzeinheit", "Roboter", "Rotationseinheit", "Keine der verfügbaren Klassen")
MERGED = "Gantry/Kombinierte Einheit"
ALIASES = {
    "lineareinheit": "Lineareinheit", "gantry": "Gantry", "multi achs system gantry": "Gantry",
    "greifer": "Greifer", "umsetzeinheit": "Umsetzeinheit", "kombinierte einheit": "Umsetzeinheit",
    "roboter": "Roboter", "rotationseinheit": "Rotationseinheit",
    "keine der verfugbaren klassen": "Keine der verfügbaren Klassen", "keine der verfügbaren klassen": "Keine der verfügbaren Klassen",
    "gantry kombinierte einheit": MERGED,
}
CONFIGS = {
    "L1": {"ambiguity": False, "undecidable": False, "merge": False},
    "L2": {"ambiguity": True,  "undecidable": False, "merge": False},
    "L3": {"ambiguity": True,  "undecidable": True,  "merge": False},
    "L4": {"ambiguity": True,  "undecidable": True,  "merge": True},
}

def clean(v): return "" if v is None or pd.isna(v) else str(v).strip()
def canon(v):
    key = re.sub(r"[^a-z0-9äöüß]+", " ", clean(v).casefold())
    return ALIASES.get(" ".join(key.split()), "")
def mapped(v, config):
    value = canon(v)
    return MERGED if CONFIGS[config]["merge"] and value in {"Gantry", "Umsetzeinheit"} else value
def parse_label_list(v):
    text = clean(v)
    if not text: return []
    return [mapped(x, "L1") for x in re.split(r"[,;\n]", text) if mapped(x, "L1")]
def default_input():
    for candidate in (SCRIPT_DIR / "input" / "classification_experiment_dataset.xlsx", SCRIPT_DIR.parent / "input" / "classification_experiment_dataset.xlsx"):
        if candidate.is_file(): return candidate
    return SCRIPT_DIR / "input" / "classification_experiment_dataset.xlsx"
def args():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--input-excel", type=Path, default=default_input())
    p.add_argument("--sheet-name", default="Experiment_Dataset")
    p.add_argument("--prompt-file", type=Path, required=False, help="Base prompt, normally your final P3/P3_optimized prompt.")
    p.add_argument("--label-config", choices=(*CONFIGS, "all"), default="all")
    p.add_argument("--repetitions", type=int, default=1, help="Runs per L configuration. Default: 1, so L1-L4 total 56 requests for 14 BGs.")
    p.add_argument("--model", default="gemini-2.5-pro")
    p.add_argument("--temperature", type=float, default=0.0)
    p.add_argument("--max-output-tokens", type=int, default=4096)
    p.add_argument("--output-dir", type=Path, default=SCRIPT_DIR / "outputs" / "label_configuration_experiment")
    p.add_argument("--evaluate-only", action="store_true", help="Evaluate existing experiment workbooks only; no model/API calls.")
    return p.parse_args()

def schema_prompt(config):
    opt=CONFIGS[config]; labels=list(BASE)
    if opt["merge"]: labels=[MERGED if x in {"Gantry","Umsetzeinheit"} else x for x in labels if x not in {"Gantry","Umsetzeinheit"}] + [MERGED]
    ambiguity = "true or false; use true only if at least two listed functional classes are genuinely plausible" if opt["ambiguity"] else "always false"
    undecidable = "true only if the supplied files are insufficient for any defensible functional classification; then primary_label must be null" if opt["undecidable"] else "always false; primary_label must never be null"
    return f'''\n\nLABEL-CONFIGURATION EXPERIMENT: {config}\nAllowed functional labels: {json.dumps(labels, ensure_ascii=False)}\nReturn ONLY this JSON object:\n{{"primary_label": "one allowed label or null", "is_ambiguous": boolean, "alternative_labels": ["allowed labels"], "is_not_decidable": boolean, "confidence_percent": number, "reasoning": "short German explanation"}}\nRules: is_ambiguous = {ambiguity}. is_not_decidable = {undecidable}. alternative_labels may contain only allowed labels different from primary_label. Never output Mehrdeutig or Nicht entscheidbar as a functional label.\n'''

def files(folder):
    root=Path(clean(folder)); return [str(p) for p in sorted(root.rglob("*")) if p.is_file() and p.suffix.casefold() in SUPPORTED] if root.is_dir() else []
def parse_json(raw):
    raw=clean(raw); candidates=[raw]
    a,b=raw.find("{"),raw.rfind("}")
    if a>=0 and b>a: candidates.append(raw[a:b+1])
    for c in candidates:
        try:
            x=json.loads(c)
            if isinstance(x,dict): return x,"SUCCESS"
        except json.JSONDecodeError: pass
    return None,"INVALID_JSON"
def value(d,*keys):
    if not d:return pd.NA
    low={str(k).casefold():v for k,v in d.items()}
    for k in keys:
        if k.casefold() in low:return low[k.casefold()]
    return pd.NA
def question(row):
    return "Klassifiziere die beigefügte Baugruppe gemäß den Systemanweisungen.\nBenennung (E): " + clean(row.get("Benennung (E)"))+"\nBenennung (D): "+clean(row.get("Benennung (D)"))+"\nGib ausschließlich das angeforderte JSON zurück."
def connector(model,key):
    from llm_connector import LLMConnector
    return LLMConnector(model,key)

def style(ws):
    fill=PatternFill("solid",fgColor="1F4E78")
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for c in ws[1]: c.fill=fill; c.font=Font(color="FFFFFF",bold=True); c.alignment=Alignment(wrap_text=True)
    for col in ws.columns:
        letter=col[0].column_letter; ws.column_dimensions[letter].width=min(max(max(len(clean(x.value)) for x in col)+2,12),45)

def save_result(frame,path):
    frame.to_excel(path,index=False)
    wb=load_workbook(path); style(wb.active); wb.save(path)

def run_one(a, dataset, base_prompt, config, repetition):
    api=os.getenv("BOSCH_FARM_SUBSCRIPTION_KEY")
    if not api: raise EnvironmentError("BOSCH_FARM_SUBSCRIPTION_KEY is missing in .env")
    llm=connector(a.model,api); outdir=a.output_dir/config; outdir.mkdir(parents=True,exist_ok=True)
    stamp=datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    path=outdir/f"label_experiment_{config}_{a.model}_run{repetition:02d}_{stamp}.xlsx"
    result=dataset.copy()
    for col in ("Predicted_Label","Is_Ambiguous","Alternative_Labels","Is_Not_Decidable","Confidence_Percent","Reasoning","Raw_Model_Response","JSON_Parse_Status","Processing_Status","Files_Used","File_Count","Run_Model","Label_Config","Configuration_Repetition","Run_Timestamp"): result[col]=pd.NA
    prompt=base_prompt+schema_prompt(config); generation={"temperature":a.temperature,"topP":0.95,"candidateCount":1,"maxOutputTokens":a.max_output_tokens}
    for i,row in result.iterrows():
        result.loc[i,["Run_Model","Label_Config","Configuration_Repetition","Run_Timestamp"]]=[a.model,config,repetition,datetime.now().isoformat(timespec="seconds")]
        try:
            used=files(row.get("Data_Folder_Path")); result.loc[i,"Files_Used"]="\n".join(used); result.loc[i,"File_Count"]=len(used)
            if not used: result.loc[i,"Processing_Status"]="SKIPPED: no supported PDF/image files"; save_result(result,path); continue
            raw=llm.ask_about_files(file_paths=used,question=question(row),system_prompt=prompt,generation_config=generation)
            data,status=parse_json(raw); result.loc[i,"Raw_Model_Response"]=clean(raw); result.loc[i,"JSON_Parse_Status"]=status
            result.loc[i,"Predicted_Label"]=value(data,"primary_label","class_label","class")
            result.loc[i,"Is_Ambiguous"]=value(data,"is_ambiguous")
            alts=value(data,"alternative_labels","possible_classes_if_ambiguous","possible_classes")
            result.loc[i,"Alternative_Labels"]=json.dumps(alts,ensure_ascii=False) if isinstance(alts,list) else alts
            result.loc[i,"Is_Not_Decidable"]=value(data,"is_not_decidable")
            result.loc[i,"Confidence_Percent"]=value(data,"confidence_percent","confidence")
            result.loc[i,"Reasoning"]=value(data,"reasoning","begründung","begruendung")
            result.loc[i,"Processing_Status"]="SUCCESS" if data else "CHECK: invalid JSON"
        except Exception as e: result.loc[i,"Processing_Status"]=f"ERROR: {e}"
        save_result(result,path)
    return path

def truth(row, config):
    primary=mapped(row.get("Ground Truth"),config)
    allowed={primary}
    for x in parse_label_list(row.get("Weitere zulässige Ground Truth")):
        allowed.add(MERGED if CONFIGS[config]["merge"] and x in {"Gantry","Umsetzeinheit"} else x)
    return primary,allowed,clean(row.get("Register"))
def yes(v): return clean(v).casefold() in {"true","yes","1","ja"}
def config_labels(config):
    labels=set(BASE)
    return (labels-{"Gantry","Umsetzeinheit"})|{MERGED} if CONFIGS[config]["merge"] else labels
def parse_output_labels(value, config):
    raw=clean(value)
    if not raw: return []
    try:
        decoded=json.loads(raw)
        values=decoded if isinstance(decoded,list) else [decoded]
    except (json.JSONDecodeError, TypeError):
        values=re.split(r"[,;\\n]",raw)
    return [canon(v) for v in values if canon(v)]
def evaluate_frame(frame, config):
    rows=[]
    valid_labels=config_labels(config)
    for _,r in frame.iterrows():
        gt,allowed,reg=truth(r,config)
        raw_primary=canon(r.get("Predicted_Label"))
        pred=mapped(r.get("Predicted_Label"),config)
        raw_alts=parse_output_labels(r.get("Alternative_Labels"),config)
        alts=[MERGED if CONFIGS[config]["merge"] and a in {"Gantry","Umsetzeinheit"} else a for a in raw_alts]
        nd=yes(r.get("Is_Not_Decidable")); amb=yes(r.get("Is_Ambiguous"))
        gt_amb=reg=="M" and len(allowed)>1
        evaluable=clean(r.get("Processing_Status"))=="SUCCESS" and bool(gt)
        primary_schema_valid=raw_primary in valid_labels
        alternatives_schema_valid=all(a in valid_labels and a != pred for a in alts)
        if config=="L1":
            schema=primary_schema_valid and not amb and not nd and not alts
        elif config=="L2":
            schema=primary_schema_valid and not nd and alternatives_schema_valid and (bool(alts) if amb else not alts)
        else:
            schema=(nd and not raw_primary and not alts and not amb) or (not nd and primary_schema_valid and alternatives_schema_valid and (bool(alts) if amb else not alts))
        strict=evaluable and not nd and pred==gt
        accepted=evaluable and not nd and pred in allowed
        expected_alts=allowed-{pred} if accepted else set()
        alternative_correct=(set(alts)==expected_alts) if gt_amb and accepted else False
        joint=(schema and strict and not amb and not alts) if not gt_amb else (schema and accepted and amb and alternative_correct)
        rows.append({"Label_Config":config,"SAP-Nummer":clean(r.get("SAP-Nummer")),"Teamcenter":clean(r.get("Teamcenter")),"Register":reg,"Ground_Truth":gt,"Accepted_Labels":", ".join(sorted(allowed)),"Prediction":pred,"Alternative_Labels":", ".join(alts),"Expected_Alternative_Labels":", ".join(sorted(expected_alts)),"Is_Ambiguous":amb,"GT_Is_Ambiguous":gt_amb,"Is_Not_Decidable":nd,"Schema_Compliant":schema,"Alternative_Labels_Correct":alternative_correct,"Strict_Correct":strict,"Accepted_Correct":accepted,"Joint_Decision_Correct":joint,"Evaluable":evaluable,"Processing_Status":clean(r.get("Processing_Status"))})
    return pd.DataFrame(rows)
def class_metrics(cases, config):
    data=cases[(cases.Label_Config==config)&cases.Evaluable&~cases.Is_Not_Decidable&cases.Prediction.ne("")].copy()
    labels=sorted(config_labels(config)); rows=[]; n=len(data)
    for label in labels:
        tp=int(((data.Ground_Truth==label)&(data.Prediction==label)).sum()); tn=int(((data.Ground_Truth!=label)&(data.Prediction!=label)).sum())
        fp=int(((data.Ground_Truth!=label)&(data.Prediction==label)).sum()); fn=int(((data.Ground_Truth==label)&(data.Prediction!=label)).sum())
        precision=tp/(tp+fp) if tp+fp else None; recall=tp/(tp+fn) if tp+fn else None
        rows.append({"Class":label,"TP":tp,"TN":tn,"FP":fp,"FN":fn,"Accuracy_(TP+TN)/N":(tp+tn)/n if n else None,"Precision":precision,"Recall":recall,"F1":2*precision*recall/(precision+recall) if precision is not None and recall is not None and precision+recall else None})
    return pd.DataFrame(rows)

def rate(x): return x.mean() if len(x) else None
def summary(cases):
    out=[]
    for config,g in cases.groupby("Label_Config"):
        ev=g[g.Evaluable]; e=ev[ev.Register.isin(["E1","E2"])]; m=ev[ev.Register.eq("M")]
        tp=int((g.Is_Ambiguous & g.GT_Is_Ambiguous).sum()); fp=int((g.Is_Ambiguous & ~g.GT_Is_Ambiguous).sum()); fn=int((~g.Is_Ambiguous & g.GT_Is_Ambiguous).sum())
        prec=tp/(tp+fp) if tp+fp else None; rec=tp/(tp+fn) if tp+fn else None
        out.append({"Label_Config":config,"Rows":len(g),"Evaluable":len(ev),"Strict_Accuracy":rate(ev.Strict_Correct),"E1_E2_Strict_Accuracy":rate(e.Strict_Correct),"M_Accepted_Set_Accuracy":rate(m.Accepted_Correct),"Joint_Decision_Accuracy":rate(ev.Joint_Decision_Correct),"E1_E2_Joint_Decision_Accuracy":rate(e.Joint_Decision_Correct),"M_Joint_Decision_Accuracy":rate(m.Joint_Decision_Correct),"Ambiguity_Precision":prec,"Ambiguity_Recall":rec,"Ambiguity_F1":2*prec*rec/(prec+rec) if prec is not None and rec is not None and prec+rec else None,"Schema_Compliance_Rate":rate(ev.Schema_Compliant),"M_Alternative_Label_Accuracy":rate(m.Alternative_Labels_Correct),"Unexpected_E1_E2_Alternatives":int((e.Alternative_Labels.ne("")).sum()),"Not_Decidable_Count":int(g.Is_Not_Decidable.sum()),"Invalid_JSON_or_Error":int((~g.Evaluable).sum())})
    return pd.DataFrame(out)
def evaluate(a):
    books=list(a.output_dir.glob("L*/label_experiment_*.xlsx"))
    if not books: raise FileNotFoundError(f"No experiment workbooks found below {a.output_dir}")
    allcases=[]
    for p in books:
        f=pd.read_excel(p,dtype=str); config=clean(f.get("Label_Config",pd.Series([p.parent.name])).iloc[0]) or p.parent.name
        allcases.append(evaluate_frame(f,config))
    cases=pd.concat(allcases,ignore_index=True); report=a.output_dir/"Label_Configuration_Experiment_Evaluation.xlsx"
    with pd.ExcelWriter(report,engine="openpyxl") as w:
        summary(cases).to_excel(w,sheet_name="Decision_Summary",index=False)
        cases.to_excel(w,sheet_name="Per_BG_Comparison",index=False)
        cases[cases.Is_Not_Decidable].to_excel(w,sheet_name="Not_Decidable_Cases",index=False)
        cases[(cases.Ground_Truth.isin(["Gantry","Umsetzeinheit",MERGED])) | (cases.Prediction.isin(["Gantry","Umsetzeinheit",MERGED]))].to_excel(w,sheet_name="Gantry_Kombi_Analysis",index=False)
        for config in sorted(cases.Label_Config.unique()):
            class_metrics(cases,config).to_excel(w,sheet_name=f"Class_Metrics_{config}",index=False)
            data=cases[(cases.Label_Config==config)&cases.Evaluable&~cases.Is_Not_Decidable&cases.Prediction.ne("")]
            pd.crosstab(data.Ground_Truth,data.Prediction,margins=True).to_excel(w,sheet_name=f"Confusion_{config}")
        cases.groupby(["Label_Config","Register"],dropna=False).agg(Rows=("Evaluable","size"),Strict_Accuracy=("Strict_Correct","mean"),Accepted_Set_Accuracy=("Accepted_Correct","mean"),Joint_Decision_Accuracy=("Joint_Decision_Correct","mean")).reset_index().to_excel(w,sheet_name="Regime_Metrics",index=False)
    wb=load_workbook(report)
    for ws in wb.worksheets: style(ws)
    wb.save(report); return report

def main():
    a=args(); load_dotenv(SCRIPT_DIR/".env")
    if a.repetitions<1: raise ValueError("--repetitions must be at least 1")
    if a.evaluate_only: print("Evaluation:",evaluate(a).resolve()); return
    if not a.input_excel.is_file(): raise FileNotFoundError(f"Input Excel not found: {a.input_excel}")
    if not a.prompt_file or not a.prompt_file.is_file(): raise FileNotFoundError("Provide your fixed base prompt with --prompt-file, e.g. prompts/P3_original.txt")
    df=pd.read_excel(a.input_excel,sheet_name=a.sheet_name,dtype=str)
    needed={"prompt_engineering","Data_Folder_Path","Ground Truth","Register"}; missing=needed-set(df.columns)
    if missing: raise ValueError(f"Missing Excel columns: {sorted(missing)}")
    df=df[df["prompt_engineering"].fillna("").str.strip().str.casefold().eq("yes")].copy()
    if df.empty: raise ValueError("No rows marked prompt_engineering = yes")
    base=a.prompt_file.read_text(encoding="utf-8").strip(); configs=list(CONFIGS) if a.label_config=="all" else [a.label_config]
    print(f"Selected {len(df)} prompt-engineering BGs. Input Excel will not be modified.")
    for c in configs:
        for r in range(1,a.repetitions+1): print("Saved:",run_one(a,df,base,c,r).resolve())
    print("Evaluation:",evaluate(a).resolve())
if __name__=="__main__":
    try: main()
    except Exception as e: print(f"ERROR: {e}",file=sys.stderr); sys.exit(1)

# PowerShell: first test only L1 once (recommended before the full experiment).
# python .\label_configuration_experiment.py --prompt-file .\prompts\P3_original.txt --label-config L1 --repetitions 1
#
# PowerShell: run all four label configurations once each (4 × 14 = 56 requests).
# python .\label_configuration_experiment.py --prompt-file .\prompts\P3_original.txt
# PowerShell: evaluate existing experiment results only (no LLM/API calls).
# python .\label_configuration_experiment.py --evaluate-only

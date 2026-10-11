"""Local text extraction and lexical retrieval for terminal document chat."""
from __future__ import annotations

import csv
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

SUPPORTED_DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".xlsm", ".csv", ".txt", ".md"}
MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_CHUNK_CHARS = 3500
MAX_CHUNKS_PER_DOCUMENT = 160
MAX_EXCEL_ROWS_PER_SHEET = 2000


def _now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _tokens(text):
    return re.findall(r"[A-Za-z_][A-Za-z0-9_]{1,}|[\u4e00-\u9fff]{2,}|\d{2,}", text.lower())


def _chunks(label, text):
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]
    result, current = [], ""
    for paragraph in paragraphs:
        if len(paragraph) > MAX_CHUNK_CHARS:
            parts = [paragraph[i:i + MAX_CHUNK_CHARS] for i in range(0, len(paragraph), MAX_CHUNK_CHARS)]
        else:
            parts = [paragraph]
        for part in parts:
            candidate = part if not current else current + "\n\n" + part
            if current and len(candidate) > MAX_CHUNK_CHARS:
                result.append({"label": label, "text": current})
                current = part
            else:
                current = candidate
    if current:
        result.append({"label": label, "text": current})
    return result


class LocalDocumentLibrary:
    """Documents and extracted chunks stay in one local JSON file."""

    def __init__(self, path: Path):
        self.path = path
        self.data = self._load()

    def _load(self):
        if self.path.exists():
            try:
                data = json.loads(self.path.read_text(encoding="utf-8"))
                if isinstance(data.get("documents"), list):
                    return data
            except (OSError, json.JSONDecodeError):
                pass
        return {"version": 1, "documents": []}

    def _save(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps(self.data, ensure_ascii=False, indent=2), encoding="utf-8")

    def add(self, value: str):
        path = Path(value.strip().strip('"')).expanduser().resolve()
        if not path.is_file():
            raise ValueError(f"File does not exist: {path}")
        if path.suffix.lower() not in SUPPORTED_DOCUMENT_EXTENSIONS:
            allowed = ", ".join(sorted(SUPPORTED_DOCUMENT_EXTENSIONS))
            raise ValueError(f"Unsupported file type {path.suffix or '(none)'}. Supported: {allowed}")
        raw = path.read_bytes()
        if len(raw) > MAX_DOCUMENT_BYTES:
            raise ValueError(f"File is larger than {MAX_DOCUMENT_BYTES // (1024 * 1024)} MB.")
        digest = hashlib.sha256(raw).hexdigest()
        chunks = self._extract(path)
        if not chunks:
            raise ValueError("No extractable text was found in this file.")
        previous = next((item for item in self.data["documents"] if item["path"] == str(path)), None)
        document = {
            "id": previous["id"] if previous else hashlib.sha256(str(path).encode()).hexdigest()[:8],
            "name": path.name,
            "path": str(path),
            "sha256": digest,
            "bytes": len(raw),
            "indexed_at": _now(),
            "chunks": chunks[:MAX_CHUNKS_PER_DOCUMENT],
        }
        self.data["documents"] = [item for item in self.data["documents"] if item["path"] != str(path)]
        self.data["documents"].append(document)
        self._save()
        return document

    def list(self):
        return self.data["documents"]

    def remove(self, document_id: str):
        before = len(self.data["documents"])
        self.data["documents"] = [item for item in self.data["documents"] if item["id"] != document_id]
        if len(self.data["documents"]) == before:
            return False
        self._save()
        return True

    def clear(self):
        count = len(self.data["documents"])
        self.data["documents"] = []
        self._save()
        return count

    def search(self, question: str, limit: int = 6):
        words = set(_tokens(question))
        chars = set(re.findall(r"[\u4e00-\u9fff]", question.lower()))
        scored = []
        for document in self.data["documents"]:
            for index, chunk in enumerate(document.get("chunks", [])):
                text = chunk["text"].lower()
                word_score = sum(1 for word in words if word in text)
                char_score = min(8, len(chars & set(re.findall(r"[\u4e00-\u9fff]", text))))
                score = word_score * 10 + char_score
                if score:
                    scored.append((score, document, index, chunk))
        scored.sort(key=lambda item: (-item[0], item[1]["name"], item[2]))
        if not scored and self.data["documents"]:
            for document in self.data["documents"][:2]:
                if document.get("chunks"):
                    scored.append((0, document, 0, document["chunks"][0]))
        return [
            {
                "id": item[1]["id"],
                "name": item[1]["name"],
                "label": item[3]["label"],
                "text": item[3]["text"],
            }
            for item in scored[:limit]
        ]

    def _extract(self, path: Path):
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self._pdf(path)
        if suffix == ".docx":
            return self._docx(path)
        if suffix in {".xlsx", ".xlsm"}:
            return self._excel(path)
        if suffix == ".csv":
            return self._csv(path)
        return _chunks(f"{path.name}", path.read_text(encoding="utf-8", errors="replace"))

    def _pdf(self, path: Path):
        try:
            import fitz
        except ImportError as exc:
            raise ValueError("PDF support requires PyMuPDF. Run: python -m pip install -r requirements.txt") from exc
        chunks = []
        with fitz.open(path) as document:
            for number, page in enumerate(document, start=1):
                chunks.extend(_chunks(f"{path.name} — page {number}", page.get_text("text")))
        return chunks

    def _docx(self, path: Path):
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("Word support requires python-docx. Run: python -m pip install -r requirements.txt") from exc
        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for number, table in enumerate(document.tables, start=1):
            rows = [" | ".join(str(cell.text).strip().replace("\n", " ") for cell in row.cells) for row in table.rows]
            if rows:
                parts.append(f"Table {number}:\n" + "\n".join(rows))
        return _chunks(path.name, "\n\n".join(parts))

    def _excel(self, path: Path):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("Excel support requires openpyxl. Run: python -m pip install -r requirements.txt") from exc
        workbook = load_workbook(path, read_only=True, data_only=False)
        chunks = []
        try:
            for sheet in workbook.worksheets:
                rows = []
                for number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if number > MAX_EXCEL_ROWS_PER_SHEET:
                        rows.append(f"[Stopped after {MAX_EXCEL_ROWS_PER_SHEET} rows]")
                        break
                    values = [str(value).replace("\n", " ")[:500] if value is not None else "" for value in row]
                    if any(values):
                        rows.append(f"Row {number}: " + " | ".join(values))
                chunks.extend(_chunks(f"{path.name} — sheet {sheet.title}", "\n".join(rows)))
        finally:
            workbook.close()
        return chunks

    def _csv(self, path: Path):
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        rows = []
        for number, row in enumerate(csv.reader(text.splitlines()), start=1):
            if number > MAX_EXCEL_ROWS_PER_SHEET:
                rows.append(f"[Stopped after {MAX_EXCEL_ROWS_PER_SHEET} rows]")
                break
            rows.append(f"Row {number}: " + " | ".join(cell[:500] for cell in row))
        return _chunks(path.name, "\n".join(rows))

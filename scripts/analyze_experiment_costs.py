#!/usr/bin/env python3
r"""Batch-calculate token costs from result Excels using model_pricing.json.

Original workbooks are updated in place. Cost columns are inserted immediately
after the token JSON column (Z by default); no new Excel files are created.

PowerShell, run in 70BG_Experimente:
python .\scripts\analyze_experiment_costs.py --input-dir .\outputs
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import column_index_from_string

BLUE = "1F4E78"
DETAIL_HEADERS = ["Pricing_Model", "Input_Price_USD_per_1M", "Output_Price_USD_per_1M", "Input_Tokens", "Visible_Output_Tokens", "Thinking_Tokens", "Billable_Output_Tokens", "Total_Tokens", "Input_Cost_USD", "Output_Cost_USD", "Total_Cost_USD", "Cost_Parse_Status"]
SUMMARY_HEADERS = ["Source_File", "Pricing_Model", "Input_Price_USD_per_1M", "Output_Price_USD_per_1M", "Valid_Result_Rows", "Skipped_Rows", "Input_Tokens", "Visible_Output_Tokens", "Thinking_Tokens", "Billable_Output_Tokens", "Total_Tokens", "Input_Cost_USD", "Output_Cost_USD", "Total_Cost_USD", "Output_File", "Status"]

def args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Calculate all experiment costs using model_pricing.json.")
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--input-excel", type=Path)
    g.add_argument("--input-dir", type=Path)
    p.add_argument("--sheet-name", default="Sheet1")
    p.add_argument("--token-column", default="Z")
    p.add_argument(
        "--pricing-json",
        type=Path,
        default=Path(__file__).parent / "model_pricing.json",
        help="Default: model_pricing.json beside this script.",
    )
    return p.parse_args()

def style(ws: Any, cols: list[int]) -> None:
    for col in cols:
        cell = ws.cell(1, col)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=BLUE)
        a = copy(cell.alignment); a.wrap_text = True; cell.alignment = a

def usage(value: Any) -> dict[str, int]:
    if value is None or not str(value).strip(): raise ValueError("empty JSON cell")
    text = str(value).strip()
    if text.startswith('"') and text.endswith('"'):
        try: text = json.loads(text)
        except json.JSONDecodeError: pass
    try: data = json.loads(text)
    except json.JSONDecodeError as e: raise ValueError(f"invalid JSON: {e.msg}") from e
    if not isinstance(data, dict): raise ValueError("JSON root is not an object")
    def n(key: str) -> int:
        try: return max(0, int(data.get(key) or 0))
        except (TypeError, ValueError) as e: raise ValueError(f"invalid {key}") from e
    inp, vis, think, total = n("promptTokenCount"), n("candidatesTokenCount"), n("thoughtsTokenCount"), n("totalTokenCount")
    billed = vis + think
    if total and inp + billed != total and total >= inp: billed = total - inp
    return {"input":inp, "visible":vis, "thinking":think, "billed":billed, "total":total or inp+billed}

def pricing(path: Path) -> tuple[str, list[dict[str, Any]]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    entries = [e for e in data["entries"] if "input_usd_per_million" in e and "output_usd_per_million" in e]
    return str(data.get("currency", "USD")).upper(), entries

def model(file: Path, entries: list[dict[str, Any]]) -> dict[str, Any] | None:
    name = file.name.lower().replace("_", "-")
    matches = [(len(str(m)), e) for e in entries for m in e.get("match", []) if str(m).lower() in name]
    return max(matches, key=lambda x:x[0])[1] if matches else None

def columns(ws: Any, after_column: int) -> dict[str, int]:
    existing = {h: next((c for c in range(1, ws.max_column+1) if ws.cell(1,c).value == h), None) for h in DETAIL_HEADERS}
    if any(existing.values()) and not all(existing.values()):
        raise ValueError("cost columns are incomplete; remove them before rerunning")
    if all(existing.values()):
        result = {h: int(existing[h]) for h in DETAIL_HEADERS}
    else:
        ws.insert_cols(after_column + 1, amount=len(DETAIL_HEADERS))
        result = {h: after_column + 1 + i for i, h in enumerate(DETAIL_HEADERS)}
    for h, col in result.items(): ws.cell(1,col).value = h
    style(ws, list(result.values()))
    return result

def process(source: Path, ns: argparse.Namespace, entry: dict[str, Any], currency: str) -> dict[str, Any]:
    wb = load_workbook(source)
    if ns.sheet_name not in wb.sheetnames: raise ValueError(f"worksheet {ns.sheet_name!r} not found")
    ws, c, token_col = wb[ns.sheet_name], None, column_index_from_string(ns.token_column)
    c = columns(ws, token_col); ip, op = float(entry["input_usd_per_million"]), float(entry["output_usd_per_million"])
    totals = {k:0.0 for k in ["input","visible","thinking","billed","total","ic","oc","cost"]}; valid=0; skipped=[]
    for row in range(2, ws.max_row+1):
        raw = ws.cell(row,token_col).value
        try: u = usage(raw)
        except ValueError as e:
            if raw is not None and str(raw).strip(): skipped.append((row,str(e))); ws.cell(row,c["Cost_Parse_Status"]).value=f"Skipped: {e}"
            continue
        ic, oc = u["input"]*ip/1e6, u["billed"]*op/1e6
        values = {"Pricing_Model":entry["display_name"],"Input_Price_USD_per_1M":ip,"Output_Price_USD_per_1M":op,"Input_Tokens":u["input"],"Visible_Output_Tokens":u["visible"],"Thinking_Tokens":u["thinking"],"Billable_Output_Tokens":u["billed"],"Total_Tokens":u["total"],"Input_Cost_USD":ic,"Output_Cost_USD":oc,"Total_Cost_USD":ic+oc,"Cost_Parse_Status":"OK"}
        for h,v in values.items():
            cell=ws.cell(row,c[h]); cell.value=v
            if "Cost" in h: cell.number_format='$ #,##0.0000'
        for k in ["input","visible","thinking","billed","total"]: totals[k]+=u[k]
        totals["ic"]+=ic; totals["oc"]+=oc; totals["cost"]+=ic+oc; valid+=1
    if "Cost_Summary" in wb.sheetnames: del wb["Cost_Summary"]
    sm=wb.create_sheet("Cost_Summary"); sm.append(["Cost calculation","Value"])
    for a,b in [("Source file",source.name),("Matched model",entry["display_name"]),(f"Input price ({currency}/1M)",ip),(f"Output price ({currency}/1M)",op),("Valid rows",valid),("Skipped rows",len(skipped)),("Total input tokens",totals["input"]),("Total visible output tokens",totals["visible"]),("Total Thinking Tokens",totals["thinking"]),("Total billable output tokens",totals["billed"]),("Total tokens",totals["total"]),(f"Total input cost ({currency})",totals["ic"]),(f"Total output cost ({currency})",totals["oc"]),(f"Total cost ({currency})",totals["cost"]),("Created",datetime.now().strftime("%Y-%m-%d %H:%M:%S"))]: sm.append([a,b])
    style(sm,[1,2]); sm.column_dimensions["A"].width=38; sm.column_dimensions["B"].width=28
    for r in [14,15]: sm.cell(r,2).number_format='$ #,##0.0000'
    wb.save(source)
    return {"Source_File":source.name,"Pricing_Model":entry["display_name"],"Input_Price_USD_per_1M":ip,"Output_Price_USD_per_1M":op,"Valid_Result_Rows":valid,"Skipped_Rows":len(skipped),"Input_Tokens":totals["input"],"Visible_Output_Tokens":totals["visible"],"Thinking_Tokens":totals["thinking"],"Billable_Output_Tokens":totals["billed"],"Total_Tokens":totals["total"],"Input_Cost_USD":totals["ic"],"Output_Cost_USD":totals["oc"],"Total_Cost_USD":totals["cost"],"Output_File":source.name,"Status":"OK"}

def summary(out: Path, rows: list[dict[str, Any]]) -> Path:
    wb=Workbook(); ws=wb.active; ws.title="All_Experiments"; ws.append(SUMMARY_HEADERS); style(ws,list(range(1,len(SUMMARY_HEADERS)+1)))
    for data in rows: ws.append([data.get(h,"") for h in SUMMARY_HEADERS])
    for r in range(2,ws.max_row+1):
        for h in ["Input_Cost_USD","Output_Cost_USD","Total_Cost_USD"]: ws.cell(r,SUMMARY_HEADERS.index(h)+1).number_format='$ #,##0.0000'
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions; ws.column_dimensions["A"].width=60; ws.column_dimensions["B"].width=30; ws.column_dimensions["O"].width=60
    gt=wb.create_sheet("Grand_Total"); gt.append(["Metric","Value"]); style(gt,[1,2])
    for h in ["Input_Tokens","Visible_Output_Tokens","Thinking_Tokens","Billable_Output_Tokens","Total_Tokens","Input_Cost_USD","Output_Cost_USD","Total_Cost_USD"]: gt.append([h,sum(float(x.get(h,0) or 0) for x in rows if x.get("Status")=="OK")])
    gt.column_dimensions["A"].width=34; gt.column_dimensions["B"].width=22
    for r in range(2,gt.max_row+1):
        if "Cost" in str(gt.cell(r,1).value): gt.cell(r,2).number_format='$ #,##0.0000'
    target=out/"all_experiments_cost_summary.xlsx"; wb.save(target); return target

def main() -> int:
    ns=args(); ns.token_column=ns.token_column.upper()
    if not re.fullmatch(r"[A-Z]{1,3}",ns.token_column): print("--token-column must be a letter, e.g. Z",file=sys.stderr); return 2
    try: currency, entries=pricing(ns.pricing_json)
    except Exception as e: print(f"Could not load pricing: {e}",file=sys.stderr); return 2
    if ns.input_excel: files=[ns.input_excel]
    else:
        if not ns.input_dir.is_dir(): print(f"Input directory not found: {ns.input_dir}",file=sys.stderr); return 2
        files=[f for f in sorted(ns.input_dir.glob("*.xlsx")) if not f.name.startswith("~$")]
    rows=[]
    for file in files:
        entry=model(file,entries)
        if not entry: rows.append({"Source_File":file.name,"Status":"ERROR: model not found in pricing JSON"}); print(f"ERROR: no pricing for {file.name}",file=sys.stderr); continue
        try:
            result=process(file,ns,entry,currency); rows.append(result); print(f"OK: updated {file.name} ({entry['display_name']})")
        except Exception as e: rows.append({"Source_File":file.name,"Pricing_Model":entry["display_name"],"Status":f"ERROR: {e}"}); print(f"ERROR: {file.name}: {e}",file=sys.stderr)
    print(f"\nUpdated: {sum(x.get('Status') == 'OK' for x in rows)}/{len(rows)} files"); return 0 if rows and all(x.get("Status")=="OK" for x in rows) else 1

if __name__ == "__main__": raise SystemExit(main())

"""Codebook (classification rule) preparation: observability, bounded output
recovery, TXT rule files and closing every failure path (synthetic data only).

The loopback mock never contacts Bosch.  A Claude model ID makes the mock
answer in Claude's shape, so ``finishReason: MAX_TOKENS`` reaches the
connector as the real ``stop_reason: max_tokens``.
"""

import json
import os
import re
from pathlib import Path
import sys
import threading
import time
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fastapi.testclient import TestClient  # noqa: E402
import chat_workbench as w  # noqa: E402
from fixture_support import FixtureEnvironment  # noqa: E402
from workbench_i18n import UIMessage, localize_message, runtime_messages  # noqa: E402

CLAUDE = "claude-haiku-5-5"
RULE_TEXT = ("Rotationseinheit: Die Hauptfunktion ist die Drehung des Objekts um eine Achse. "
             "Größe, Übersetzung und Schwenkwinkel sind Hinweise. Regel rotation-1 gilt.")


def make_pdf(path, pages):
    import fitz
    with fitz.open() as pdf:
        for text in pages:
            page = pdf.new_page()
            page.insert_textbox(fitz.Rect(40, 40, 560, 800), text, fontsize=7)
        pdf.save(path)
    return path


def long_page(label, words=260):
    return f"{label}: " + " ".join(f"Regel{index} Rotationseinheit Drehung" for index in range(words))


class CodebookPreparation(unittest.TestCase):
    def setUp(self):
        self.f = FixtureEnvironment().__enter__()
        self.addCleanup(self.f.__exit__, None, None, None)
        self.client = TestClient(w.app, raise_server_exceptions=False)
        self.addCleanup(self.client.close)
        with w.CODEBOOK_JOB_GUARD:
            w.CODEBOOK_JOB.clear()
        self.post("/api/setup", {"workspace": str(self.f.workspace), "model": "gemini-2.5-pro", "language": "en"})

    def post(self, path, body, status=200):
        response = self.client.post(path, json=body)
        self.assertEqual(response.status_code, status, response.text)
        return response.json()

    def task(self):
        response = self.client.get("/api/codebook/task")
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["task"]

    def wait_task(self, timeout=20):
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            task = self.task()
            if task and task["state"] != "running":
                return task
            time.sleep(.05)
        self.fail("Codebook preparation did not finish")

    def keys(self, task):
        return [entry["message"]["message_key"] for entry in task["log"]]

    def english(self, payload):
        return str(localize_message(UIMessage(payload["message_key"], "zh", **{
            key: (UIMessage(value["message_key"], "zh", **value.get("params", {})) if isinstance(value, dict) else value)
            for key, value in payload.get("params", {}).items()}), "en"))

    def codebook_calls(self, start=0):
        return [call for call in self.f.calls[start:] if "transcribe exactly one" in json.dumps(call["body"])]

    def assert_closed(self):
        self.assertFalse(w.MODEL_WORK_IN_PROGRESS.is_set())
        self.assertFalse(self.client.get("/api/codebook/task").json()["running"])

    # 1. PDF: background job, visible stages, view, lock, restore after a restart
    def test_pdf_background_job_reports_stages_then_view_lock_and_restore(self):
        path = make_pdf(self.f.root / "Codebook 3 Seiten.pdf", [RULE_TEXT, "Gantry: zwei Achsen.", "Greifer: greift."])
        started = self.client.post("/api/codebook/prepare", json={"path": str(path), "model": CLAUDE,
                                                                  "request_id": "c" * 32, "background": True})
        self.assertEqual(started.status_code, 202, started.text)
        self.assertEqual(started.json()["task"]["state"], "running")
        self.assertEqual(started.json()["task"]["model"], CLAUDE)
        task = self.wait_task()
        self.assertEqual(task["state"], "succeeded", task.get("failure"))
        self.assertEqual(task["source"], {"path": str(path), "name": path.name, "type": "pdf", "bytes": path.stat().st_size})
        self.assertEqual((task["result"]["page_count"], task["result"]["unit_count"]), (3, 3))
        self.assertTrue(task["finished_at"])
        keys = self.keys(task)
        for key in ["rulejob.log.started", "rulejob.stage.extracting_pdf", "rulejob.stage.extracting_pdf_page",
                    "rulejob.stage.parsing_page", "rulejob.stage.merging", "rulejob.stage.saving", "rulejob.log.succeeded"]:
            self.assertIn(key, keys)
        parsing = [entry["message"]["params"] for entry in task["log"] if entry["message"]["message_key"] == "rulejob.stage.parsing_page"]
        self.assertEqual([(item["current"], item["total"]) for item in parsing], [(1, 3), (2, 3), (3, 3)])
        self.assertTrue(all(call["path"].endswith(":rawPredict") for call in self.f.calls))
        self.assertTrue(all(call["body"]["max_tokens"] == w.CODEBOOK_BUDGET.unit_output_tokens for call in self.codebook_calls()))
        profile = self.client.get("/api/codebook").json()
        self.assertEqual((profile["status"], profile["source_type"], profile["requested_model"]), ("draft", "pdf", CLAUDE))
        self.assertEqual(task["result"]["profile_version"], profile["version"])
        self.assertEqual(task["context_after"], w.context_token(w.ACTIVE))
        self.assertEqual(self.post("/api/codebook/lock", {})["status"], "locked")
        # A restarted program has no live job: the persisted final state is shown.
        with w.CODEBOOK_JOB_GUARD:
            w.CODEBOOK_JOB.clear()
        restored = self.client.get("/api/codebook/task").json()
        self.assertEqual(restored["task"]["state"], "succeeded")
        self.assertEqual(restored["task"]["task_id"], task["task_id"])
        self.assertEqual(restored["default_model"], CLAUDE)
        self.assert_closed()

    # 2. DOCX: succeeds and cites text blocks instead of invented pages
    def test_docx_prepares_with_text_block_references(self):
        profile = self.post("/api/codebook/prepare", {"path": str(self.f.rules)})
        self.assertEqual((profile["status"], profile["source_type"], profile["unit_kind"]), ("draft", "docx", "block"))
        rule = profile["draft"]["classes"][0]["decision_rules"][0]
        self.assertEqual((rule["pages"], rule["blocks"]), ([], [1]))
        self.assertEqual(profile["draft"]["classes"][0]["evidence_pages"], [])
        self.assertEqual(self.task()["state"], "succeeded")
        self.assertIn("rulejob.stage.extracting_docx", self.keys(self.task()))

    # 3. TXT: common encodings, upper-case suffix, local reading only
    def test_txt_rule_files_decode_common_encodings_locally(self):
        cases = [("utf-8", "utf-8", ".txt"), ("utf-8-sig", "utf-8-sig", ".TXT"),
                 ("utf-16", "utf-16", ".txt"), ("cp1252", "cp1252", ".Txt")]
        text = RULE_TEXT + "\n\nGantry: Bewegung über zwei translatorische Achsen.\n\n" + "Grenzfall: " + "Prüfung " * 40
        for encoding, expected, suffix in cases:
            with self.subTest(encoding=encoding):
                path = self.f.root / f"Regeln {encoding}{suffix}"
                path.write_bytes(text.encode(encoding))
                before = len(self.f.calls)
                profile = self.post("/api/codebook/prepare", {"path": str(path)})
                self.assertEqual((profile["status"], profile["source_type"], profile["source_encoding"]), ("draft", "txt", expected))
                self.assertEqual(profile["page_records"], [])
                self.assertEqual(profile["ocr_runtime"]["status"], "not applicable to TXT")
                calls = self.f.calls[before:]
                prompts = json.dumps([call["body"] for call in calls], ensure_ascii=False)
                self.assertIn("Größe, Übersetzung", prompts)  # decoded, not mojibake
                self.assertIn("no physical pages", prompts)
                self.assertNotRegex(prompts, r"Codebook page \d|page \d+ of")
                self.assertNotIn("inline_data", prompts)  # no page images for TXT
                task = self.task()
                self.assertEqual(task["result"]["encoding"], expected)
                self.assertIn("rulejob.stage.extracting_txt", self.keys(task))
                self.assertIn("rulejob.stage.parsing_blocks", self.keys(task))

    def test_txt_empty_binary_and_unknown_encoding_fail_clearly_without_model_calls(self):
        cases = {"empty": (b"", "empty_file"), "blank": (b"\r\n  \n", "empty_file"),
                 "binary": (b"\x00\x01\x02 PK\x03\x04" * 80, "binary"),
                 "unknown": (b"\x81\x8d\x8f\x90\x9d rules " * 30, "encoding")}
        for name, (raw, reason) in cases.items():
            with self.subTest(case=name):
                path = self.f.root / f"rules {name}.txt"
                path.write_bytes(raw)
                before = len(self.f.calls)
                response = self.client.post("/api/codebook/prepare", json={"path": str(path)})
                self.assertEqual(response.status_code, 400, response.text)
                failure = response.json()["failure"]
                self.assertEqual((failure["reason"], failure["stage"]), (reason, "extracting"))
                self.assertEqual(len(self.f.calls), before)
                self.assertEqual(w.load_codebook_profile(), {})
                self.assertNotRegex(self.english(failure["message"]), r"[一-鿿]|\{\w+\}")
                self.assert_closed()

    def test_file_type_whitelist_is_case_insensitive(self):
        for name in ["rules.rtf", "rules.doc", "rules"]:
            path = self.f.root / name
            path.write_text("x", encoding="utf-8")
            response = self.client.post("/api/codebook/prepare", json={"path": str(path)})
            self.assertEqual(response.status_code, 400)
            self.assertEqual(response.json()["message_key"], "error.rules_format")
        upper = self.f.root / "RULES.DOCX"
        upper.write_bytes(self.f.rules.read_bytes())
        self.assertEqual(self.post("/api/codebook/prepare", {"path": str(upper)})["source_type"], "docx")

    # 4. Claude max_tokens on page 3: bounded, smaller retries; never half a profile
    def test_pdf_page_three_max_tokens_recovers_by_splitting_only_that_page(self):
        pages = [RULE_TEXT, "Gantry.", long_page("Seite drei Tabelle"), "Greifer."]
        path = make_pdf(self.f.root / "Codebook_Baugruppenklassifikation_v2_3.pdf", pages)

        def page_three_too_long(response, prompt, system):
            if "Codebook page 3 of 4" in prompt and "part " not in prompt:
                response["candidates"][0]["finishReason"] = "MAX_TOKENS"
            return response
        self.f.response_transform = page_three_too_long
        profile = self.post("/api/codebook/prepare", {"path": str(path), "model": CLAUDE})
        self.assertEqual(profile["status"], "draft")
        task = self.task()
        retries = [entry for entry in task["log"] if entry["message"]["message_key"] == "rulejob.log.retry_split"]
        self.assertEqual(len(retries), 1)
        self.assertEqual(retries[0]["message"]["params"]["unit"]["params"], {"number": 3})
        page_three = [call for call in self.codebook_calls() if "Codebook page 3 of 4" in json.dumps(call["body"])]
        self.assertEqual(len(page_three), 3)  # whole page, then two halves
        self.assertTrue(all(call["body"]["max_tokens"] == 2200 for call in page_three))
        parts = json.loads(profile["page_transcriptions"][2])
        self.assertEqual([note["source"] for note in parts], ["page 3 part 1/2", "page 3 part 2/2"])

    def test_pdf_page_three_persistent_max_tokens_fails_with_page_and_attempts(self):
        pages = [RULE_TEXT, "Gantry.", long_page("Seite drei Tabelle"), "Greifer."]
        path = make_pdf(self.f.root / "Codebook_Baugruppenklassifikation_v2_3.pdf", pages)

        def always_too_long(response, prompt, system):
            if "Codebook page 3 of 4" in prompt:
                response["candidates"][0]["finishReason"] = "MAX_TOKENS"
            return response
        self.f.response_transform = always_too_long
        before = len(self.f.calls)
        response = self.client.post("/api/codebook/prepare", json={"path": str(path), "model": CLAUDE})
        self.assertEqual(response.status_code, 400, response.text)
        self.assertIn("incomplete Claude response (max_tokens)", response.text)
        failure = self.task()["failure"]
        budget = w.CODEBOOK_BUDGET
        self.assertEqual((failure["reason"], failure["stage"], failure["unit"]),
                         ("output_limit", "parsing", {"kind": "page", "number": 3}))
        self.assertEqual(failure["attempts"], budget.max_attempts_per_unit)
        self.assertEqual(self.english(failure["message"]).split(";")[0], "Page 3 is too long")
        self.assertIn("after 4 attempts", self.english(failure["message"]))
        self.assertEqual(self.english(failure["short"]), "Page 3: output limit reached")
        self.assertIn("incomplete Claude response (max_tokens)", failure["technical"])
        calls = self.f.calls[before:]
        page_three = [call for call in calls if "Codebook page 3 of 4" in json.dumps(call["body"])]
        # Depth-first and fail-fast: whole page, first half, its first quarter,
        # one escalation of that quarter.  The worst case (earlier parts
        # succeeding) is still capped by max_calls_per_unit.
        self.assertEqual(len(page_three), budget.max_attempts_per_unit)
        self.assertLessEqual(len(page_three), budget.max_calls_per_unit)
        self.assertEqual([call["body"]["max_tokens"] for call in page_three], [2200, 2200, 2200, 4400])
        self.assertEqual(sum(key == "rulejob.log.retry_split" for key in self.keys(self.task())), 2)
        self.assertFalse(any("consolidate page-by-page" in json.dumps(call["body"]) for call in calls))  # no merge
        self.assertFalse(any("Codebook page 4 of 4" in json.dumps(call["body"]) for call in calls))
        self.assertEqual(w.load_codebook_profile(), {})  # no half profile
        self.assertEqual(w.ACTIVE["state"].data["usage"]["requests"], len(calls))  # failed calls still counted
        self.assert_closed()

    def test_large_merge_switches_to_sections_instead_of_longer_responses(self):
        # Field case: every page parsed, but one merged response did not fit;
        # 8000/16000-token single responses were cut by the corporate proxy.
        path = make_pdf(self.f.root / "Codebook 6 Seiten.pdf", [RULE_TEXT] + [f"Seite {n}." for n in range(2, 7)])

        def merge_too_long(response, prompt, system):
            if "consolidate page-by-page" in system:
                response["candidates"][0]["finishReason"] = "MAX_TOKENS"
            return response
        self.f.response_transform = merge_too_long
        before = len(self.f.calls)
        profile = self.post("/api/codebook/prepare", {"path": str(path), "model": CLAUDE})
        self.assertEqual(profile["draft"]["classes"][0]["name"], "Rotationseinheit")
        calls = [json.dumps(call["body"], ensure_ascii=False) for call in self.f.calls[before:]]
        merges = [call for call in self.f.calls[before:] if "consolidate page-by-page" in json.dumps(call["body"])]
        sections = [call for call in self.f.calls[before:] if "one section per request" in json.dumps(call["body"])]
        self.assertEqual([call["body"]["max_tokens"] for call in merges], [3500])  # one try, no escalation
        self.assertEqual([re.search(r"SECTION (\w+)", json.dumps(call["body"]["messages"])).group(1) for call in sections],
                         ["outline", "class", "general", "tables"])
        self.assertTrue(all(call["body"]["max_tokens"] == 3500 for call in sections))
        self.assertLessEqual(max(call["body"]["max_tokens"] for call in self.f.calls[before:]), 4400)
        # The notes travel in the cached system prompt, not the user message.
        self.assertTrue(all("CODEBOOK NOTES" in json.dumps(call["body"]["system"]) for call in sections))
        self.assertEqual(len(calls), 6 + 1 + 4)
        keys = self.keys(self.task())
        self.assertIn("rulejob.log.merge_sectioned", keys)
        self.assertEqual((keys.count("rulejob.log.merge_outline"), keys.count("rulejob.log.merge_section")), (1, 3))
        progress = [entry["message"]["params"] for entry in self.task()["log"]
                    if entry["message"]["message_key"] == "rulejob.log.merge_section"]
        self.assertEqual([(item["current"], item["total"]) for item in progress], [(1, 3), (2, 3), (3, 3)])

    def test_section_output_limit_fails_with_section_name_and_saves_nothing(self):
        def class_too_long(response, prompt, system):
            if "consolidate page-by-page" in system or "SECTION class" in prompt:
                response["candidates"][0]["finishReason"] = "MAX_TOKENS"
            return response
        self.f.response_transform = class_too_long
        response = self.client.post("/api/codebook/prepare", json={"path": str(self.f.rules)})
        self.assertEqual(response.status_code, 400, response.text)
        failure = self.task()["failure"]
        self.assertEqual((failure["reason"], failure["stage"], failure["attempts"]), ("section_output_limit", "merging", 2))
        self.assertIn("Merging “class Rotationseinheit” produced too much output", self.english(failure["message"]))
        class_calls = [call for call in self.f.calls if "SECTION class" in json.dumps(call["body"])]
        self.assertEqual([call["body"]["generationConfig"]["maxOutputTokens"] for call in class_calls], [3500, 7000])
        self.assertIn("Be more concise", json.dumps(class_calls[1]["body"]))
        self.assertEqual(w.load_codebook_profile(), {})
        self.assert_closed()

    def test_retry_reuses_completed_pages_and_only_redoes_the_merge(self):
        path = make_pdf(self.f.root / "Codebook Wiederholung.pdf", [RULE_TEXT, "Gantry.", "Greifer.", "Grenzfall."])

        def merge_fails(response, prompt, system):
            if "consolidate page-by-page" in system:
                response["candidates"][0]["finishReason"] = "SAFETY"
            return response
        self.f.response_transform = merge_fails
        self.assertEqual(self.client.post("/api/codebook/prepare", json={"path": str(path)}).status_code, 400)
        self.assertEqual(len(self.codebook_calls()), 4)
        self.f.response_transform = None
        before = len(self.f.calls)
        profile = self.post("/api/codebook/prepare", {"path": str(path)})
        self.assertEqual((profile["status"], profile["reused_units"]), ("draft", 4))
        self.assertEqual(len(self.codebook_calls(before)), 0)  # no page is paid for twice
        self.assertEqual(self.keys(self.task()).count("rulejob.log.unit_reused"), 4)
        # Another model or changed content never reuses notes.
        before = len(self.f.calls)
        self.post("/api/codebook/prepare", {"path": str(path), "model": "gemini-2.5-flash"})
        self.assertEqual(len(self.codebook_calls(before)), 4)

    def test_retry_reuses_completed_merge_sections(self):
        def merge_cut_and_tables_fail(response, prompt, system):
            if "consolidate page-by-page" in system:
                response["candidates"][0]["finishReason"] = "MAX_TOKENS"
            if "SECTION tables" in prompt and merge_cut_and_tables_fail.fail:
                response["candidates"][0]["finishReason"] = "SAFETY"
            return response
        merge_cut_and_tables_fail.fail = True
        self.f.response_transform = merge_cut_and_tables_fail
        self.assertEqual(self.client.post("/api/codebook/prepare", json={"path": str(self.f.rules)}).status_code, 400)
        merge_cut_and_tables_fail.fail = False
        before = len(self.f.calls)
        self.assertEqual(self.post("/api/codebook/prepare", {"path": str(self.f.rules)})["status"], "draft")
        sections = [re.search(r"SECTION (\w+)", json.dumps(call["body"])).group(1)
                    for call in self.f.calls[before:] if "one section per request" in json.dumps(call["body"])]
        self.assertEqual(sections, ["tables"])  # outline, class and general rules are reused
        self.assertEqual(self.keys(self.task()).count("rulejob.log.section_reused"), 3)
        self.assertEqual(len(self.codebook_calls(before)), 0)  # pages reused as well

    def test_connection_setup_timeout_is_reported_as_unsent_and_retryable(self):
        reason, _params = w.classify_model_failure(
            "Error after 3 attempts: ConnectTimeout (connection or proxy tunnel setup took longer than 15 s).")
        self.assertEqual(reason, "connect_timeout")
        full, short = w.CodebookPrepareError("connect_timeout", "merging").messages()
        self.assertIn("the request was not sent and costs nothing", str(localize_message(full, "en")))
        self.assertIn("WORKBENCH_CONNECT_TIMEOUT_SECONDS", str(localize_message(full, "de")))
        self.assertEqual(str(localize_message(short, "zh")), "合并并校验规则：连接超时（网络或代理）")
        # A genuine read timeout still means "maybe generating": not this reason.
        self.assertEqual(w.classify_model_failure("Error: no complete response within 300 s (ReadTimeout)")[0], "timeout")

    def test_proxy_interruption_is_reported_as_such(self):
        reason, _params = w.classify_model_failure("Error after 3 attempts: ProxyError")
        self.assertEqual(reason, "proxy")
        full, short = w.CodebookPrepareError("proxy", "merging").messages()
        self.assertIn("network proxy interrupted", str(localize_message(full, "en")))
        self.assertEqual(str(localize_message(short, "zh")), "合并并校验规则：被网络代理中断")

    def test_total_call_limit_stops_a_run(self):
        from docx import Document
        document = Document()
        for index in range(9):
            document.add_paragraph(f"Regelabschnitt {index}: Rotationseinheit. " + "Detail " * 520)
        path = self.f.root / "Viele Abschnitte.docx"
        document.save(path)
        with patch.object(w, "CODEBOOK_BUDGET", w.CodebookParseBudget(max_total_calls=2)):
            before = len(self.f.calls)
            response = self.client.post("/api/codebook/prepare", json={"path": str(path)})
        self.assertEqual(response.status_code, 400, response.text)
        self.assertEqual(response.json()["failure"]["reason"], "call_limit")
        self.assertEqual(len(self.f.calls) - before, 2)
        self.assert_closed()

    def test_budget_is_explicit_and_separate_from_chat(self):
        budget = w.CodebookParseBudget()
        self.assertEqual((budget.unit_output_tokens, budget.escalated_output_tokens, budget.max_split_depth), (2200, 4400, 2))
        self.assertEqual((budget.max_attempts_per_unit, budget.max_calls_per_unit), (4, 11))
        self.assertEqual(budget.call_limit(3), 3 * 11 + 1 + 2 * (40 + 3) + 1)
        self.assertEqual((budget.merge_output_tokens, budget.section_output_tokens), ((3500,), (3500, 7000)))
        self.assertEqual(budget.call_limit(1000), budget.max_total_calls)
        self.assertLessEqual(max(budget.merge_output_tokens), 32000)  # connector ceiling
        self.post("/api/codebook/prepare", {"path": str(self.f.rules)})
        self.post("/api/codebook/lock", {})
        self.post("/api/target", {"target": str(self.f.bg_a), "selected_paths": ["facts.txt"]})
        before = len(self.f.calls)
        self.post("/api/chat", {"question": "Classify this Baugruppe"})
        configs = [call["body"].get("generationConfig", {}) for call in self.f.calls[before:]]
        self.assertFalse(any(config.get("maxOutputTokens") in {2200, 4400} for config in configs), configs)

    # 5. A failed new preparation keeps the previously locked rules usable
    def test_failed_new_rules_keep_old_locked_rules_in_effect(self):
        self.post("/api/codebook/prepare", {"path": str(self.f.rules)})
        locked = self.post("/api/codebook/lock", {})
        path = make_pdf(self.f.root / "Neue Regeln.pdf", [RULE_TEXT, "Gantry.", long_page("Seite drei")])
        self.f.response_transform = lambda response, prompt, system: (
            response["candidates"][0].update(finishReason="MAX_TOKENS") or response
            if "Codebook page 3 of 3" in prompt else response)
        response = self.client.post("/api/codebook/prepare", json={"path": str(path), "model": CLAUDE})
        self.assertEqual(response.status_code, 400)
        self.f.response_transform = None
        task = self.task()
        self.assertEqual(task["state"], "failed")
        self.assertEqual(task["previous_profile"]["status"], "locked")
        self.assertEqual(task["previous_profile"]["source_name"], self.f.rules.name)
        current = self.client.get("/api/codebook").json()
        self.assertEqual((current["status"], current["version"], current["source_path"]),
                         ("locked", locked["version"], str(self.f.rules)))
        self.post("/api/target", {"target": str(self.f.bg_a), "selected_paths": ["facts.txt"]})
        result = self.post("/api/chat", {"question": "Classify this Baugruppe"})
        self.assertEqual(result["classification"]["label"], "Rotationseinheit")

    # 6. Network, timeout, invalid JSON and cancellation all close the job
    def test_network_failure_closes_job(self):
        with patch.dict(os.environ, {"BOSCH_FARM_BASE_URL": "http://127.0.0.1:9/api", "WORKBENCH_MAX_ATTEMPTS": "1",
                                     "WORKBENCH_RETRY_BACKOFF_SECONDS": "0"}):
            response = self.client.post("/api/codebook/prepare", json={"path": str(self.f.rules)})
        self.assertEqual(response.status_code, 400, response.text)
        failure = self.task()["failure"]
        self.assertEqual((failure["reason"], failure["unit"]), ("network", {"kind": "blocks", "first": 1, "last": 1}))
        self.assertIn("Bosch Farm could not be reached", self.english(failure["message"]))
        self.assert_closed()
        self.assertEqual(self.post("/api/codebook/prepare", {"path": str(self.f.rules)})["status"], "draft")

    def test_model_timeout_closes_job(self):
        def slow(response, prompt, system):
            time.sleep(.6)
            return response
        self.f.response_transform = slow
        with patch.dict(os.environ, {"WORKBENCH_RESPONSE_TIMEOUT_SECONDS": ".2", "WORKBENCH_RETRY_BACKOFF_SECONDS": "0"}):
            response = self.client.post("/api/codebook/prepare", json={"path": str(self.f.rules)})
        self.assertEqual(response.status_code, 400, response.text)
        self.assertEqual(self.task()["failure"]["reason"], "timeout")
        self.assertEqual(w.load_codebook_profile(), {})
        time.sleep(.7)
        self.assert_closed()

    def test_complete_but_malformed_page_note_is_merged_as_text_and_flagged(self):
        # Field case: page 6 returned a complete note with a JSON syntax error.
        path = make_pdf(self.f.root / "Codebook 6 Seiten.pdf", [RULE_TEXT] + [f"Seite {n}." for n in range(2, 7)])
        broken = '{"source":"page 6","decision_rules":["Greifer auf einer Hubachse: R = 0, T = 1 "Klasse 1""]}'

        def page_six_malformed(response, prompt, system):
            if "transcribe exactly one" in system and "Codebook page 6 of 6" in prompt:
                response["candidates"][0]["content"]["parts"][0]["text"] = broken
            return response
        self.f.response_transform = page_six_malformed
        before = len(self.f.calls)
        profile = self.post("/api/codebook/prepare", {"path": str(path), "model": CLAUDE})
        self.assertEqual(profile["status"], "draft")
        self.assertEqual(profile["unparsed_units"], [{"kind": "page", "number": 6}])
        self.assertEqual(len(self.f.calls) - before, 6 + 1)  # no repeated page call, no repair call
        merge = next(json.dumps(call["body"], ensure_ascii=False) for call in self.f.calls[before:]
                     if "consolidate page-by-page" in json.dumps(call["body"]))
        self.assertIn("unparsed_note", merge)
        self.assertIn("Greifer auf einer Hubachse", merge)
        task = self.task()
        self.assertIn("rulejob.log.note_unparsed", self.keys(task))
        self.assertIn("Expecting", task["unparsed_units"][0]["detail"])  # parser position for diagnostics
        self.assertIn("Klasse 1", task["unparsed_units"][0]["detail"])
        # A retry reuses the note and keeps the review flag.
        self.f.response_transform = None
        before = len(self.f.calls)
        again = self.post("/api/codebook/prepare", {"path": str(path), "model": CLAUDE})
        self.assertEqual((again["reused_units"], again["unparsed_units"]), (6, [{"kind": "page", "number": 6}]))
        self.assertEqual(len(self.codebook_calls(before)), 0)
        status = [entry["message"]["params"] for entry in self.task()["log"]
                  if entry["message"]["message_key"] == "rulejob.log.cache_status"]
        self.assertEqual((status[0]["reusable"], status[0]["total"]), (6, 6))

    def test_page_note_parsing_tolerates_common_wrappers_but_not_empty_text(self):
        self.assertEqual(w.parse_codebook_note('```json\n{"a": [1, 2,],}\n```'), {"a": [1, 2]})
        self.assertEqual(w.parse_codebook_note('{"a": 1} Hinweis: weitere Regeln folgen.'), {"a": 1})
        self.assertIsNone(w.parse_codebook_note("Sorry, prose."))

    def test_invalid_json_from_the_merge_is_still_rejected(self):
        def bad_merge(response, prompt, system):
            if "consolidate page-by-page" in system or "repair its JSON syntax" in prompt:
                response["candidates"][0]["content"]["parts"][0]["text"] = '{"classes": [{"name": '
            return response
        self.f.response_transform = bad_merge
        response = self.client.post("/api/codebook/prepare", json={"path": str(self.f.rules)})
        self.assertEqual(response.status_code, 400)
        failure = self.task()["failure"]
        self.assertEqual((failure["reason"], failure["stage"]), ("invalid_json", "merging"))
        self.assertEqual(w.load_codebook_profile(), {})
        self.assert_closed()

    def test_user_cancel_of_background_job_is_recorded_and_saves_nothing(self):
        from docx import Document
        document = Document()
        for index in range(12):
            document.add_paragraph(f"Regelabschnitt {index}: " + "Detail " * 520)
        path = self.f.root / "Abbruch.docx"
        document.save(path)
        gate = threading.Event()
        self.addCleanup(gate.set)

        def slow(response, prompt, system):
            gate.wait(.3)
            return response
        self.f.response_transform = slow
        request_id = "d" * 32
        started = self.client.post("/api/codebook/prepare", json={"path": str(path), "request_id": request_id, "background": True})
        self.assertEqual(started.status_code, 202)
        busy = self.client.post("/api/codebook/prepare", json={"path": str(self.f.rules), "background": True})
        self.assertEqual(busy.status_code, 409)  # no second job while one runs
        self.assertEqual(self.client.delete("/api/codebook").status_code, 409)  # no file change meanwhile
        deadline = time.monotonic() + 5
        while not self.codebook_calls() and time.monotonic() < deadline:
            time.sleep(.02)
        self.assertTrue(self.post("/api/operation/cancel", {"request_id": request_id})["cancelled"])
        task = self.wait_task()
        self.assertEqual((task["state"], task["failure"]["reason"]), ("cancelled", "cancelled"))
        self.assertEqual(w.load_codebook_profile(), {})
        self.assertLessEqual(len(self.codebook_calls()), 2)
        self.assert_closed()
        self.f.response_transform = None
        self.assertEqual(self.post("/api/codebook/prepare", {"path": str(self.f.rules)})["status"], "draft")

    def test_interrupted_running_record_is_closed_after_restart(self):
        record = {"task_id": "x", "state": "running", "stage": "parsing", "source": {"name": "a.pdf"}, "log": [],
                  "updated_at": "2026-10-25T10:00:00+01:00"}
        w.write_json_atomically(w.codebook_job_path(), record)
        task = self.task()
        self.assertEqual((task["state"], task["failure"]["reason"]), ("failed", "interrupted"))
        self.assertEqual(json.loads(w.codebook_job_path().read_text(encoding="utf-8"))["state"], "failed")

    def test_invalid_model_is_rejected_before_any_job(self):
        for model in ["not a model", "unknown-family-1", "claude?x"]:
            response = self.client.post("/api/codebook/prepare", json={"path": str(self.f.rules), "model": model,
                                                                       "background": True})
            self.assertEqual(response.status_code, 400, response.text)
            self.assertEqual(response.json()["message_key"], "rulejob.model_invalid")
        self.assertIsNone(self.task())
        self.assertEqual(self.f.calls, [])
        self.assert_closed()

    # 10. Every new message is complete in all three languages
    def test_every_failure_message_is_complete_in_three_languages(self):
        messages = runtime_messages()
        for reason in sorted(w.CODEBOOK_FAILURE_REASONS):
            for unit in [None, {"kind": "page", "number": 3}, {"kind": "blocks", "first": 4, "last": 6}]:
                error = w.CodebookPrepareError(reason, "parsing", unit=unit, attempts=4,
                                               params={"status": "503", "limit": 9, "detail": UIMessage("rules.classes_empty"),
                                                       "section": UIMessage("rulejob.section.class", name="Gantry")})
                for message in error.messages():
                    for language in ["zh", "en", "de"]:
                        text = str(localize_message(message, language))
                        self.assertNotRegex(text, r"\{\w+\}", (reason, language, text))
                        if language != "zh":
                            self.assertNotRegex(text, r"[一-鿿]", (reason, language, text))
        for key in messages["zh"]:
            if key.startswith(("rulejob.", "answer.", "confirm.")):
                self.assertTrue(all(messages[language][key].strip() for language in ["en", "de"]), key)

    def test_frontend_has_no_native_model_prompt_or_dead_mode_dropdown(self):
        source = w.HTML
        prepare = source[source.index("async function prepareCodebook()"):source.index("function retryCodebookPreparation()")]
        self.assertNotIn("prompt(", prepare)
        self.assertNotIn("alert(", prepare)
        for name in ["lockCodebook", "changeCodebook"]:
            body = source[source.index(f"async function {name}()"):]
            body = body[:body.index("\n")]
            self.assertNotIn("confirm(t(", body.replace("appConfirm(t(", ""))
        self.assertNotIn("createElement('select');mode.id='answermode'", source)
        self.assertIn('id="answermodegroup"', source)
        self.assertNotRegex(source, r"rulews\.model_prompt")


if __name__ == "__main__":
    unittest.main(verbosity=2)

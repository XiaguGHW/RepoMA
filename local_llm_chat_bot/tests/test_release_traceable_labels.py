"""F-07 traceable evidence labels and F-11 ignored runtime state (synthetic only)."""

import json
from pathlib import Path
import shutil
import subprocess
import sys
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fastapi.testclient import TestClient
import chat_workbench as w
from fixture_support import FixtureEnvironment
from local_document_library import LABEL_VERSION, LocalDocumentLibrary

REPOSITORY = Path(__file__).resolve().parents[2]


class Extraction(unittest.TestCase):
    def setUp(self):
        self.f = FixtureEnvironment().__enter__()
        self.addCleanup(self.f.__exit__, None, None, None)
        self.library = LocalDocumentLibrary(self.f.root / "index" / "documents.json")

    def test_nonpaged_chunks_get_unique_section_labels(self):
        text = self.f.root / "long notes.txt"
        text.write_text("\n\n".join(f"Paragraph {i}: " + "detail " * 400 for i in range(6)), encoding="utf-8")
        labels = [chunk["label"] for chunk in self.library.add(str(text))["chunks"]]
        self.assertGreater(len(labels), 1)
        self.assertEqual(len(labels), len(set(labels)))
        self.assertEqual(labels[0], f"long notes.txt — section 1/{len(labels)}")
        from docx import Document
        document = Document()
        for index in range(6):
            document.add_paragraph(f"Section {index}: " + "regel " * 400)
        path = self.f.root / "Rules long.docx"
        document.save(path)
        labels = [chunk["label"] for chunk in self.library.add(str(path))["chunks"]]
        self.assertTrue(all(label.startswith("Rules long.docx — section ") for label in labels))
        self.assertEqual(len(labels), len(set(labels)))

    def test_single_chunk_keeps_its_plain_name_and_csv_uses_row_ranges(self):
        self.assertEqual(self.library.add(str(self.f.bg_a / "facts.txt"))["chunks"][0]["label"], "facts.txt")
        csv = self.f.root / "parts.csv"
        csv.write_text("id,name\n" + "\n".join(f"{i},part {i} " + "x" * 120 for i in range(1, 80)), encoding="utf-8")
        labels = [chunk["label"] for chunk in self.library.add(str(csv))["chunks"]]
        self.assertGreater(len(labels), 1)
        self.assertEqual(labels[0].split(" — ")[0], "parts.csv")
        self.assertRegex(labels[0], r"^parts\.csv — rows 1-\d+$")
        self.assertEqual(len(labels), len(set(labels)))

    def test_old_index_records_are_relabelled_without_changing_ids(self):
        path = self.f.root / "legacy.txt"
        path.write_text("\n\n".join("Legacy " + "word " * 500 for _ in range(4)), encoding="utf-8")
        record = self.library.add(str(path))
        record.pop("label_version")
        for chunk in record["chunks"]:
            chunk["label"] = "legacy.txt"  # the previous file-level format
        self.library._save()
        reloaded = LocalDocumentLibrary(self.library.path)
        upgraded = reloaded.add(str(path))
        self.assertEqual(upgraded["id"], record["id"])
        self.assertEqual(upgraded["label_version"], LABEL_VERSION)
        self.assertTrue(upgraded["chunks"][0]["label"].startswith("legacy.txt — section 1/"))


class Workbench(unittest.TestCase):
    def setUp(self):
        self.f = FixtureEnvironment().__enter__()
        self.addCleanup(self.f.__exit__, None, None, None)
        self.client = TestClient(w.app, raise_server_exceptions=False)
        self.addCleanup(self.client.close)
        self.post("/api/setup", {"workspace": str(self.f.workspace), "model": "gemini-2.5-pro", "language": "en"})
        self.post("/api/codebook/prepare", {"path": str(self.f.rules)})
        self.post("/api/codebook/lock", {})

    def post(self, path, body):
        response = self.client.post(path, json=body)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def test_same_named_files_in_subfolders_get_distinct_validated_citations(self):
        for folder, text in (("sub1", "FIXTURE motor rotates the table."), ("sub2", "FIXTURE gripper holds the part.")):
            (self.f.bg_a / folder).mkdir()
            (self.f.bg_a / folder / "notes.txt").write_text(text + "\n\n" + ("filler " * 700), encoding="utf-8")
        self.post("/api/target", {"target": str(self.f.bg_a), "selected_paths": ["facts.txt", "sub1/notes.txt", "sub2/notes.txt"]})
        before = len(self.f.calls)
        result = self.post("/api/chat", {"question": "Classify this Baugruppe"})
        labels = [item["label"] for item in result["sources"]]
        self.assertEqual(len(labels), len(set(labels)), labels)
        self.assertTrue(any(label.startswith("sub1/notes.txt — section ") for label in labels), labels)
        self.assertTrue(any(label.startswith("sub2/notes.txt — section ") for label in labels), labels)
        self.assertIn("facts.txt", labels)  # a file at the top level keeps its plain name
        cited = result["classification"]["bg_evidence_pages"]
        self.assertTrue(set(cited) <= set(labels))
        # Pre-reading summaries name the files the same way the evidence does.
        prompts = json.dumps(self.f.calls[before:], ensure_ascii=False)
        self.assertIn("--- sub1/notes.txt — section 1/", prompts)
        self.assertIn("--- sub2/notes.txt — section 1/", prompts)

    def test_nonpaged_codebook_references_name_their_section(self):
        from docx import Document
        document = Document()
        for index in range(6):
            document.add_paragraph(f"Rotationseinheit rule section {index}: rotation-1 applies. " + "regel " * 400)
        rules = self.f.root / "Sectioned rules.docx"
        document.save(rules)
        self.post("/api/codebook/prepare", {"path": str(rules)})
        self.post("/api/codebook/lock", {})
        self.post("/api/target", {"target": str(self.f.bg_a), "selected_paths": ["facts.txt"]})
        result = self.post("/api/chat", {"question": "Classify this Baugruppe rotation"})
        refs = result["classification"]["codebook_rule_refs"]
        self.assertTrue(refs and all(ref.startswith("Sectioned rules.docx — section ") for ref in refs), refs)


@unittest.skipUnless(shutil.which("git") and (REPOSITORY / ".git").exists(), "requires the git checkout")
class IgnoredRuntimeFiles(unittest.TestCase):
    def ignored(self, path):
        return subprocess.run(["git", "check-ignore", "-q", path], cwd=REPOSITORY).returncode == 0

    def test_state_caches_secrets_and_outputs_are_ignored(self):
        for path in ["local_llm_chat_bot/AI_Bot_State/workbench_settings.json",
                     "local_llm_chat_bot/AI_Bot_State/projects/abc/default/transcript.jsonl",
                     "local_llm_chat_bot/.bosch_chat/session.json",
                     "local_llm_chat_bot/.excel_chat_agent/session.json",
                     "local_llm_chat_bot/.env", "local_llm_chat_bot/Farm ü.env", "local_llm_chat_bot/.env.local",
                     "Results/workshop_ai_updated_20261025_120000.xlsx"]:
            with self.subTest(path=path):
                self.assertTrue(self.ignored(path))

    def test_program_files_and_templates_are_not_ignored(self):
        for path in ["local_llm_chat_bot/chat_workbench.py", "local_llm_chat_bot/workbench_i18n.json",
                     "local_llm_chat_bot/requirements.txt", ".env.example"]:
            with self.subTest(path=path):
                self.assertFalse(self.ignored(path))


if __name__ == "__main__":
    unittest.main(verbosity=2)

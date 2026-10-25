"""F-02: realistic request bounds and Stop for long operations (synthetic loopback only)."""

import os
from pathlib import Path
import sys
import threading
import time
from types import SimpleNamespace
import unittest
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fastapi.testclient import TestClient
import chat_workbench as w
from fixture_support import FixtureEnvironment
from llm_connector_with_prompt_caching import LLMConnector
from request_runtime import RequestControl, RequestPolicy


class PolicyDefaults(unittest.TestCase):
    def test_defaults_separate_complete_response_from_stream_gap(self):
        with patch.dict(os.environ, {}, clear=False):
            for name in ("REQUEST_TIMEOUT_SECONDS", "READ_TIMEOUT_SECONDS", "RESPONSE_TIMEOUT_SECONDS"):
                os.environ.pop("WORKBENCH_" + name, None)
            policy = RequestPolicy.from_env()
        self.assertEqual((policy.timeout, policy.read_timeout, policy.response_timeout), (3600, 60, 300))
        control = RequestControl(policy)
        self.assertEqual(control.timeouts()[1], 60)
        self.assertEqual(control.timeouts(complete=True)[1], 300)

    def test_complete_wait_is_still_capped_by_the_operation_deadline(self):
        control = RequestControl(RequestPolicy(timeout=.5, response_timeout=300))
        self.assertLessEqual(control.timeouts(complete=True)[1], .5)

    def test_environment_bounds(self):
        with patch.dict(os.environ, {"WORKBENCH_RESPONSE_TIMEOUT_SECONDS": "900",
                                     "WORKBENCH_REQUEST_TIMEOUT_SECONDS": "7200"}):
            policy = RequestPolicy.from_env()
        self.assertEqual((policy.response_timeout, policy.timeout), (900, 7200))
        for name, value in (("RESPONSE_TIMEOUT_SECONDS", "0"), ("RESPONSE_TIMEOUT_SECONDS", "3601"),
                            ("RESPONSE_TIMEOUT_SECONDS", "nan"), ("REQUEST_TIMEOUT_SECONDS", "14401")):
            with self.subTest(name=name, value=value), patch.dict(os.environ, {"WORKBENCH_" + name: value}):
                with self.assertRaises(ValueError):
                    RequestPolicy.from_env()

    def test_openai_full_response_uses_response_timeout_and_stream_uses_gap(self):
        for stream, expected in ((False, 300), (True, 60)):
            with self.subTest(stream=stream):
                captured = {}
                client = MagicMock()
                client._client.event_hooks = {}

                def create(**kwargs):
                    captured.update(kwargs)
                    raise ValueError("fixture stops here")
                client.chat.completions.create.side_effect = create
                connector = LLMConnector("gpt-4.1", "fixture-only")
                connector.request_control = RequestControl(RequestPolicy())
                with patch("llm_connector_with_prompt_caching.OpenAI", return_value=client):
                    self.assertTrue(connector.ask_about_files([], "test", stream=stream).startswith("Error:"))
                self.assertEqual(captured["timeout"].read, expected)


class ApiBounds(unittest.TestCase):
    def setUp(self):
        self.f = FixtureEnvironment().__enter__()
        self.addCleanup(self.f.__exit__, None, None, None)
        self.client = TestClient(w.app, raise_server_exceptions=False)
        self.addCleanup(self.client.close)
        self.post("/api/setup", {"workspace": str(self.f.workspace), "model": "gemini-2.5-pro", "language": "en"})

    def post(self, path, body, client=None):
        response = (client or self.client).post(path, json=body)
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def delay(self, seconds):
        def transform(response, prompt, system):
            time.sleep(seconds)
            return response
        self.f.response_transform = transform

    def test_slow_but_healthy_complete_response_succeeds_once(self):
        """Slower than the stream gap, within the complete-response wait: one dispatch, success."""
        self.delay(.6)
        before = len(self.f.calls)
        with patch.dict(os.environ, {"WORKBENCH_READ_TIMEOUT_SECONDS": ".2",
                                     "WORKBENCH_RESPONSE_TIMEOUT_SECONDS": "5"}):
            result = self.post("/api/chat", {"question": "Explain the workspace"})
        self.assertTrue(result["answer"])
        self.assertEqual(len(self.f.calls) - before, 1)
        self.assertEqual(w.ACTIVE["state"].data["usage"]["requests"], 1)

    def test_read_timeout_is_reported_and_never_redispatched(self):
        self.delay(1.0)
        before = len(self.f.calls)
        with patch.dict(os.environ, {"WORKBENCH_RESPONSE_TIMEOUT_SECONDS": ".3",
                                     "WORKBENCH_RETRY_BACKOFF_SECONDS": "0"}):
            response = self.client.post("/api/chat", json={"question": "Explain the workspace"})
        self.assertEqual(response.status_code, 400, response.text)
        self.assertIn("WORKBENCH_RESPONSE_TIMEOUT_SECONDS", response.text)
        time.sleep(1.2)  # let the fixture finish the abandoned generation
        self.assertEqual(len(self.f.calls) - before, 1)
        self.assertEqual(w.ACTIVE["state"].data["usage"]["requests"], 1)
        self.assertEqual(w.ACTIVE["state"].data["history"], [])

    def test_many_section_codebook_and_multi_file_first_question_fit_defaults(self):
        from docx import Document
        rules = self.f.root / "Long rules.docx"
        document = Document()
        for index in range(24):
            document.add_paragraph(f"Rule section {index}: Rotationseinheit rotates. " + "detail " * 450)
        document.save(rules)
        self.delay(.05)
        before = len(self.f.calls)
        profile = self.post("/api/codebook/prepare", {"path": str(rules)})
        self.assertEqual(profile["status"], "draft")
        self.assertGreaterEqual(len(self.f.calls) - before, 9)  # >= 8 section calls + merge
        self.post("/api/codebook/lock", {})
        for index in range(6):
            (self.f.bg_a / f"part_{index}.txt").write_text(f"FIXTURE part {index} rotates.", encoding="utf-8")
        self.post("/api/target", {"target": str(self.f.bg_a), "selected_paths": ["."]})
        result = self.post("/api/chat", {"question": "Classify this Baugruppe"})
        self.assertEqual(result["classification"]["label"], "Rotationseinheit")

    def run_and_stop(self, path, body):
        """Start a long operation, stop it after the first provider dispatch, return (status, calls)."""
        request_id = "a" * 32
        body = dict(body, request_id=request_id)
        self.delay(.4)
        before = len(self.f.calls)
        outcome = {}
        worker_client = TestClient(w.app, raise_server_exceptions=False)
        self.addCleanup(worker_client.close)

        def run():
            outcome["response"] = worker_client.post(path, json=body)
        thread = threading.Thread(target=run)
        thread.start()
        deadline = time.monotonic() + 5
        while len(self.f.calls) == before and time.monotonic() < deadline:
            time.sleep(.01)
        self.assertGreater(len(self.f.calls), before, "operation never dispatched")
        stopped = self.post("/api/operation/cancel", {"request_id": request_id})
        self.assertTrue(stopped["cancelled"])
        thread.join(10)
        self.assertFalse(thread.is_alive())
        time.sleep(.5)
        self.assertFalse(w.MODEL_WORK_IN_PROGRESS.is_set())
        self.assertEqual(self.client.post("/api/operation/cancel", json={"request_id": request_id}).status_code, 404)
        return outcome["response"], len(self.f.calls) - before

    def test_stop_codebook_preparation_saves_no_profile(self):
        from docx import Document
        rules = self.f.root / "Stop rules.docx"
        document = Document()
        for index in range(12):
            document.add_paragraph(f"Rule section {index}: " + "detail " * 450)
        document.save(rules)
        response, calls = self.run_and_stop("/api/codebook/prepare", {"path": str(rules)})
        self.assertEqual(response.status_code, 400, response.text)
        self.assertIn("cancelled", response.text.lower())
        self.assertLessEqual(calls, 2)  # the in-flight call; nothing after the Stop
        self.assertEqual(w.load_codebook_profile(), {})

    def test_stop_comparison_dispatches_no_further_model(self):
        self.post("/api/codebook/prepare", {"path": str(self.f.rules)})
        self.post("/api/codebook/lock", {})
        self.post("/api/target", {"target": str(self.f.bg_a), "selected_paths": ["facts.txt"]})
        self.post("/api/coverage", {})
        models = ["gemini-2.5-pro", "gemini-2.5-flash", "gemini-3.5-flash", "gemini-3.6-flash"]
        response, calls = self.run_and_stop("/api/compare", {"question": "classify", "models": models})
        self.assertEqual(response.status_code, 400, response.text)
        self.assertLessEqual(calls, 2)
        self.assertGreaterEqual(w.ACTIVE["state"].data["usage"]["requests"], 1)  # consumed usage kept

    def test_stop_coverage_writes_no_report(self):
        self.post("/api/codebook/prepare", {"path": str(self.f.rules)})
        self.post("/api/codebook/lock", {})
        for index in range(5):
            (self.f.bg_a / f"part_{index}.txt").write_text(f"FIXTURE part {index} rotates.", encoding="utf-8")
        self.post("/api/target", {"target": str(self.f.bg_a), "selected_paths": ["."]})
        response, calls = self.run_and_stop("/api/coverage", {})
        self.assertEqual(response.status_code, 400, response.text)
        self.assertLessEqual(calls, 2)
        self.assertFalse(w.coverage_path(w.ACTIVE["state"]).exists())

    def test_operation_cancel_validates_its_id(self):
        self.assertEqual(self.client.post("/api/operation/cancel", json={"request_id": "x"}).status_code, 422)
        self.assertEqual(self.client.post("/api/operation/cancel", json={"request_id": "b" * 32}).status_code, 404)


class TerminalDefaults(unittest.TestCase):
    def test_terminal_connector_gets_the_same_realistic_defaults(self):
        connector = LLMConnector("gemini-2.5-pro", "fixture-only")
        good = {"candidates": [{"content": {"parts": [{"text": "ok"}]}, "finishReason": "STOP"}]}
        seen = {}

        def post(*_args, **kwargs):
            seen["timeout"] = kwargs["timeout"]
            return SimpleNamespace(ok=True, json=lambda: good)
        with patch("requests.post", side_effect=post), patch.dict(os.environ, {}, clear=False):
            for name in ("REQUEST_TIMEOUT_SECONDS", "RESPONSE_TIMEOUT_SECONDS"):
                os.environ.pop("WORKBENCH_" + name, None)
            self.assertEqual(connector.ask_about_files([], "fixture"), "ok")
        self.assertGreater(seen["timeout"][1], 299)


if __name__ == "__main__":
    unittest.main(verbosity=2)

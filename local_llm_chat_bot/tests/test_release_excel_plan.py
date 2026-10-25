"""F-10: Excel edits are confirmed against a deterministic local plan summary (offline)."""

import contextlib
import hashlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import excel_chat_agent as chat


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class ExcelPlan(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="Excel Plan Prüfung 中文 ")
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        app_dir = patch.object(chat, "APP_DIR", self.root / "state")
        app_dir.start()
        self.addCleanup(app_dir.stop)
        self.workbook = self.root / "workshop.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "Workshop"
        sheet.append(["BG-ID", "Jonas", None, "Check"])
        for index in range(1, 13):
            sheet.append([f"BG-{index}", f"J{index}", None, f"=B{index + 1}"])
        book.save(self.workbook)
        self.results = self.root / "berk.csv"
        self.results.write_text("BG-ID,Judgement\n" + "\n".join(f"BG-{i},Rotation {i}" for i in range(1, 13)),
                                encoding="utf-8")
        self.payload = {"sheet": "Workshop", "header_row": 1, "id_header": "BG-ID", "template_header": "Jonas",
                        "target_column": "C", "new_header": "Berk", "results_id_column": "BG-ID",
                        "results_value_column": "Judgement"}

    def tools(self):
        session = chat.Session("plan-test")
        session.data.update({"workbook": str(self.workbook), "results": str(self.results)})
        session.save()
        return chat.WorkbookTools(session)

    def test_summary_is_built_from_the_plan_file(self):
        tools = self.tools()
        result = tools.plan_participants(self.payload)
        plan_path = Path(result["plan"])
        summary = tools.last_plan_summary
        self.assertEqual(result["plan_sha256"], digest(plan_path))
        self.assertEqual(tools.session.data["pending_plan_sha256"], digest(plan_path))
        for expected in ["工作表：Workshop", "目标列：C", '新表头："Berk"（Workshop!C1）',
                         "写入单元格总数：13（表头 1 个，判断值 12 个）", "前 10 条（单元格 → 值）：",
                         'Workshop!C2 → "Rotation 1"', 'Workshop!C11 → "Rotation 10"', "…其余 2 条见计划文件",
                         f"计划文件：{plan_path}", f"SHA-256：{digest(plan_path)}"]:
            self.assertIn(expected, summary)
        self.assertNotIn("Workshop!C12", summary)  # only the first ten value edits are listed

    def test_wrong_model_wording_cannot_change_the_printed_local_summary(self):
        session = chat.Session("main-loop")
        session.data.update({"workbook": str(self.workbook), "results": str(self.results)})
        session.save()
        decisions = iter([
            (json.dumps({"tool_calls": [{"name": "plan_participants", "arguments": self.payload}]}), False),
            ("计划已生成：将在 Z 列写入 999 个单元格，表头为 Alice。", False),
        ])
        lines = iter(["在 Workshop 表中按 BG-ID 添加 Berk 的判断", "/quit"])
        output = io.StringIO()
        with patch.object(sys, "argv", ["excel_chat_agent.py", "--session", "main-loop", "--model", "gpt-5.5"]), \
                patch.dict(os.environ, {"BOSCH_FARM_SUBSCRIPTION_KEY": "offline-fixture-only"}), \
                patch.object(chat, "load_dotenv"), \
                patch.object(chat, "ask_with_progress", side_effect=lambda *a, **k: next(decisions)), \
                patch("builtins.input", side_effect=lambda *_: next(lines)), contextlib.redirect_stdout(output):
            chat.main()
        text = output.getvalue()
        plan_path = Path(chat.Session("main-loop").data["pending_plan"])
        self.assertIn("将在 Z 列写入 999 个单元格", text)  # the model's (wrong) wording is shown ...
        summary_at = text.index("【本地计划摘要")
        self.assertGreater(summary_at, text.index("将在 Z 列写入 999 个单元格"))  # ... then the authoritative summary
        local = text[summary_at:]
        self.assertIn("目标列：C", local)
        self.assertIn("写入单元格总数：13", local)
        self.assertIn('新表头："Berk"', local)
        self.assertIn(f"SHA-256：{digest(plan_path)}", local)

    def test_tampered_plan_is_refused_and_nothing_is_written(self):
        tools = self.tools()
        plan_path = Path(tools.plan_participants(self.payload)["plan"])
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        plan["edits"][1]["value"] = "UNREVIEWED"
        plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
        before = digest(self.workbook)
        with self.assertRaisesRegex(ValueError, "SHA-256 不一致"):
            tools.apply_pending()
        self.assertEqual(list(self.root.glob("workshop_ai_updated_*")), [])
        self.assertEqual(digest(self.workbook), before)
        self.assertIsNone(tools.session.data["pending_plan"])
        with self.assertRaisesRegex(ValueError, "没有待执行计划"):
            tools.apply_pending()

    def test_plan_without_a_shown_summary_is_refused(self):
        tools = self.tools()
        tools.plan_participants(self.payload)
        tools.session.data["pending_plan_sha256"] = None  # e.g. a session saved by an older version
        with self.assertRaisesRegex(ValueError, "重新生成计划"):
            tools.apply_pending()
        self.assertEqual(list(self.root.glob("workshop_ai_updated_*")), [])

    def test_confirmed_plan_writes_a_verified_copy_and_keeps_the_source(self):
        tools = self.tools()
        tools.plan_participants(self.payload)
        before = digest(self.workbook)
        result = tools.apply_pending()
        output = Path(result["output"])
        self.assertTrue(output.is_file())
        self.assertNotEqual(output, self.workbook)
        self.assertTrue(result["verification"]["ok"], result["verification"])
        self.assertEqual(result["verification"]["checked_cells"], 13)
        self.assertEqual(digest(self.workbook), before)
        sheet = load_workbook(output)["Workshop"]
        self.assertEqual(sheet["C1"].value, "Berk")
        self.assertEqual(sheet["C13"].value, "Rotation 12")
        self.assertEqual(sheet["D2"].value, "=B2")  # formulas untouched
        self.assertEqual(list(tools.session.root.glob(".reviewed_plan_*")), [])

    def test_safety_refusals_are_reported_without_ending_the_agent(self):
        tools = self.tools()
        tools.plan_participants(self.payload)
        book = load_workbook(self.workbook)
        book["Workshop"]["B2"] = "edited after planning"
        book.save(self.workbook)
        with self.assertRaisesRegex(ValueError, "Workbook changed after planning"):
            tools.apply_pending()  # previously excel_agent.die() raised SystemExit here
        occupied = dict(self.payload, target_column="B")
        with self.assertRaisesRegex(ValueError, "not empty"):
            tools.plan_participants(occupied)


if __name__ == "__main__":
    unittest.main(verbosity=2)

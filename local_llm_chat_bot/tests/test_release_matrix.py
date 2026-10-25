"""Offline transport, HTTP API, parser and write-failure release matrix."""
import contextlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import threading
from types import SimpleNamespace
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from fastapi.testclient import TestClient
from openpyxl import Workbook
import requests
import chat_workbench as w
import excel_agent as excel
from code_chat_agent import PricingCatalog
from llm_connector_with_prompt_caching import LLMConnector
from release_safety import normalize_usage
import test_release_regressions as regression


class TransportMatrix(unittest.TestCase):
    def test_http_status_retry_matrix(self):
        for status in [400, 401, 403, 404, 413, 429, 500, 502, 503, 504]:
            with self.subTest(status=status):
                llm = LLMConnector('gemini-test', 'private-fixture-key')
                response = SimpleNamespace(ok=False, status_code=status, text='private-fixture-key')
                with patch('requests.post', return_value=response) as post, patch('request_runtime.RequestControl.pause') as sleep:
                    answer = llm.ask_about_files([], 'fixture')
                self.assertEqual(post.call_count, 3 if status in {429, 500, 502, 503, 504} else 1)
                self.assertEqual(sleep.call_count, post.call_count - 1)
                self.assertIsNotNone(w.connector_failure(answer))
                self.assertNotIn('private-fixture-key', answer)

    def test_timeout_connection_invalid_json_empty_and_recovery(self):
        # The request never reached the provider: bounded retries are safe.
        # A read timeout may still be generating/billing: never re-dispatch it.
        for error, dispatches in ((requests.ConnectionError('connection'), 3),
                                  (requests.exceptions.ConnectTimeout('connect'), 3),
                                  (requests.exceptions.ReadTimeout('read'), 1),
                                  (requests.Timeout('timeout'), 1)):
            llm = LLMConnector('gemini-test', 'fixture')
            with patch('requests.post', side_effect=error) as post, patch('request_runtime.RequestControl.pause'):
                self.assertIsNotNone(w.connector_failure(llm.ask_about_files([], 'fixture')))
                self.assertEqual(post.call_count, dispatches)
        llm = LLMConnector('gemini-test', 'fixture')
        with patch('requests.post', return_value=SimpleNamespace(ok=True, json=lambda: [])) as post:
            self.assertIsNotNone(w.connector_failure(llm.ask_about_files([], 'fixture')))
            self.assertEqual(post.call_count, 1)
        good = {'candidates': [{'content': {'parts': [{'text': 'recovered'}]}, 'finishReason': 'STOP'}]}
        with patch('requests.post', return_value=SimpleNamespace(ok=True, json=lambda: good)):
            self.assertEqual(llm.ask_about_files([], 'fixture'), 'recovered')

    def test_provider_usage_missing_zero_invalid_and_thoughts(self):
        for raw in ({'prompt_tokens': 12, 'completion_tokens': 3}, {'input_tokens': 12, 'output_tokens': 3},
                    {'promptTokenCount': 12, 'candidatesTokenCount': 3}):
            u = normalize_usage(raw); self.assertEqual((u['input_tokens'], u['output_tokens'], u['total_tokens']), (12, 3, 15))
        self.assertIsNone(normalize_usage({})['input_tokens'])
        self.assertEqual(normalize_usage({'prompt_tokens': 0, 'completion_tokens': 0})['total_tokens'], 0)
        for value in (-1, 'bad', float('nan'), float('inf')):
            self.assertIsNone(normalize_usage({'input_tokens': value})['input_tokens'])
        self.assertEqual(normalize_usage({'candidatesTokenCount': 3, 'thoughtsTokenCount': 9})['output_tokens'], 12)

    def test_provider_truncation_and_empty_not_valid(self):
        for reason in ('MAX_TOKENS', 'SAFETY'):
            answer = LLMConnector._extract_gemini_text({'candidates': [{'finishReason': reason, 'content': {'parts': [{'text': '{"x":1}'}]}}]})
            self.assertIsNotNone(w.connector_failure(answer))
        llm = LLMConnector('claude-test', 'fixture')
        self.assertIsNotNone(w.connector_failure(llm._extract_claude_text({'stop_reason': 'max_tokens', 'content': [{'text': '{}'}]})))
        self.assertIsNotNone(w.connector_failure(LLMConnector._extract_gemini_text({})))


class ApiAndFileMatrix(unittest.TestCase):
    setUp = regression.ReleaseRegressionTests.setUp
    tearDown = regression.ReleaseRegressionTests.tearDown

    def test_invalid_paths_blank_question_long_question_and_missing_key(self):
        with TestClient(w.app, raise_server_exceptions=False) as client:
            for payload in ({'target': ''}, {'target': str(self.root/'missing')}):
                self.assertEqual(client.post('/api/target', json=payload).status_code, 400)
            self.assertEqual(client.post('/api/chat', json={'question': ' '}).status_code, 400)
            self.assertEqual(client.post('/api/chat', json={'question': 'x'*50001}).status_code, 422)
            with patch.dict(os.environ, {'WORKBENCH_ALLOW_TEST_HOST': '1'}, clear=True):
                result = client.post('/api/setup', json={'workspace': str(self.workspace), 'model': 'gemini-test'})
                self.assertEqual(result.status_code, 400)
                self.assertIn('BOSCH_FARM_SUBSCRIPTION_KEY', result.json()['detail'])

    def test_model_switch_canonicalization_and_request(self):
        with TestClient(w.app) as client:
            for model in [' GEMINI-TEST ', 'claude-fixture', 'gpt-fixture']:
                result = client.post('/api/model', json={'model': model})
                self.assertEqual(result.status_code, 200)
                self.assertEqual(w.ACTIVE['llm'].model_name, model.strip().lower())
            for model in ['', 'gpt with spaces', 'unknown']:
                self.assertEqual(client.post('/api/model', json={'model': model}).status_code, 400)

    def test_model_slot_rejects_context_mutation_and_recovers(self):
        with TestClient(w.app) as client:
            w.MODEL_WORK_IN_PROGRESS.set()
            for route, payload in [('/api/model', {'model':'gemini-test'}), ('/api/target', {'target':str(self.root)}), ('/api/session/new', {}), ('/api/chat', {'question':'x'})]:
                self.assertEqual(client.post(route, json=payload).status_code, 409)
            w.MODEL_WORK_IN_PROGRESS.clear()
            with patch.object(w, 'call_model', side_effect=RuntimeError('fixture failure')):
                self.assertEqual(client.post('/api/chat', json={'question':'x'}).status_code, 400)
                self.assertFalse(w.MODEL_WORK_IN_PROGRESS.is_set())

    def test_cross_origin_and_host_rejected(self):
        with TestClient(w.app) as client:
            self.assertEqual(client.post('/api/model', json={'model':'gemini-test'}, headers={'Origin':'https://untrusted.example'}).status_code, 403)
            self.assertEqual(client.get('/api/settings', headers={'Host':'untrusted.example'}).status_code, 403)

    def test_all_declared_file_formats_and_source_immutability(self):
        import fitz
        from docx import Document
        from PIL import Image
        source = self.root/'材料 ü'; source.mkdir()
        pdf = fitz.open(); pdf.new_page().insert_text((30,30), 'Linear motor'); pdf.save(source/'fixture.pdf'); pdf.close()
        doc = Document(); doc.add_paragraph('Linear motor'); doc.save(source/'fixture.docx')
        book = Workbook(); book.active['A1']='part'; book.active['B2']='=1+1'; book.active.merge_cells('C1:D1'); book.save(source/'fixture.xlsx'); book.close()
        for ext in ['png','jpg','jpeg','webp']: Image.new('RGB', (64,64), 'white').save(source/f'fixture.{ext}')
        for ext in ['csv','txt','md']: (source/f'fixture.{ext}').write_text('motor,linear\n1,2', encoding='utf-8')
        before = {p.name:p.read_bytes() for p in source.iterdir()}
        w.analyze_target(w.TargetRequest(target=str(source), selected_paths=['.']))
        self.assertEqual(len(w.ACTIVE['state'].data['active_source_ids']), len(before))
        self.assertEqual(before, {p.name:p.read_bytes() for p in source.iterdir()})

    def test_empty_corrupt_oversize_nested_duplicate_selection(self):
        folder=self.root/'BG'; (folder/'nested').mkdir(parents=True)
        (folder/'nested'/'good.txt').write_text('technical motor')
        (folder/'bad.pdf').write_bytes(b'broken PDF')
        (folder/'empty.txt').write_text('')
        book=Workbook(); book.save(folder/'blank.xlsx'); book.close()
        x=w.analyze_target(w.TargetRequest(target=str(folder),selected_paths=['.','nested/good.txt']))
        self.assertEqual(x['sources'],1); self.assertEqual(x['skipped_count'],3)
        self.assertTrue(any('bad.pdf' in e for e in x['skipped']))
        with self.assertRaises(w.HTTPException): w.analyze_target(w.TargetRequest(target=str(folder),selected_paths=[]))
        from local_document_library import MAX_DOCUMENT_BYTES
        large=folder/'large.txt'
        with large.open('wb') as handle: handle.truncate(MAX_DOCUMENT_BYTES+1)
        with self.assertRaisesRegex(ValueError,'larger'): w.ACTIVE['documents'].add(str(large))

    def test_no_source_scope_means_no_old_documents(self):
        source=self.root/'old.txt';source.write_text('SECRET_OLD motor')
        w.ACTIVE['documents'].add(str(source))
        self.assertEqual(w.ACTIVE['documents'].search('motor', allowed_ids=[]), [])

    def test_pending_cancel_confirm_backup_and_rollback(self):
        a=self.workspace/'a.txt'; b=self.workspace/'b.txt'; a.write_text('A'); b.write_text('B')
        state=w.ACTIVE['state']
        changes=w.build_diff(state,[{'path':'a.txt','content':'new A'},{'path':'b.txt','content':'new B'}])
        plan_id = w.publish_pending(w.ACTIVE, changes)
        with self.assertRaises(w.HTTPException): w.apply(w.ApplyRequest(confirmed=False))
        self.assertEqual(a.read_text(),'A')
        real_replace=os.replace
        def fail_second(source,dest):
            if Path(dest)==b: raise PermissionError('file locked')
            return real_replace(source,dest)
        with patch.object(w.os,'replace',side_effect=fail_second), self.assertRaises(PermissionError):
            w.apply(w.ApplyRequest(confirmed=True, plan_id=plan_id))
        self.assertEqual(a.read_text(),'A'); self.assertEqual(b.read_text(),'B')
        self.assertTrue(list((state.root/'backups').rglob('a.txt')))
        w.cancel_pending(); self.assertEqual(w.ACTIVE['pending'],[])

    def test_atomic_excel_publication_preserves_formulas_and_failure(self):
        source=self.root/'source.xlsx'; output=self.root/'result.xlsx'
        book=Workbook(); sheet=book.active; sheet.title='Data'; sheet['A1']='ID';sheet['A2']='1';sheet['B1']='Template';sheet['B2']='x';sheet['D2']='=A2';sheet.merge_cells('E1:F1');book.save(source);book.close()
        book=excel.load_book(source)
        plan={'source_sha256':excel.sha256(source),'source_formula_fingerprint':excel.formula_fingerprint(book),
              'sheet':'Data','target_column':'C','template_column':'B','header_row':1,
              'edits':[{'cell':'Data!C1','value':'New'},{'cell':'Data!C2','value':'Gantry'}]}
        book.close(); path=self.root/'plan.json';path.write_text(json.dumps(plan))
        args=SimpleNamespace(workbook=str(source),output=str(output),plan=str(path),force=False)
        with contextlib.redirect_stdout(io.StringIO()): excel.command_apply(args)
        self.assertEqual(excel.sha256(source),plan['source_sha256']);self.assertTrue(excel.verify_plan(output,plan)['ok'])
        output.unlink()
        with patch('openpyxl.workbook.workbook.Workbook.save',side_effect=PermissionError('locked')), self.assertRaises(PermissionError): excel.command_apply(args)
        self.assertFalse(output.exists()); self.assertEqual(excel.sha256(source),plan['source_sha256'])
        self.assertFalse(list(self.root.glob('.result.*')))


if __name__=='__main__': unittest.main()

"""Chinese conversational Excel agent with exact, locally controlled workbook tools.

The LLM may inspect data through a small allow-list of tools.  It never receives
permission to write a workbook directly: it can only create a reviewable plan;
the user must type "确认执行" before a new output workbook is produced.
"""
from __future__ import annotations

import argparse
import contextlib
import hashlib
import io
import json
import os
import re
import sys
import threading
import time
import uuid
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace
from typing import Any

from openpyxl.utils.cell import range_boundaries

from excel_agent import (
    cell_record,
    command_apply,
    command_plan_participants,
    excel_json_value,
    load_book,
    resolve_sheet,
    special_formula_ranges,
    stored_cells,
    trace_cell,
    verify_plan,
    worksheet_overview,
)

try:
    from dotenv import load_dotenv
except ImportError:
    def load_dotenv(*_args, **_kwargs):
        return False


APP_DIR = Path(__file__).resolve().parent / ".excel_chat_agent"
SYSTEM = "你是一个谨慎的 Excel 助手。始终用中文回答。只能依据本地精确工具返回的单元格、公式和计划；不得猜测单元格位置或文件内容。写入前必须清楚展示计划并等待用户输入‘确认执行’。"
TOOL_LABELS = {
    "discover_workbooks": "搜索目录", "open_workbook": "打开选中的文件",
    "overview": "读取工作簿概况", "sheet_preview": "读取工作表预览",
    "read_range": "读取指定单元格", "search": "搜索单元格内容",
    "trace": "检查公式引用", "plan_participants": "生成参与者修改计划",
}
QUOTED_RE = re.compile(r'["“]([^"”\n]+)["”]')
WINDOWS_PATH_RE = re.compile(r'([A-Za-z]:\\[^"“”\r\n]+)')
NAME_HINT_RE = re.compile(r'(?:名字叫|名为|叫)\s*["“]?([^"“”\s]+?)(?:的)?(?:excel|xlsx|xlsm|工作簿|表格|文件)', re.IGNORECASE)


PLAN_SUMMARY_EDITS = 10


def run_exact(command, args) -> str:
    """Run an excel_agent command; turn its die() exit into an ordinary refusal.

    excel_agent.die() raises SystemExit for every safety refusal. Inside the
    interactive agent that must become a reported "not written" error, never
    an exit of the whole conversation.
    """
    output, errors = io.StringIO(), io.StringIO()
    try:
        with contextlib.redirect_stdout(output), contextlib.redirect_stderr(errors):
            command(args)
    except SystemExit as exc:
        if exc.code not in (None, 0):
            message = errors.getvalue().strip()
            raise ValueError(message[len("ERROR: "):] if message.startswith("ERROR: ") else message or f"exit {exc.code}") from None
    return output.getvalue()


def plan_digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def plan_summary(plan: dict[str, Any], plan_path: Path, digest: str) -> str:
    """Deterministic review text built from the plan file itself, never by the model."""
    edits = plan.get("edits", [])
    header = next((edit for edit in edits if edit.get("kind") == "header"), None)
    values = [edit for edit in edits if edit.get("kind") != "header"]

    def shown(value: Any) -> str:
        return json.dumps(excel_json_value(value), ensure_ascii=False)
    lines = ["【本地计划摘要｜由程序直接读取计划文件生成，以此为准；模型的描述仅供参考】",
             f"工作表：{plan.get('sheet', '—')}",
             f"目标列：{plan.get('target_column', '—')}",
             f"新表头：{shown(header['value']) if header else '—'}（{header['cell'] if header else '—'}）",
             f"写入单元格总数：{len(edits)}（表头 {1 if header else 0} 个，判断值 {len(values)} 个）",
             f"前 {min(PLAN_SUMMARY_EDITS, len(values))} 条（单元格 → 值）："]
    lines += [f"  {edit['cell']} → {shown(edit['value'])}" for edit in values[:PLAN_SUMMARY_EDITS]]
    if len(values) > PLAN_SUMMARY_EDITS:
        lines.append(f"  …其余 {len(values) - PLAN_SUMMARY_EDITS} 条见计划文件")
    lines += [f"计划文件：{plan_path}", f"SHA-256：{digest}",
              "核对无误后输入“确认执行”生成新副本；输入 /cancel 取消。原文件不会被修改。"]
    return "\n".join(lines)


def compact(value: Any, length: int = 220) -> str:
    text = "" if value is None else str(value)
    return text if len(text) <= length else text[: length - 1] + "…"


def parse_json_reply(text: str) -> dict[str, Any] | None:
    cleaned = text.strip().replace("```json", "").replace("```", "")
    decoder = json.JSONDecoder()
    for offset, char in enumerate(cleaned):
        if char != "{":
            continue
        try:
            value, _ = decoder.raw_decode(cleaned[offset:])
            if isinstance(value, dict):
                return value
        except json.JSONDecodeError:
            pass
    return None


class TerminalProgress:
    """Show actual stages and elapsed time, independently of model output."""

    def __init__(self, interval: float = 10.0):
        self.interval = interval
        self.output = sys.stdout
        self._lock = threading.Lock()
        self._stream_kind: str | None = None
        self._last_activity = time.monotonic()

    def _end_stream(self) -> None:
        if self._stream_kind is not None:
            print(file=self.output, flush=True)
            self._stream_kind = None

    def note(self, message: str) -> None:
        with self._lock:
            self._end_stream()
            print(f"[进度] {message}", file=self.output, flush=True)
            self._last_activity = time.monotonic()

    def chunk(self, kind: str, text: str) -> None:
        if not text:
            return
        with self._lock:
            if self._stream_kind != kind:
                self._end_stream()
                print("接口思考> " if kind == "reasoning" else "assistant> ", end="", file=self.output, flush=True)
                self._stream_kind = kind
            print(text, end="", file=self.output, flush=True)
            self._last_activity = time.monotonic()

    @contextlib.contextmanager
    def phase(self, label: str):
        self.note(label)
        started = time.monotonic()
        stopped = threading.Event()

        def heartbeat():
            while not stopped.wait(self.interval):
                with self._lock:
                    if time.monotonic() - self._last_activity < self.interval:
                        continue
                    self._end_stream()
                    elapsed = time.monotonic() - started
                    print(f"[进度] {label}，已用 {elapsed:.0f} 秒，仍在等待完成…", file=self.output, flush=True)
                    self._last_activity = time.monotonic()

        worker = threading.Thread(target=heartbeat, daemon=True)
        worker.start()
        state = "完成"
        try:
            yield
        except BaseException:
            state = "中止"
            raise
        finally:
            stopped.set()
            worker.join(timeout=0.5)
            self.note(f"{label}：{state}，{time.monotonic() - started:.1f} 秒")


def ask_with_progress(llm, prompt: str, label: str, tokens: int, args, progress: TerminalProgress,
                      *, show_answer: bool = False) -> tuple[str, bool]:
    rendered_answer = False

    def on_event(kind, text):
        nonlocal rendered_answer
        if kind == "notice":
            progress.note(text)
        elif kind == "reasoning" and args.show_reasoning:
            progress.chunk(kind, text)
        elif kind == "content" and show_answer:
            progress.chunk(kind, text)
            rendered_answer = True

    with progress.phase(label):
        answer = llm.ask_about_files(
            [], prompt, SYSTEM, {"maxOutputTokens": tokens},
            stream=args.stream or (args.show_reasoning and not args.no_stream),
            event_handler=on_event, thinking=args.thinking,
        )
        if answer.startswith(("Error:", "HTTP error", "Error after")):
            raise RuntimeError(answer)
    if args.show_reasoning and not llm.last_reasoning:
        progress.note("本次接口没有返回 reasoning_content；无法显示本次思考文本。")
    return answer, rendered_answer


class Session:
    def __init__(self, name: str):
        self.root = APP_DIR / name
        self.root.mkdir(parents=True, exist_ok=True)
        self.path = self.root / "session.json"
        self.data = json.loads(self.path.read_text(encoding="utf-8")) if self.path.exists() else {"workbook": None, "results": None, "pending_plan": None, "discovered_files": [], "history": []}

    def save(self) -> None:
        self.path.write_text(json.dumps(self.data, ensure_ascii=False, indent=2), encoding="utf-8")

    def add(self, role: str, text: str) -> None:
        self.data["history"].append({"role": role, "text": text})
        self.data["history"] = self.data["history"][-12:]
        self.save()


class WorkbookTools:
    def __init__(self, session: Session):
        self.session = session
        self.allowed_roots: list[Path] = []
        # Set when a plan is created; the main loop prints it after the model's reply.
        self.last_plan_summary: str | None = None

    def clear_pending(self) -> None:
        self.session.data["pending_plan"] = None
        self.session.data["pending_plan_sha256"] = None

    @property
    def path(self) -> Path:
        raw = self.session.data.get("workbook")
        if not raw:
            raise ValueError("尚未打开工作簿。请先输入：打开 \"C:\\路径\\文件.xlsx\"")
        path = Path(raw)
        if not path.is_file():
            raise ValueError(f"当前工作簿已不存在：{path}")
        return path

    def book(self):
        return load_book(self.path, data_only=False)

    def overview(self) -> dict[str, Any]:
        return worksheet_overview(self.book())

    def discover_workbooks(self, root: str, name_hint: str = "", purpose: str = "workbook") -> dict[str, Any]:
        folder = Path(root).expanduser().resolve()
        if not folder.is_dir():
            raise ValueError(f"该目录不存在：{folder}")
        if not self.allowed_roots:
            raise ValueError("请先在当前消息中提供需要搜索的目录路径。")
        try:
            permitted = any(os.path.commonpath([str(folder), str(allowed)]) == str(allowed) for allowed in self.allowed_roots)
        except ValueError:
            permitted = False
        if not permitted:
            raise ValueError("只能搜索你在当前消息中明确提供的目录或其子目录。")
        include_csv = purpose == "results"
        candidates = workbook_candidates(folder, name_hint, include_csv=include_csv)
        self.session.data["discovered_files"] = [str(item.resolve()) for item in candidates]
        self.session.save()
        return {"root": str(folder), "name_hint": name_hint, "purpose": purpose, "candidates": [{"path": str(item.resolve()), "name": item.name} for item in candidates[:30]], "candidate_count": len(candidates)}

    def open_workbook(self, path: str, role: str = "workbook") -> dict[str, Any]:
        candidate = Path(path).expanduser().resolve()
        known = {Path(item).resolve() for item in self.session.data.get("discovered_files", [])}
        if candidate not in known:
            raise ValueError("只能打开本轮目录搜索返回的候选文件。")
        if role == "results":
            if candidate.suffix.lower() not in {".xlsx", ".xlsm", ".csv"}:
                raise ValueError("结果文件必须是 .xlsx、.xlsm 或 .csv。")
            self.session.data["results"] = str(candidate)
            self.session.save()
            return {"opened_results_file": candidate.name, "path": str(candidate)}
        if candidate.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError("工作簿必须是 .xlsx 或 .xlsm 文件。")
        book = load_book(candidate, data_only=False)
        self.session.data["workbook"] = str(candidate)
        self.clear_pending()
        self.session.save()
        return {"opened_workbook": candidate.name, "path": str(candidate), "overview": worksheet_overview(book)}

    def sheet_preview(self, selector: str, rows: int = 35, cols: int = 20) -> dict[str, Any]:
        book = self.book()
        sheet = resolve_sheet(book, str(selector))
        cells = []
        definitions = special_formula_ranges(sheet)
        for row in sheet.iter_rows(min_row=1, max_row=min(rows, sheet.max_row), max_col=min(cols, sheet.max_column)):
            for cell in row:
                if cell.value is not None:
                    cells.append(cell_record(cell, definitions))
        return {"sheet": sheet.title, "dimensions": sheet.calculate_dimension(), "preview": f"A1:{sheet.cell(min(rows, sheet.max_row), min(cols, sheet.max_column)).coordinate}", "cells": cells}

    def read_range(self, selector: str, address: str) -> dict[str, Any]:
        book = self.book()
        sheet = resolve_sheet(book, str(selector))
        try:
            min_col, min_row, max_col, max_row = range_boundaries(address)
        except ValueError as exc:
            raise ValueError(f"无效单元格区域：{exc}") from exc
        if (max_col - min_col + 1) * (max_row - min_row + 1) > 500:
            raise ValueError("一次最多读取 500 个单元格。请缩小区域。")
        cells = []
        definitions = special_formula_ranges(sheet)
        for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                cells.append(cell_record(cell, definitions))
        return {"sheet": sheet.title, "range": address, "cells": cells}

    def trace(self, qualified_cell: str) -> dict[str, Any]:
        if "!" not in qualified_cell:
            raise ValueError("公式追踪必须使用 Sheet!A1，例如 Workshop!D37。")
        sheet_name, address = qualified_cell.rsplit("!", 1)
        book = self.book()
        sheet = resolve_sheet(book, sheet_name)
        return trace_cell(sheet, address)

    def search(self, query: str, selector: str | None = None, limit: int = 40) -> dict[str, Any]:
        query = query.strip().lower()
        if len(query) < 2:
            raise ValueError("搜索词至少需要两个字符。")
        book = self.book()
        sheets = [resolve_sheet(book, str(selector))] if selector else book.worksheets
        matches = []
        for sheet in sheets:
            definitions = special_formula_ranges(sheet)
            for cell in stored_cells(sheet):
                if cell.value is not None and query in json.dumps(excel_json_value(cell.value), ensure_ascii=False).lower():
                    matches.append(cell_record(cell, definitions))
                    if len(matches) >= limit:
                        return {"query": query, "matches": matches, "truncated": True}
        return {"query": query, "matches": matches, "truncated": False}

    def plan_participants(self, payload: dict[str, Any]) -> dict[str, Any]:
        results = self.session.data.get("results")
        if not results:
            raise ValueError("还没有参与者结果文件。请先输入：结果文件 \"C:\\路径\\results.xlsx\"")
        required = ["sheet", "header_row", "id_header", "template_header", "target_column", "new_header", "results_id_column", "results_value_column"]
        missing = [key for key in required if not payload.get(key)]
        if missing:
            raise ValueError("模型尚不能安全确定这些字段：" + ", ".join(missing))
        plan_path = self.session.root / f"pending_plan_{uuid.uuid4().hex[:8]}.json"
        args = SimpleNamespace(
            workbook=str(self.path), results=str(results), plan=str(plan_path), force=False,
            sheet=str(payload["sheet"]), header_row=int(payload["header_row"]), id_header=str(payload["id_header"]),
            template_header=str(payload["template_header"]), target_column=str(payload["target_column"]), new_header=str(payload["new_header"]),
            results_id_column=str(payload["results_id_column"]), results_value_column=str(payload["results_value_column"]),
        )
        run_exact(command_plan_participan